from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


BASE = Path("D:/table")
EXCEL_PATH = BASE / "\u6570\u7267\u589e\u6548-\u786c\u4ef6\u91c7\u8d2d\u6e05\u53555.xlsx"
TEMPLATE_PATH = BASE / "\u7269\u6599\u7533\u8bf7\u8868(3).docx"
OUTPUT_PATH = BASE / "\u7269\u6599\u7533\u8bf7\u8868(3)-\u6570\u7267\u589e\u6548\u5b8c\u6210\u7248.docx"


PURPOSE_FALLBACKS = {
    "MQ-135 空气质量传感器模块": "检测烟雾、挥发性气体及空气质量变化，与氨气、火焰等数据共同形成牛舍环境安全预警依据。",
    "MQ137氨气模块": "采集牛舍氨气浓度，用于判断通风不足、垫料污染或环境异常，并触发通风和告警联动。",
    "INMP441麦克风模块": "采集牛舍环境声音，用于后续进行异常叫声、设备噪声和活动状态识别算法验证。",
    "1.3寸OLED4针（蓝色显示）": "本地显示联网状态、传感器读数和告警信息，便于现场调试和演示时快速查看节点运行状态。",
    "新款 ADXL345 加速度传感器": "采集运动和倾斜数据，用于活动量、姿态变化、跌倒或异常行为判断等原型验证。",
    "GY-MAX30102": "用于心率、血氧等生理指标采集原型验证，为后续扩展牛只健康监测功能提供硬件基础。",
    "安服优B-1红外测温传感器模块": "进行非接触体温或物体温度采集，辅助牛只健康状态判断和环境异常排查。",
    "聚合物锂电池": "为移动采集节点和便携式原型设备供电，便于脱离固定电源进行场景测试和演示。",
    "万用表": "用于电压、电阻、线路连通性等基础检测，保障传感器、供电模块和执行器调试安全可靠。",
    "四路继电器模块，带光耦隔离": "控制风扇、报警器、雾化片等外设通断，并通过光耦隔离提升控制侧与负载侧的安全性。",
    "830孔面包板 带背胶可拼接": "用于原型阶段快速搭建电路、验证接线和调整传感器组合，减少反复焊接成本。",
    "杜邦线21CM 公对公 2.54mm（1排40P）": "用于开发板、面包板和传感器模块之间快速连接，满足原型搭建和调试需要。",
    "杜邦线21CM 母对母 2.54mm（1排40P）": "用于开发板、面包板和传感器模块之间快速连接，满足原型搭建和调试需要。",
    "杜邦线21CM 公对母 2.54mm（1排40P）": "用于开发板、面包板和传感器模块之间快速连接，满足原型搭建和调试需要。",
    "电烙铁BY-55W": "用于焊接排针、接口和供电线路，满足传感器模块装配、返修及现场调试需要。",
    "60w刀头": "作为电烙铁耗材备件，用于焊接、返修和调试过程中替换磨损刀头，保证焊接质量。",
}


def load_items():
    from openpyxl import load_workbook

    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col = {name: idx for idx, name in enumerate(headers)}

    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = clean(row[col["名称"]]) if col.get("名称") is not None else ""
        if not name:
            continue
        link = clean(row[col["链接"]]) if col.get("链接") is not None else ""
        spec = clean(row[col["颜色分类"]]) if col.get("颜色分类") is not None else ""
        qty = clean(row[col["个数"]]) if col.get("个数") is not None else ""
        purpose = extract_purpose(link) or PURPOSE_FALLBACKS.get(name, "用于数牧增效项目硬件原型搭建、功能调试和比赛展示。")
        items.append({"name": name, "spec": spec, "qty": qty, "purpose": purpose})
    return items


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return re.sub(r"\s+", " ", text)


def extract_purpose(text):
    if not text:
        return ""
    match = re.search(r"\u7528\u9014\s*[:\uff1a]\s*(.+)", text)
    if not match:
        return ""
    purpose = match.group(1).strip()
    purpose = re.split(r"(?:\r?\n|\s{2,}|https?://)", purpose)[0].strip(" ;；")
    return purpose.rstrip("。") + "。"


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_paragraph_before(anchor, style_name, p_pr=None):
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    paragraph.style = style_name
    if p_pr is not None:
        existing = paragraph._p.pPr
        if existing is not None:
            paragraph._p.remove(existing)
        paragraph._p.insert(0, deepcopy(p_pr))
    return paragraph


def format_lead(item):
    if item["spec"]:
        return f"{item['name']}\uff08{item['spec']}\uff0c\u91c7\u8d2d\u6570\u91cf\uff1a{item['qty']}\uff09"
    return f"{item['name']}\uff08\u91c7\u8d2d\u6570\u91cf\uff1a{item['qty']}\uff09"


def build_document(items):
    doc = Document(TEMPLATE_PATH)
    paragraphs = doc.paragraphs

    start_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "\u91c7\u8d2d\u786c\u4ef6\u53ca\u4f7f\u7528\u65b9\u5f0f")
    end_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "\u9884\u8ba1\u53c2\u4e0e\u6bd4\u8d5b")
    sample = paragraphs[start_idx + 1] if end_idx > start_idx + 1 else None
    sample_ppr = deepcopy(sample._p.pPr) if sample is not None and sample._p.pPr is not None else None
    style_name = sample.style.name if sample is not None else "Normal"

    for paragraph in list(paragraphs[start_idx + 1 : end_idx]):
        delete_paragraph(paragraph)

    anchor = next(p for p in doc.paragraphs if p.text.strip() == "\u9884\u8ba1\u53c2\u4e0e\u6bd4\u8d5b")
    for item in items:
        paragraph = insert_paragraph_before(anchor, style_name, sample_ppr)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = None

        lead = paragraph.add_run(format_lead(item))
        lead.bold = True
        body = item["purpose"].strip()
        paragraph.add_run(f"\uff1a{body}")

    doc.save(OUTPUT_PATH)


def verify(items):
    doc = Document(OUTPUT_PATH)
    text = "\n".join(p.text for p in doc.paragraphs)
    missing = [item["name"] for item in items if item["name"] not in text]
    competition_rows = []
    if doc.tables:
        competition_rows = [row.cells[0].text.strip() for row in doc.tables[0].rows]
    print(f"output={OUTPUT_PATH}")
    print(f"items={len(items)}")
    print(f"paragraphs={len(doc.paragraphs)}")
    print(f"tables={len(doc.tables)}")
    print(f"competition_rows={len([r for r in competition_rows if r])}")
    if missing:
        raise RuntimeError("Missing items: " + ", ".join(missing))


if __name__ == "__main__":
    items = load_items()
    build_document(items)
    verify(items)
