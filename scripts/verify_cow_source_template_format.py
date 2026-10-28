from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document
from lxml import etree as ET
from PIL import Image, ImageDraw
from pypdf import PdfReader

from cow_source_template_format import REFERENCE, NS
from reformat_cow_source_materials import WORK


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clean(text):
    return re.sub(r'\s+', '', text)


def body_rows(page):
    positions = {}
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            if line['bbox'][1] <= 65 or line['bbox'][0] < 80:
                continue
            for span in line['spans']:
                baseline = round(span['origin'][1], 1)
                positions.setdefault(baseline, []).append(span)
    return [(y, ''.join(span['text'] for span in sorted(spans, key=lambda s: s['origin'][0])),
             max(span['bbox'][2] for span in spans)) for y, spans in sorted(positions.items())]


def contact(paths, destination):
    width, height = 270, 410
    sheet = Image.new('RGB', (width * 4, height * ((len(paths) + 3) // 4)), '#DEE4ED')
    draw = ImageDraw.Draw(sheet)
    for index, path in enumerate(paths):
        picture = Image.open(path).convert('RGB')
        picture.thumbnail((width - 10, height - 27))
        x, y = index % 4 * width, index // 4 * height
        sheet.paste(picture, (x + (width - picture.width) // 2, y + 24))
        draw.text((x + 8, y + 6), path.stem, fill='black')
    sheet.save(destination)


def verify(item, published=False, render=False):
    path = Path(item['target'] if published else item['staged'])
    listing = json.loads((Path(item['target']).parent / '源程序文件与行号清单.json').read_text('utf-8'))
    expected = [row['text'].expandtabs(2) for row in listing['selected']]
    document = Document(path)
    actual = [paragraph.text for paragraph in document.paragraphs]
    assert actual == expected, (path, 'Source content changed')
    assert not any(re.match(r'^\d{5}\s*\|', row) for row in actual)
    reference = Document(REFERENCE)
    for name in ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin', 'header_distance', 'footer_distance']:
        assert getattr(document.sections[0], name) == getattr(reference.sections[0], name), (path, name)
    numbering = document.sections[0]._sectPr.find('w:lnNumType', NS)
    assert numbering is not None and numbering.get('{'+NS['w']+'}restart') == 'continuous'
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith('.xml'):
                root = ET.fromstring(archive.read(name))
                if name.startswith('word/header'):
                    text = ''.join(root.itertext())
                    assert 'Instagram' not in text and 'V1.0V1.0' not in text
    pdf = fitz.open(path.with_suffix('.pdf'))
    assert len(pdf) == item['expectedPages'], (path.parent.name, 'Pages', len(pdf), item['expectedPages'])
    line_counts = []
    alternate_reader = None
    for index, page in enumerate(pdf):
        needed = min(50, len(expected) - index * 50)
        native = [word[4] for word in page.get_text('words') if word[2] < 80 and word[1] > 65 and word[4].isdigit()]
        assert list(map(int, native)) == list(range(index * 50 + 1, index * 50 + needed + 1)), (path.parent.name, index + 1, native)
        rows = body_rows(page)
        assert len(rows) == needed, (path.parent.name, index + 1, 'Wrapped rows', len(rows), needed)
        for row_index, (y, text, right) in enumerate(rows):
            expected_row = clean(expected[index * 50 + row_index])
            if clean(text) != expected_row:
                # MuPDF can collapse repeated punctuation in compressed code rows.
                if alternate_reader is None:
                    alternate_reader = PdfReader(path.with_suffix('.pdf'))
                assert expected_row in clean(alternate_reader.pages[index].extract_text()), (path.parent.name, index + 1, row_index + 1, 'PDF text mismatch', text)
            assert right <= 513, (path.parent.name, index + 1, 'Code outside margin', right)
            assert y < 780
        line_counts.append(needed)
    if render:
        key = 'web' if '多模态' in item['name'] else 'hm' if '智能巡检' in item['name'] else 'wx'
        destination = WORK / 'review' / key
        destination.mkdir(parents=True, exist_ok=True)
        images = []
        for index, page in enumerate(pdf):
            output = destination / f'page-{index+1:02d}.png'
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(output)
            images.append(output)
        for start in range(0, len(images), 12):
            contact(images[start:start+12], destination / f'contact-{start//12+1}.png')
    return dict(software=item['name'], sourceLines=len(expected), codePages=len(pdf), lineCountsPerPage=line_counts,
                inlineNumbersRemoved=True, automaticMarginLineNumbers=True, sourceTextUnchanged=True,
                templateGeometryMatches=True, noWrappedSourceRows=True, pdfTextMatches=True,
                sourceDocxSha256=sha(path), sourcePdfSha256=sha(path.with_suffix('.pdf')))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--published', action='store_true')
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    items = json.loads((WORK / '待检查清单.json').read_text('utf-8'))
    reports = []
    for item in items:
        report = verify(item, args.published, args.render)
        reports.append(report)
        print(item['name'], 'verified', report['codePages'], 'pages;', report['sourceLines'], 'rows')
    (WORK / ('最终检查.json' if args.published else '排版检查.json')).write_text(json.dumps(reports, ensure_ascii=False, indent=2), 'utf-8')


if __name__ == '__main__':
    main()
