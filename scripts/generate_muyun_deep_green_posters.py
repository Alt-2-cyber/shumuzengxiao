#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate deep-green image2 poster slides and package them as a PPTX.

The workflow intentionally separates visual generation from text rendering:
image2 creates a no-text background for each page, then PIL overlays the exact
Chinese copy from the source file. This avoids model hallucinated Chinese text
while keeping the slides poster-like.
"""

from __future__ import annotations

import argparse
import base64
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
import io
import json
import os
from pathlib import Path
import random
import re
import sys
import time
from typing import Iterable
from urllib.request import Request, urlopen

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches
import requests


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "ppt_pages_source.txt"
DEFAULT_OUT_DIR = ROOT / "outputs" / "muyun_deep_green_posters"

MODEL = os.getenv("GPT_IMAGE_MODEL_NAME", "gpt-image-2")
BASE_URL = os.getenv("OPENAI_BASE_URL", "http://ctoken.top/v1").rstrip("/")
API_URL = os.getenv("GPT_IMAGE_API_URL", f"{BASE_URL}/images/generations")
IMAGE_SIZE = os.getenv("IMAGE_SIZE", "1536x864")
IMAGE_QUALITY = os.getenv("IMAGE_QUALITY", "medium")

W, H = 1536, 864
SLIDE_W, SLIDE_H = 13.333333, 7.5

COLORS = {
    "ink": (4, 20, 13),
    "bg": (5, 38, 25),
    "deep": (8, 53, 36),
    "panel": (8, 65, 45),
    "panel2": (12, 82, 56),
    "emerald": (29, 185, 115),
    "mint": (129, 230, 188),
    "lime": (172, 220, 84),
    "gold": (241, 196, 84),
    "white": (244, 252, 246),
    "soft": (199, 223, 207),
    "muted": (150, 184, 160),
    "line": (45, 116, 78),
}


@dataclass(frozen=True)
class SlidePage:
    number: int
    title: str
    lines: list[str]

    @property
    def page_id(self) -> str:
        return f"{self.number:02d}"


def clean_line(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\ufeff", "").strip())


def parse_source(path: Path) -> list[SlidePage]:
    text = path.read_text(encoding="utf-8-sig")
    slides: list[SlidePage] = []
    current_number: int | None = None
    current_title = ""
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_number, current_title, current_lines
        if current_number is not None:
            slides.append(SlidePage(current_number, current_title, current_lines))
        current_number = None
        current_title = ""
        current_lines = []

    for raw in text.splitlines():
        line = clean_line(raw)
        if not line or line == "PPT逐页内容":
            continue
        match = re.match(r"^P(\d{2})\s+(.+)$", line)
        if match:
            flush()
            current_number = int(match.group(1))
            current_title = clean_line(match.group(2))
            continue
        if current_number is not None:
            current_lines.append(line)

    flush()
    if not slides:
        raise ValueError(f"No slides parsed from {path}")
    expected = list(range(1, len(slides) + 1))
    actual = [slide.number for slide in slides]
    if actual != expected:
        raise ValueError(f"Slide numbering is not contiguous: {actual[:8]} ... {actual[-8:]}")
    return slides


def select_slides(all_slides: list[SlidePage], spec: str) -> list[SlidePage]:
    if spec.lower() in {"all", "*"}:
        return all_slides
    wanted: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = part.split("-", 1)
            wanted.update(range(int(start), int(end) + 1))
        else:
            wanted.add(int(part))
    found = [slide for slide in all_slides if slide.number in wanted]
    missing = sorted(wanted.difference({slide.number for slide in found}))
    if missing:
        raise ValueError(f"Slides not found: {missing}")
    return found


def section_for(slide: SlidePage) -> str:
    n = slide.number
    if n == 1:
        return "封面"
    if 2 <= n <= 10:
        return "项目背景"
    if 11 <= n <= 19:
        return "项目创新"
    if 20 <= n <= 32:
        return "产业价值"
    if 33 <= n <= 36:
        return "团队协作"
    if 37 <= n <= 41:
        return "个人成长"
    if n == 42:
        return "发展规划"
    if n == 43:
        return "封底"
    return slide.title


def visual_focus(slide: SlidePage) -> str:
    joined = " ".join([slide.title, *slide.lines])
    n = slide.number
    if n == 1:
        return "modern smart cattle barn at night, healthy beef cattle, IoT sensors, edge camera, cloud data streams, mobile dashboard, cinematic hero composition"
    if n == 43:
        return "warm closing scene, smart cattle barn under sunrise, protective digital canopy, cattle silhouettes, AI data constellation, polished final slide"
    if "政策" in joined:
        return "policy support, digital agriculture roadmap, abstract document sheets, national agriculture data map silhouette, green data network"
    if "市场" in joined or "财务" in joined or "融资" in joined:
        return "market analysis dashboard, growth charts, customer segmentation, investment planning, SaaS business metrics, green technology finance"
    if "产业链" in joined:
        return "upstream middle downstream industry chain, sensor modules, assembly, cloud service, farm users, distribution network"
    if "商业模式" in joined or "营销" in joined:
        return "business model canvas, channel marketing, exhibition booth, online promotion, hardware and SaaS service blocks"
    if "研发历程" in joined or "202" in joined or "2030" in joined:
        return "horizontal product roadmap, milestones, prototype hardware, trial farm deployment, future expansion timeline"
    if "竞品" in joined or "比较" in joined:
        return "comparison matrix between manual inspection, camera, single sensor, large smart farm system, and lightweight smart barn platform"
    if "核心技术" in joined or "传感器" in joined or "视觉" in joined or "声音" in joined:
        return "multimodal sensing system, ESP32-S3 nodes, MaixCam edge vision, microphone waveform, environmental sensors, AI risk fusion"
    if "解决方案" in joined or "平台" in joined or "闭环" in joined:
        return "end-edge-cloud-application closed-loop architecture, alert confirmation, command dispatch, review records, data flow arrows"
    if "知识产权" in joined or "创新性" in joined or "认可" in joined:
        return "innovation validation wall, certificates placeholders without text, software modules, patent concept, expert review cards"
    if "客户" in joined or "试用" in joined or "应用领域" in joined:
        return "customer scenario collage, small beef cattle farms, vocational college training base, agriculture demonstration site, farm service agency"
    if "团队" in joined or "教师" in joined or "专家" in joined or "创始人" in joined:
        return "multidisciplinary team collaboration, hardware engineer, AI developer, livestock expert, business planner, classroom and farm fieldwork"
    if "院校" in joined or "教育" in joined or "就业" in joined or "产教" in joined:
        return "education practice scene, students testing smart barn devices, classroom training, industry-education integration, career pathway map"
    if "痛点" in joined or "初心" in joined or "调研" in joined:
        return "field investigation in cattle barn, manual patrol pain points, night inspection blind spots, risk alerts, practical agriculture technology"
    return "smart agriculture pitch deck background, cattle barn telemetry, IoT devices, AI analytics dashboard, clean data visualization"


def background_prompt(slide: SlidePage) -> str:
    content_summary = "；".join(slide.lines[:5])
    if len(content_summary) > 260:
        content_summary = content_summary[:260] + "..."
    return f"""
Create a premium 16:9 widescreen PPT poster BACKGROUND ONLY, no readable text.

Topic: 牧云智舍, a smart monitoring and closed-loop warning platform for small and medium beef cattle barns.
Slide {slide.page_id}: {slide.title}
Visual focus: {visual_focus(slide)}
Content context for visual planning only: {content_summary}

Style requirements:
- Deep green theme: forest green, dark emerald, moss green, black-green base.
- Accents: cyan-green data lines, fresh grass green highlights, a small amount of warm gold for energy.
- Chinese college innovation competition / business pitch deck aesthetic.
- Realistic agriculture + clean technology infographic hybrid.
- Leave large clean dark areas for later Chinese text overlay, especially top-left and center-left.
- High-end, polished, cinematic, not cartoonish, not childish.
- Use abstract dashboards, sensor icons, cattle barn silhouettes, data streams, subtle maps, timeline or charts when appropriate.

Hard restrictions:
- NO readable text, NO letters, NO numbers, NO labels.
- NO watermark, NO QR code, NO brand logo, NO government emblem.
- Do not imitate an existing template page.
- Landscape widescreen only, width much larger than height.
""".strip()


def request_image(prompt: str, api_key: str) -> bytes:
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json; charset=utf-8",
    }
    payload = {
        "model": MODEL,
        "prompt": prompt,
        "n": 1,
        "size": IMAGE_SIZE,
        "quality": IMAGE_QUALITY,
    }
    last_error = ""
    for attempt in range(1, 5):
        try:
            response = requests.post(API_URL, headers=headers, json=payload, timeout=520)
            if response.ok:
                body = response.json()
                data = body.get("data") or []
                if data and data[0].get("b64_json"):
                    return base64.b64decode(data[0]["b64_json"])
                if data and data[0].get("url"):
                    req = Request(data[0]["url"], headers={"User-Agent": "Mozilla/5.0"})
                    with urlopen(req, timeout=180) as remote:
                        return remote.read()
                last_error = f"No image payload in response: {response.text[:500]}"
            else:
                last_error = f"HTTP {response.status_code}: {response.text[:500]}"
        except Exception as exc:  # noqa: BLE001 - want retry detail in log
            last_error = repr(exc)
        if attempt < 4:
            wait = 10 * attempt
            print(f"  retry after failure: {last_error}; waiting {wait}s", flush=True)
            time.sleep(wait)
    raise RuntimeError(last_error)


def normalize_image(data: bytes) -> Image.Image:
    image = Image.open(io.BytesIO(data)).convert("RGB")
    ratio = W / H
    iw, ih = image.size
    if abs((iw / ih) - ratio) > 0.02:
        if iw / ih > ratio:
            new_w = int(ih * ratio)
            left = (iw - new_w) // 2
            image = image.crop((left, 0, left + new_w, ih))
        else:
            new_h = int(iw / ratio)
            top = (ih - new_h) // 2
            image = image.crop((0, top, iw, top + new_h))
    return image.resize((W, H), Image.Resampling.LANCZOS)


def fallback_background(slide: SlidePage) -> Image.Image:
    rng = random.Random(slide.number * 9973)
    img = Image.new("RGB", (W, H), COLORS["bg"])
    draw = ImageDraw.Draw(img, "RGBA")
    for y in range(H):
        t = y / max(1, H - 1)
        r = int(4 + 8 * t)
        g = int(26 + 44 * t)
        b = int(17 + 24 * t)
        draw.line((0, y, W, y), fill=(r, g, b, 255))
    for _ in range(120):
        x = rng.randint(-60, W)
        y = rng.randint(-40, H)
        length = rng.randint(80, 320)
        color = rng.choice([
            (*COLORS["emerald"], 70),
            (*COLORS["mint"], 54),
            (*COLORS["lime"], 42),
            (*COLORS["gold"], 36),
        ])
        draw.line((x, y, x + length, y + rng.randint(-26, 26)), fill=color, width=rng.choice([1, 2, 3]))
        if rng.random() < 0.55:
            s = rng.randint(5, 18)
            draw.rectangle((x + length, y - s // 2, x + length + s, y + s // 2), fill=color)
    for _ in range(15):
        x = rng.randint(120, W - 120)
        y = rng.randint(120, H - 120)
        rad = rng.randint(50, 190)
        color = rng.choice([COLORS["emerald"], COLORS["mint"], COLORS["lime"]])
        draw.ellipse((x - rad, y - rad, x + rad, y + rad), outline=(*color, 32), width=2)
    draw.polygon(
        [(0, H), (W, H), (W, H - 160), (int(W * 0.78), H - 228), (int(W * 0.48), H - 175), (0, H - 260)],
        fill=(5, 35, 22, 178),
    )
    return img.filter(ImageFilter.GaussianBlur(radius=0.25))


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyhbd.ttf" if bold else r"C:\Windows\Fonts\msyhl.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    if not text:
        return (0, 0)
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if text_size(draw, candidate, fnt)[0] <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines or [""]


def split_heading(line: str) -> tuple[str, str] | None:
    for sep in ("：", ":"):
        if sep in line:
            head, body = line.split(sep, 1)
            if 1 <= len(head) <= 16:
                return head + "：", body.strip()
    return None


def estimate_height(draw: ImageDraw.ImageDraw, lines: list[str], body_font: ImageFont.FreeTypeFont, max_width: int, gap: int) -> int:
    total = 0
    for line in lines:
        wrapped = wrap_text(draw, line, body_font, max_width)
        total += len(wrapped) * (text_size(draw, "数牧", body_font)[1] + 6) + gap
    return total


def choose_text_layout(draw: ImageDraw.ImageDraw, lines: list[str], box: tuple[int, int, int, int]) -> tuple[int, int]:
    x1, y1, x2, y2 = box
    available_h = y2 - y1
    for cols in (1, 2):
        max_width = (x2 - x1 - (cols - 1) * 30) // cols - 10
        for size in range(27, 15, -1):
            fnt = font(size)
            column_groups = [[] for _ in range(cols)]
            if cols == 1:
                column_groups[0] = lines
            else:
                # Greedy balance by estimated wrapped height.
                heights = [0, 0]
                for line in lines:
                    line_h = estimate_height(draw, [line], fnt, max_width, 10)
                    idx = 0 if heights[0] <= heights[1] else 1
                    column_groups[idx].append(line)
                    heights[idx] += line_h
            max_h = max(estimate_height(draw, group, fnt, max_width, 12) for group in column_groups)
            if max_h <= available_h:
                return cols, size
    return 2, 16


def draw_rich_lines(
    draw: ImageDraw.ImageDraw,
    lines: list[str],
    box: tuple[int, int, int, int],
    size: int,
    cols: int,
) -> None:
    x1, y1, x2, _ = box
    col_gap = 34
    col_w = (x2 - x1 - col_gap * (cols - 1)) // cols
    f_body = font(size)
    f_bold = font(size, True)
    f_small = font(max(14, size - 2))
    accent_cycle = [COLORS["mint"], COLORS["lime"], COLORS["gold"], COLORS["emerald"]]

    groups = [[] for _ in range(cols)]
    if cols == 1:
        groups[0] = lines
    else:
        heights = [0, 0]
        for line in lines:
            line_h = estimate_height(draw, [line], f_body, col_w - 20, 10)
            idx = 0 if heights[0] <= heights[1] else 1
            groups[idx].append(line)
            heights[idx] += line_h

    for col_idx, group in enumerate(groups):
        x = x1 + col_idx * (col_w + col_gap)
        y = y1
        for line_idx, line in enumerate(group):
            color = accent_cycle[(line_idx + col_idx) % len(accent_cycle)]
            heading = split_heading(line)
            draw.rounded_rectangle((x, y + 6, x + 12, y + 23), radius=3, fill=(*color, 220))
            if heading:
                head, body = heading
                draw.text((x + 24, y), head, font=f_bold, fill=color)
                head_w = text_size(draw, head, f_bold)[0]
                first_x = x + 24 + head_w + 4
                max_first = col_w - (first_x - x) - 8
                wrapped = wrap_text(draw, body, f_body, max_first)
                line_h = text_size(draw, "数牧", f_body)[1] + 6
                if wrapped:
                    draw.text((first_x, y), wrapped[0], font=f_body, fill=COLORS["white"])
                    yy = y + line_h
                    for extra in wrapped[1:]:
                        draw.text((x + 24, yy), extra, font=f_body, fill=COLORS["white"])
                        yy += line_h
                    y = yy + 11
                else:
                    y += line_h + 11
            else:
                wrapped = wrap_text(draw, line, f_body, col_w - 32)
                yy = y
                for wrapped_line in wrapped:
                    draw.text((x + 24, yy), wrapped_line, font=f_body if len(line) > 14 else f_small, fill=COLORS["white"])
                    yy += text_size(draw, "数牧", f_body)[1] + 6
                y = yy + 11


def add_vignette_and_panels(img: Image.Image) -> Image.Image:
    img = ImageEnhance.Color(img).enhance(0.88)
    img = ImageEnhance.Contrast(img).enhance(0.96)
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay, "RGBA")
    draw.rectangle((0, 0, W, H), fill=(0, 22, 12, 76))
    for i in range(0, 260):
        alpha = int(120 * (1 - i / 260))
        draw.rectangle((i, i, W - i, H - i), outline=(0, 0, 0, alpha))
    return Image.alpha_composite(img.convert("RGBA"), overlay)


def draw_common_header(draw: ImageDraw.ImageDraw, slide: SlidePage) -> None:
    section = section_for(slide)
    f_tag = font(20, True)
    f_num = font(28, True)
    draw.rounded_rectangle((58, 34, 238, 78), radius=22, fill=(7, 61, 39, 210), outline=(*COLORS["emerald"], 150), width=2)
    draw.text((78, 43), f"P{slide.page_id}", font=f_num, fill=COLORS["gold"])
    draw.text((150, 47), section, font=f_tag, fill=COLORS["mint"])
    draw.line((260, 56, 1468, 56), fill=(*COLORS["line"], 160), width=2)
    for x in (1220, 1265, 1310, 1355):
        draw.rectangle((x, 45, x + 14, 59), fill=(*COLORS["emerald"], 190))


def draw_footer(draw: ImageDraw.ImageDraw) -> None:
    f = font(17)
    draw.text((58, 822), "牧云智舍 | 中小肉牛养殖智能监测与闭环预警平台", font=f, fill=(*COLORS["muted"], 225))
    draw.text((1290, 822), "AI 赋能 · 守护牛舍", font=f, fill=(*COLORS["muted"], 225))


def compose_cover(slide: SlidePage, bg: Image.Image) -> Image.Image:
    img = add_vignette_and_panels(bg)
    draw = ImageDraw.Draw(img, "RGBA")
    draw_common_header(draw, slide)
    draw.rounded_rectangle((70, 145, 1025, 668), radius=34, fill=(2, 32, 19, 180), outline=(*COLORS["emerald"], 145), width=2)
    draw.rectangle((70, 145, 86, 668), fill=(*COLORS["gold"], 230))
    title = slide.lines[0] if slide.lines else "牧云智舍"
    subtitle = slide.lines[1] if len(slide.lines) > 1 else "中小肉牛养殖智能监测与闭环预警平台"
    draw.text((125, 205), title, font=font(92, True), fill=COLORS["white"])
    draw.text((130, 330), subtitle, font=font(38, True), fill=COLORS["mint"])
    y = 430
    for line in slide.lines[2:]:
        draw.text((132, y), line, font=font(28), fill=COLORS["soft"])
        y += 58
    draw.rounded_rectangle((112, 606, 772, 650), radius=22, fill=(*COLORS["panel2"], 220))
    draw.text((136, 616), "连续感知  |  主动预警  |  闭环管理", font=font(23, True), fill=COLORS["gold"])
    draw_footer(draw)
    return img.convert("RGB")


def compose_closing(slide: SlidePage, bg: Image.Image) -> Image.Image:
    img = add_vignette_and_panels(bg)
    draw = ImageDraw.Draw(img, "RGBA")
    draw_common_header(draw, slide)
    draw.rounded_rectangle((190, 150, 1345, 728), radius=40, fill=(2, 32, 19, 172), outline=(*COLORS["emerald"], 150), width=2)
    y = 214
    sizes = [42, 54, 40, 32]
    for idx, line in enumerate(slide.lines[:4]):
        fnt = font(sizes[idx] if idx < len(sizes) else 28, idx in {1, 2})
        color = COLORS["gold"] if idx == 0 else COLORS["white"] if idx == 1 else COLORS["mint"]
        tw = text_size(draw, line, fnt)[0]
        draw.text(((W - tw) // 2, y), line, font=fnt, fill=color)
        y += sizes[idx] + 28
    rest = slide.lines[4:]
    if rest:
        draw.rounded_rectangle((270, 520, 1265, 655), radius=22, fill=(*COLORS["panel"], 185))
        draw_rich_lines(draw, rest, (315, 545, 1220, 635), 19, 1)
    draw_footer(draw)
    return img.convert("RGB")


def compose_general(slide: SlidePage, bg: Image.Image) -> Image.Image:
    img = add_vignette_and_panels(bg)
    draw = ImageDraw.Draw(img, "RGBA")
    draw_common_header(draw, slide)

    title = slide.title
    lines = list(slide.lines)
    headline = ""
    if lines and (len(lines[0]) <= 46 or slide.number in {2, 3, 4, 5, 7, 8, 9, 11, 15, 16, 18, 20, 21, 22, 31, 32, 38, 39}):
        headline = lines.pop(0)

    draw.text((72, 100), title, font=font(46, True), fill=COLORS["white"])
    if headline:
        wrapped_head = wrap_text(draw, headline, font(27, True), 1180)
        y = 164
        for idx, item in enumerate(wrapped_head[:2]):
            draw.text((76, y), item, font=font(27, True), fill=COLORS["mint"] if idx == 0 else COLORS["soft"])
            y += 36

    panel_top = 230 if headline else 176
    panel_bottom = 786
    draw.rounded_rectangle((58, panel_top, 1478, panel_bottom), radius=28, fill=(2, 32, 19, 188), outline=(*COLORS["emerald"], 135), width=2)
    draw.rectangle((58, panel_top, 73, panel_bottom), fill=(*COLORS["emerald"], 230))
    draw.rectangle((74, panel_top, 80, panel_bottom), fill=(*COLORS["gold"], 195))

    if not lines:
        lines = [headline] if headline else []
    text_box = (112, panel_top + 38, 1434, panel_bottom - 34)
    cols, size = choose_text_layout(draw, lines, text_box)
    draw_rich_lines(draw, lines, text_box, size, cols)

    # Right-side decorative metric rail for visual rhythm. It intentionally
    # contains only short generic labels, not source facts.
    rail_x = 1330
    for idx, (label, color) in enumerate([("端", COLORS["mint"]), ("边", COLORS["emerald"]), ("云", COLORS["lime"]), ("闭", COLORS["gold"])]):
        y = 102 + idx * 38
        draw.rounded_rectangle((rail_x, y, rail_x + 48, y + 28), radius=14, fill=(*color, 210))
        draw.text((rail_x + 24, y + 2), label, font=font(18, True), fill=COLORS["ink"], anchor="ma")

    draw_footer(draw)
    return img.convert("RGB")


def compose_poster(slide: SlidePage, bg_path: Path, poster_path: Path) -> None:
    try:
        bg = normalize_image(bg_path.read_bytes())
    except Exception:
        bg = fallback_background(slide)
    if slide.number == 1:
        poster = compose_cover(slide, bg)
    elif slide.number == 43:
        poster = compose_closing(slide, bg)
    else:
        poster = compose_general(slide, bg)
    poster.save(poster_path, quality=95)


def write_prompts(slides: Iterable[SlidePage], out_dir: Path) -> None:
    prompt_path = out_dir / "image2_prompts.jsonl"
    with prompt_path.open("w", encoding="utf-8") as fh:
        for slide in slides:
            fh.write(json.dumps({
                "slide": slide.number,
                "title": slide.title,
                "prompt": background_prompt(slide),
            }, ensure_ascii=False) + "\n")


def generate_background(slide: SlidePage, out_dir: Path, api_key: str, overwrite: bool) -> dict:
    bg_dir = out_dir / "backgrounds"
    bg_dir.mkdir(parents=True, exist_ok=True)
    bg_path = bg_dir / f"slide-{slide.page_id}-bg.png"
    if bg_path.exists() and not overwrite:
        return {"slide": slide.number, "status": "existing", "path": str(bg_path)}
    prompt = background_prompt(slide)
    started = time.time()
    try:
        image_bytes = request_image(prompt, api_key)
        image = normalize_image(image_bytes)
        image.save(bg_path, quality=95)
        return {
            "slide": slide.number,
            "status": "generated",
            "path": str(bg_path),
            "seconds": round(time.time() - started, 1),
        }
    except Exception as exc:  # noqa: BLE001
        fallback_background(slide).save(bg_path, quality=95)
        return {
            "slide": slide.number,
            "status": "fallback",
            "path": str(bg_path),
            "error": str(exc),
            "seconds": round(time.time() - started, 1),
        }


def compose_all(slides: Iterable[SlidePage], out_dir: Path, overwrite: bool) -> list[Path]:
    poster_dir = out_dir / "posters"
    bg_dir = out_dir / "backgrounds"
    poster_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for slide in slides:
        bg_path = bg_dir / f"slide-{slide.page_id}-bg.png"
        if not bg_path.exists():
            fallback_background(slide).save(bg_path, quality=95)
        poster_path = poster_dir / f"slide-{slide.page_id}.png"
        if overwrite or not poster_path.exists():
            compose_poster(slide, bg_path, poster_path)
        paths.append(poster_path)
    return paths


def build_pptx(posters: list[Path], out_dir: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    while len(prs.slides) > 0:
        # Newly-created python-pptx presentations normally have no slides, but
        # keep this for future compatibility with custom defaults.
        r_id = prs.slides._sldIdLst[0].rId  # noqa: SLF001
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]  # noqa: SLF001
    for poster in posters:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(poster), 0, 0, width=prs.slide_width, height=prs.slide_height)
    pptx_path = out_dir / "牧云智舍-深绿海报图PPT.pptx"
    prs.save(pptx_path)
    return pptx_path


def make_contact_sheet(posters: list[Path], out_dir: Path) -> Path:
    thumb_w, thumb_h = 320, 180
    cols = 5
    rows = (len(posters) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * thumb_w, rows * (thumb_h + 34)), (4, 20, 13))
    draw = ImageDraw.Draw(sheet)
    fnt = font(18, True)
    for idx, poster in enumerate(posters):
        row, col = divmod(idx, cols)
        x, y = col * thumb_w, row * (thumb_h + 34)
        with Image.open(poster) as img:
            thumb = img.convert("RGB").resize((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        sheet.paste(thumb, (x, y))
        label = f"P{idx + 1:02d}"
        draw.rectangle((x, y + thumb_h, x + thumb_w, y + thumb_h + 34), fill=(7, 61, 39))
        draw.text((x + 12, y + thumb_h + 6), label, font=fnt, fill=COLORS["gold"])
    path = out_dir / "contact_sheet.png"
    sheet.save(path, quality=92)
    return path


def verify_outputs(slides: list[SlidePage], posters: list[Path], pptx_path: Path) -> dict:
    image_sizes = []
    for poster in posters:
        with Image.open(poster) as image:
            image.load()
            image_sizes.append({"file": poster.name, "size": image.size})
    prs = Presentation(str(pptx_path))
    result = {
        "slide_count_expected": len(slides),
        "slide_count_actual": len(prs.slides),
        "poster_count": len(posters),
        "all_images_16_9": all(abs((w / h) - (16 / 9)) < 0.01 for item in image_sizes for w, h in [item["size"]]),
        "image_sizes": image_sizes,
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--slides", default="all", help="all, 1, 1-5, or comma-separated list")
    parser.add_argument("--workers", type=int, default=int(os.getenv("IMAGE_WORKERS", "1")))
    parser.add_argument("--overwrite-bg", action="store_true")
    parser.add_argument("--overwrite-posters", action="store_true")
    parser.add_argument("--skip-api", action="store_true", help="Use existing/fallback backgrounds only")
    parser.add_argument("--dry-run", action="store_true", help="Parse source and write prompts only")
    args = parser.parse_args()

    all_slides = parse_source(args.source)
    selected = select_slides(all_slides, args.slides)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    write_prompts(all_slides, args.out_dir)
    (args.out_dir / "slides_parsed.json").write_text(
        json.dumps([slide.__dict__ for slide in all_slides], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"Parsed {len(all_slides)} slides; selected {len(selected)} slides.", flush=True)
    print(f"Output: {args.out_dir}", flush=True)
    print(f"Image endpoint: {API_URL} | model={MODEL} | size={IMAGE_SIZE} | quality={IMAGE_QUALITY}", flush=True)

    if args.dry_run:
        return 0

    if not args.skip_api:
        api_key = os.getenv("OPENAI_API_KEY") or os.getenv("EASYNET_API_KEY")
        if not api_key:
            raise SystemExit("Set OPENAI_API_KEY before running, or use --skip-api.")
        statuses: list[dict] = []
        workers = max(1, args.workers)
        print(f"Generating image2 backgrounds with {workers} worker(s)...", flush=True)
        if workers == 1:
            for idx, slide in enumerate(selected, start=1):
                print(f"[{idx:02d}/{len(selected):02d}] P{slide.page_id} {slide.title}", flush=True)
                status = generate_background(slide, args.out_dir, api_key, args.overwrite_bg)
                statuses.append(status)
                print(f"  -> {status['status']} {status.get('seconds', 0)}s", flush=True)
        else:
            with ThreadPoolExecutor(max_workers=workers) as executor:
                futures = {
                    executor.submit(generate_background, slide, args.out_dir, api_key, args.overwrite_bg): slide
                    for slide in selected
                }
                done = 0
                for future in as_completed(futures):
                    slide = futures[future]
                    done += 1
                    status = future.result()
                    statuses.append(status)
                    print(
                        f"[{done:02d}/{len(selected):02d}] P{slide.page_id} {slide.title} -> {status['status']} {status.get('seconds', 0)}s",
                        flush=True,
                    )
        (args.out_dir / "image2_status.json").write_text(json.dumps(statuses, ensure_ascii=False, indent=2), encoding="utf-8")

    posters = compose_all(all_slides, args.out_dir, args.overwrite_posters)
    pptx_path = build_pptx(posters, args.out_dir)
    contact_sheet = make_contact_sheet(posters, args.out_dir)
    verification = verify_outputs(all_slides, posters, pptx_path)
    verification["pptx"] = str(pptx_path)
    verification["contact_sheet"] = str(contact_sheet)
    (args.out_dir / "verification.json").write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PPTX: {pptx_path}", flush=True)
    print(f"Contact sheet: {contact_sheet}", flush=True)
    print(f"Verification: {verification['slide_count_actual']} slides, 16:9={verification['all_images_16_9']}", flush=True)
    if verification["slide_count_actual"] != verification["slide_count_expected"] or not verification["all_images_16_9"]:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
