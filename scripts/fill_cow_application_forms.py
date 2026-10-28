from __future__ import annotations

import json
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn


ROOT = Path(r'D:\table\软著申请\牛')
WORK = Path(__file__).resolve().parents[1] / '.run' / 'application-form-fill-20261024'


def feature_text(folder: Path) -> str:
    path = folder / '功能描述-填表用.txt'
    text = path.read_text('utf-8').strip()
    assert len(text) >= 501, (folder.name, len(text))
    return text


def set_paragraph_text(paragraph, text: str, size: float = 10.5, bold: bool = False) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.font.name = '宋体'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.font.bold = bold


def set_cell_text(cell, text: str, size: float = 10.5, bold: bool = False, remove_extra: bool = True) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        paragraphs = [cell.add_paragraph()]
    set_paragraph_text(paragraphs[0], text, size=size, bold=bold)
    if remove_extra:
        for paragraph in paragraphs[1:]:
            paragraph._element.getparent().remove(paragraph._element)


def fill_form(folder: Path, data: dict) -> dict:
    source = folder / '申请表模板-待本人填写.docx'
    assert source.exists(), source
    backup = WORK / '修改前备份' / folder.name / source.name
    backup.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        shutil.copy2(source, backup)

    document = Document(source)
    table = document.tables[0]

    simple_fields = {
        2: data['full_name'],
        3: data['short_name'],
        4: 'V1.0',
        5: '1 应用软件',
        11: data['development_hardware'],
        12: data['runtime_hardware'],
        13: 'Windows 10/11（开发机）',
        14: data['development_tools'],
        15: data['runtime_platform'],
        16: data['support_environment'],
        17: data['languages'],
        18: data['program_quantity'],
        19: data['purpose'],
        20: data['industry'],
        23: data['tags'],
        24: data['technical_detail'],
    }
    for row_index, value in simple_fields.items():
        set_cell_text(table.cell(row_index, 1), value, size=10 if row_index in (11, 12, 14, 15, 16, 17, 18, 19, 20, 24) else 10.5)

    # The template deliberately leaves this row open for the required 500+ word feature description.
    set_cell_text(table.cell(21, 1), feature_text(folder), size=9.5)
    # Keep the visual tag palette in row 22 unchanged; row 23 records the selected tag serials.

    # Project fields only: identity, dates, publication, company and contact rows remain untouched/blank.
    document.core_properties.title = data['full_name'] + ' V1.0 申请表'
    document.core_properties.subject = '软件著作权登记项目资料'
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.save(source)

    reopened = Document(source)
    check = reopened.tables[0]
    for row_index, expected in simple_fields.items():
        assert expected in check.cell(row_index, 1).text, (folder.name, row_index, check.cell(row_index, 1).text)
    assert feature_text(folder) in check.cell(21, 1).text
    for row_index in range(6, 11):
        assert check.cell(row_index, 1).text.strip() == '', (folder.name, row_index, check.cell(row_index, 1).text)
    for row_index in range(26, 35):
        assert check.cell(row_index, 1).text.strip() == '', (folder.name, row_index, check.cell(row_index, 1).text)
    return {
        'software': data['full_name'],
        'path': str(source),
        'featureCharacters': len(feature_text(folder)),
        'projectRowsFilled': sorted(list(simple_fields.keys()) + [21]),
        'identityAndContactRowsBlank': list(range(6, 11)) + list(range(26, 35)),
        'backup': str(backup),
    }


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    projects = [
        {
            'folder': '数牧增效牛舍多模态监测管理系统',
            'full_name': '数牧增效牛舍多模态监测管理系统',
            'short_name': '数牧增效Web',
            'development_hardware': 'Windows或Linux x86-64开发计算机，满足JDK、Node.js和数据库运行要求，具备稳定网络接口。',
            'runtime_hardware': '服务器或普通PC（x86-64），具备稳定网络连接；浏览器终端用于访问平台。',
            'development_tools': 'JDK 17、Maven、Node.js、Vite、代码编辑器。',
            'runtime_platform': 'Windows或Linux服务端；Chrome或Edge浏览器。',
            'support_environment': 'Spring Boot 3.3.5、Vue 3、MySQL或H2、JPA持久化、浏览器、HTTP/WebSocket服务。',
            'languages': 'Java 17、JavaScript ES模块、HTML、CSS、SQL。',
            'program_quantity': '非空源代码约14558行；本次交存3000行，共60页，每页50行。',
            'purpose': '面向肉牛养殖场建立统一的环境、行为、声音、视觉、健康和设备管理入口，帮助巡检人员集中查看多源数据、追踪异常、核对证据并留存处置记录。',
            'industry': '智慧畜牧、肉牛养殖、农业物联网、动物健康辅助管理。',
            'tags': '9、10、14（大数据软件、人工智能软件、物联网软件）。',
            'technical_detail': '采用前后端分离架构，通过HTTP/WebSocket接入多源数据，结合持久化、角色权限、证据质量校验、风险筛查和设备回执审计。',
        },
        {
            'folder': '数牧增效牛舍智能巡检与预警平台',
            'full_name': '数牧增效牛舍智能巡检与预警平台',
            'short_name': '数牧增效APP',
            'development_hardware': 'Windows x86-64开发计算机，配置DevEco Studio和HarmonyOS SDK，并使用HarmonyOS手机或匹配SDK的模拟器测试。',
            'runtime_hardware': 'HarmonyOS手机或匹配SDK的手机模拟器。',
            'development_tools': 'DevEco Studio、HarmonyOS SDK 6.0.2（API 22）。',
            'runtime_platform': 'HarmonyOS手机或匹配SDK的手机模拟器。',
            'support_environment': 'ArkUI、AbilityKit、ArkData Preferences、MediaLibraryKit、平台HTTP服务。',
            'languages': 'ArkTS；JSON/JSON5配置。',
            'program_quantity': '非空源代码2390行；全部交存，共48页，每页50行，末页40行。',
            'purpose': '为养殖人员提供鸿蒙移动端巡检入口，用于查看巡检概览、选择牛只照片、保存本地记录、查看设备演示状态，并连接平台使用AI问答、专家咨询和养牛资讯服务。',
            'industry': '智慧畜牧、肉牛养殖移动巡检、农业物联网辅助管理。',
            'tags': '1、10、14（APP、人工智能软件、物联网软件）。',
            'technical_detail': '采用ArkUI声明式界面与底部图标导航，结合系统选图器、本机Preferences持久化、平台HTTP接口、AI问答、咨询回执和资讯服务。',
        },
        {
            'folder': '数牧增效牛舍移动巡检与设备管理小程序平台',
            'full_name': '数牧增效牛舍移动巡检与设备管理小程序平台',
            'short_name': '数牧增效小程序',
            'development_hardware': 'Windows开发计算机，安装微信开发者工具和TypeScript编译插件，可连接微信测试设备。',
            'runtime_hardware': '支持微信客户端的Android或iOS手机，以及微信开发者工具模拟器。',
            'development_tools': '微信开发者工具、TypeScript编译插件。',
            'runtime_platform': '微信小程序运行环境；微信开发者工具。',
            'support_environment': '微信基础库3.16.0、原生Page/Component、自定义tabBar、StorageSync和平台HTTP接口。',
            'languages': 'TypeScript 5.8、WXML、WXSS；JSON配置。',
            'program_quantity': '非空源代码2633行；全部交存，共53页，每页50行，末页33行。',
            'purpose': '为养殖人员提供轻量化微信移动巡检工具，用于查看牛舍信息、选择照片进行本地演示分析、保存识别记录、浏览设备管理入口，并使用AI助手、专家咨询和养牛资讯功能。',
            'industry': '智慧畜牧、肉牛养殖移动服务、农业物联网辅助管理。',
            'tags': '10、13、14（人工智能软件、小程序、物联网软件）。',
            'technical_detail': '采用微信原生Page/Component结构，通过chooseMedia选图、StorageSync缓存、自定义tabBar和平台接口实现巡检、问答、咨询及资讯管理。',
        },
    ]
    reports = []
    for project in projects:
        folder = ROOT / project['folder']
        reports.append(fill_form(folder, project))
    (WORK / '填写检查.json').write_text(json.dumps(reports, ensure_ascii=False, indent=2), 'utf-8')
    print(json.dumps(reports, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
