from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import load_workbook


BASE = Path("D:/table")
PURCHASE_XLSX = BASE / "数牧增效-硬件采购清单1.xlsx"
TEMPLATE_DOCX = BASE / "物料申请表(3)(2).docx"
OUT_DOCX = BASE / "数牧增效-物料申请表1.docx"


COMPETITIONS = [
    "天津市大学生创新训练计划（大创）",
    "中国国际大学生创新大赛（互联网+）",
    "“挑战杯”大学生课外学术科技作品竞赛",
    "仁爱杯校级创新创业大赛",
    "开放原子开源大赛",
    "卓越杯校级科技竞赛",
    "中国高校计算机大赛--嵌入式设计赛",
    "全国大学生电子设计竞赛",
    "全国大学生物联网设计竞赛",
    "华北五省（市、自治区）大学生计算机应用大赛天津赛区竞赛",
]


def clean_text(value):
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def extract_use(link_text):
    text = clean_text(link_text)
    marker = "用途："
    if marker in text:
        return text.split(marker, 1)[1].strip()
    return ""


def normalize_item_use(name, spec, qty, link_text):
    use = extract_use(link_text)
    if use:
        return use

    fallback = {
        "MAX30102": "用于项目扩展健康体征采集验证，可作为单牛体征或可穿戴节点的预研模块，辅助展示系统后续可扩展到个体健康监测。",
        "OLED": "用于本地显示环境数据、网络状态、告警状态和调试信息，便于现场演示与故障排查。",
        "继电器": "用于控制风扇、雾化、补光、声光报警等外部执行器，实现环境异常时的本地联动。",
        "MOSFET": "用于直流风扇、雾化器等低压负载开关控制或 PWM 调速，避免主控 GPIO 直接带载。",
        "风扇": "作为牛舍通风执行器，在高温、高湿或氨气偏高时由主控自动联动。",
        "雾化": "作为降温或加湿执行器，用于演示高温干燥场景下的环境调节闭环。",
        "报警": "用于火情、氨气超阈值、视觉异常或声音异常等场景下的本地声光提示。",
        "电源": "用于为控制箱、执行器和边缘设备提供稳定低压供电，保证系统连续运行。",
        "降压": "用于将 12V 电源转换为外设所需电压，给 MaixCam、继电器或其他模块提供稳定供电。",
        "控制箱": "用于安装主控、电源、继电器、端子排和保护器件，提升牛舍粉尘、水汽环境下的可靠性。",
        "支架": "用于固定 MaixCam 或摄像头模块，便于调整牛舍高位俯视/侧视角度。",
        "补光": "用于提升夜间或弱光环境下视觉识别稳定性，减少因光照不足造成的漏检。",
        "端子": "用于控制箱内部布线、传感器连接、执行器接线和后期维护。",
        "保险": "用于电源输入保护、感性负载续流和浪涌抑制，降低主控复位或损坏风险。",
    }

    merged = f"{name} {spec}"
    for key, desc in fallback.items():
        if key in merged:
            return desc
    return f"用于“数牧增效”项目硬件原型搭建，采购数量为 {qty}，规格按表格中“颜色分类”执行。"


def load_purchase_items():
    workbook = load_workbook(PURCHASE_XLSX, data_only=True)
    worksheet = workbook.active
    items = []

    for row_index in range(2, worksheet.max_row + 1):
        name = clean_text(worksheet.cell(row_index, 1).value)
        link = clean_text(worksheet.cell(row_index, 2).value)
        spec = clean_text(worksheet.cell(row_index, 3).value)
        qty = worksheet.cell(row_index, 4).value

        if not name or name == "备注":
            continue
        if qty in (None, ""):
            continue

        try:
            qty_value = int(qty)
        except (TypeError, ValueError):
            qty_value = clean_text(qty)

        items.append(
            {
                "name": name,
                "link": link,
                "spec": spec,
                "qty": qty_value,
                "use": normalize_item_use(name, spec, qty_value, link),
            }
        )

    return items


def set_run_font(run, size=None, bold=None):
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_heading(doc, text, size=14, center=False):
    paragraph = doc.add_paragraph()
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)
    return paragraph


def add_body_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def apply_table_font(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=10.5)


def build_doc(items):
    # Start from a fresh document to avoid carrying stale project text from the old template.
    doc = Document()
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    add_heading(doc, "物料采购申请报告", 18, True)
    add_heading(doc, "采购硬件及使用方式", 14)

    add_body_paragraph(
        doc,
        "本次物料采购依据《数牧增效-硬件采购清单1.xlsx》整理，面向“数牧增效”肉牛养殖智能监测项目硬件原型搭建。采购内容覆盖 ESP32-S3 主控、MaixCam 边缘 AI 摄像头、环境与健康传感器、声音采集模块、执行器联动模块、电源与现场安装辅材等部分，用于完成牛舍环境监测、视觉识别、声音异常采集、本地告警和云端数据展示闭环。",
    )
    add_body_paragraph(
        doc,
        "ESP32-S3 开发板作为设备端核心主控，负责采集温湿度、氨气、光照、火情/烟雾等环境数据，并通过 WiFi/MQTT 上传至后端平台；同时通过 GPIO、继电器或 MOSFET 驱动风扇、雾化器、声光报警器等执行器，实现本地自动控制。表格中该开发板采购 3 块，主要用于主机调试、备用替换和后续扩展节点验证。",
    )
    add_body_paragraph(
        doc,
        "Sipeed MaixCam 边缘 AI 摄像头开发板用于视觉健康监测节点，在牛舍端执行目标检测、姿态识别、卧倒/低活动等异常事件判断，并保存截图证据。采购配置选择带 32G 镜像卡版本，可减少系统烧录、模型部署和本地缓存配置工作量。",
    )
    add_body_paragraph(
        doc,
        "SHT31 温湿度传感器、MQ-137 氨气传感器、BH1750 光照传感器、MQ-135 空气质量/烟雾传感器和火焰传感器共同构成环境感知层，用于采集牛舍温湿度、氨气趋势、光照强度和安全风险信号。MAX30102 心率血氧传感器作为扩展体征采集模块，可用于后续单牛健康监测或可穿戴节点预研展示。",
    )
    add_body_paragraph(
        doc,
        "执行器、电源、防护和布线类物料用于保障样机真正可运行。继电器、MOSFET、风扇、雾化器、声光报警器等模块用于搭建环境异常后的本地联动控制；电源、控制箱、端子排、线材、支架、补光灯和保护元件用于完成控制箱装配、现场固定、弱光补偿与电气保护，避免样机只停留在软件演示层面。",
    )

    add_heading(doc, "核心物料简表", 14)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["序号", "物料名称", "规格/颜色分类", "数量", "主要用途"]
    for col, header in enumerate(headers):
        table.rows[0].cells[col].text = header

    for index, item in enumerate(items, 1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = item["name"]
        cells[2].text = item["spec"]
        cells[3].text = str(item["qty"])
        cells[4].text = item["use"]

    add_heading(doc, "预计参与比赛", 14)
    competition_table = doc.add_table(rows=len(COMPETITIONS), cols=1)
    competition_table.style = "Table Grid"
    for row_index, competition in enumerate(COMPETITIONS):
        competition_table.rows[row_index].cells[0].text = competition

    for table in doc.tables:
        apply_table_font(table)

    doc.save(OUT_DOCX)


def main():
    if not PURCHASE_XLSX.exists():
        raise FileNotFoundError(PURCHASE_XLSX)
    if not TEMPLATE_DOCX.exists():
        raise FileNotFoundError(TEMPLATE_DOCX)

    items = load_purchase_items()
    build_doc(items)
    print(f"generated={OUT_DOCX}")
    print(f"items={len(items)}")


if __name__ == "__main__":
    main()
