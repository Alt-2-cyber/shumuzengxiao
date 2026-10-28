from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side


BASE = Path("D:/table")
XLSX_TEMPLATE = next(p for p in BASE.glob("*.xlsx") if "(5)(2)" in p.name)
OUT_DOC = BASE / "数牧增效-物料采购申请报告.docx"
OUT_XLSX = BASE / "数牧增效-硬件采购清单.xlsx"


ITEMS = [
    {
        "name": "ESP32-S3 DevKitC-1 / N16R8 开发板，带 USB 下载调试接口",
        "keyword": "ESP32-S3 DevKitC N16R8 开发板 USB Type-C",
        "spec": "ESP32-S3-WROOM-1-N16R8，开发板款，带排针",
        "qty": 2,
        "use": "作为牛舍环境控制节点主控，负责温湿度、氨气、光照、火情/烟雾数据采集，执行风扇、雾化、声光报警联动，并通过 WiFi/MQTT 上传数据。采购 2 块便于调试和备份。",
    },
    {
        "name": "Sipeed MaixCam 边缘 AI 摄像头开发板",
        "keyword": "Sipeed MaixCam 边缘AI摄像头 开发板",
        "spec": "MaixCam 标准套装，含摄像头/外壳优先",
        "qty": 1,
        "use": "作为视觉 AI 节点，在牛舍端运行轻量 YOLO 模型，实现牛只站立、卧倒、活动异常和事件截图留证，不依赖云端视频流长期上传。",
    },
    {
        "name": "microSD 存储卡",
        "keyword": "microSD 32GB 高速 TF 卡 工业级",
        "spec": "32GB 及以上，Class10/U1，优先耐高温款",
        "qty": 1,
        "use": "用于 MaixCam 本地保存模型文件、事件截图、短视频和运行日志，网络异常时可缓存关键证据。",
    },
    {
        "name": "SHT31 / AHT20 温湿度传感器模块",
        "keyword": "SHT31 AHT20 温湿度传感器 I2C 模块",
        "spec": "I2C 数字输出，优先 SHT31；预算紧张可选 AHT20",
        "qty": 2,
        "use": "采集牛舍温湿度，作为高温高湿告警和风扇/雾化联动依据；采购 2 个便于备份和对比标定。",
    },
    {
        "name": "MQ-137 氨气传感器模块",
        "keyword": "MQ-137 氨气传感器 模块 NH3",
        "spec": "MQ-137 模拟量输出模块，带调理板",
        "qty": 2,
        "use": "用于原型阶段氨气趋势监测和阈值告警，提示粪污区通风风险；该模块用于趋势判断，后续正式版预留工业 NH3 传感器接口。",
    },
    {
        "name": "BH1750 光照传感器模块",
        "keyword": "BH1750 光照传感器 I2C 模块",
        "spec": "GY-302/BH1750FVI，I2C 数字光照强度",
        "qty": 2,
        "use": "采集牛舍光照强度，辅助判断夜间补光、摄像头曝光和视觉识别可靠性。",
    },
    {
        "name": "烟雾传感器模块",
        "keyword": "MQ-2 烟雾传感器 模块 火灾报警",
        "spec": "MQ-2/MQ-135 均可，模拟量/数字量输出",
        "qty": 2,
        "use": "与火焰/温升信号组合用于牛舍安全告警，降低单一火焰传感器误报风险。",
    },
    {
        "name": "火焰传感器模块",
        "keyword": "火焰传感器 模块 红外 火情检测",
        "spec": "数字量输出，带电位器阈值调节",
        "qty": 2,
        "use": "用于火情快速检测，与烟雾数据共同构成本地声光报警触发条件。",
    },
    {
        "name": "INMP441 I2S 数字麦克风模块",
        "keyword": "INMP441 I2S 数字麦克风 模块",
        "spec": "I2S 输出，3.3V 供电",
        "qty": 2,
        "use": "采集牛舍咳嗽、喘鸣、异常叫声等声音信号，供 ESP32-S3 做 VAD/事件截取和后端多模态风险融合。",
    },
    {
        "name": "0.96 寸 OLED 显示屏模块",
        "keyword": "0.96寸 OLED I2C 显示屏 SSD1306",
        "spec": "SSD1306，I2C，白/蓝色均可",
        "qty": 1,
        "use": "用于本地显示温湿度、氨气趋势、网络状态和告警状态，方便比赛展示和现场调试。",
    },
    {
        "name": "四路继电器模块，带光耦隔离",
        "keyword": "5V 四路继电器模块 光耦隔离",
        "spec": "5V 4 路，光耦隔离，低电平/高电平触发按开发板适配",
        "qty": 1,
        "use": "控制风扇、雾化器、声光报警、补光灯等低压执行器，强弱电隔离便于安全演示。",
    },
    {
        "name": "MOSFET 大功率开关驱动模块",
        "keyword": "MOSFET 开关驱动模块 PWM 大功率 直流电机",
        "spec": "支持 5V/12V 直流负载，带接线端子",
        "qty": 2,
        "use": "用于风扇 PWM 调速、雾化片或低压直流负载开关控制，避免直接由 ESP32-S3 GPIO 带载。",
    },
    {
        "name": "12V 直流散热风扇",
        "keyword": "12V 直流风扇 12025 散热风扇",
        "spec": "12V，8cm/12cm 均可，带防护网优先",
        "qty": 1,
        "use": "模拟牛舍通风执行器，在高温、高湿或氨气偏高时由主控自动联动。",
    },
    {
        "name": "超声波雾化器模块 / 12V 雾化片套件",
        "keyword": "12V 超声波雾化器模块 雾化片 套件",
        "spec": "12V 雾化片套件，带驱动板",
        "qty": 1,
        "use": "模拟牛舍降温/加湿执行器，在高温干燥场景下与风扇形成联动控制演示。",
    },
    {
        "name": "声光报警器 / 蜂鸣器警示灯模块",
        "keyword": "12V 声光报警器 蜂鸣器 警示灯 模块",
        "spec": "12V 小型声光报警器或蜂鸣器+LED 警示灯组合",
        "qty": 1,
        "use": "本地即时告警输出，用于火情、氨气超阈值、视觉异常或声音异常事件提示。",
    },
    {
        "name": "12V 防水开关电源",
        "keyword": "12V 5A 防水开关电源 监控电源",
        "spec": "AC220V 转 DC12V，5A 左右，带过流保护",
        "qty": 1,
        "use": "为控制箱、风扇、雾化、继电器、补光等低压设备提供主电源，容量预留调试余量。",
    },
    {
        "name": "DC-DC 降压模块 12V 转 5V",
        "keyword": "DC-DC 降压模块 12V转5V 3A USB",
        "spec": "5V/3A，优先带 USB 或接线端子",
        "qty": 2,
        "use": "给 MaixCam、继电器模块和部分 5V 外设供电；MaixCam 建议独立一路 5V/2A 以上。",
    },
    {
        "name": "DC-DC 降压模块 12V 转 3.3V",
        "keyword": "DC-DC 降压模块 12V转3.3V 1A",
        "spec": "3.3V/1A，低纹波，带接线端子优先",
        "qty": 1,
        "use": "为 I2C 传感器、I2S 麦克风和逻辑电平外设提供稳定 3.3V 电源。",
    },
    {
        "name": "TTL 转 RS485 模块",
        "keyword": "MAX485 TTL转RS485 模块 3.3V 5V",
        "spec": "MAX485/SP3485 模块，3.3V 兼容优先",
        "qty": 2,
        "use": "为后续接入工业级 NH3/温湿度 RS485 传感器预留接口，也可用于现场长线通信实验。",
    },
    {
        "name": "4G 无线路由器 / CPE",
        "keyword": "4G 无线路由器 CPE 插卡 工业 物联网",
        "spec": "支持 SIM 卡，LAN/WiFi 输出；如现场 WiFi 稳定可暂缓",
        "qty": 1,
        "use": "用于牛舍 WiFi 覆盖不足时提供 MQTT/HTTP 网络回传通道，保障现场试点数据上传。",
    },
    {
        "name": "IP65 防水控制箱",
        "keyword": "IP65 防水控制箱 塑料 配电箱 透明盖",
        "spec": "约 200×300mm 或按实物尺寸选型，带安装底板优先",
        "qty": 1,
        "use": "安装 ESP32-S3、电源、继电器、端子排和保护器件，提升牛舍粉尘、水汽环境下的可靠性。",
    },
    {
        "name": "摄像头万向支架 / 监控支架",
        "keyword": "摄像头万向支架 监控支架 壁装",
        "spec": "可调角度，壁装/顶装均可",
        "qty": 1,
        "use": "用于固定 MaixCam，便于实现 2.5m 到 3.5m 高位俯视或侧视安装角度。",
    },
    {
        "name": "12V LED 补光灯 / 红外补光灯",
        "keyword": "12V LED 补光灯 红外补光灯 监控 夜视",
        "spec": "白光或红外按镜头适配选择，小功率演示款即可",
        "qty": 1,
        "use": "提升夜间或弱光环境下 MaixCam 视觉识别稳定性，避免因光照不足造成漏检。",
    },
    {
        "name": "接线端子排、杜邦线、硅胶线、热缩管组合",
        "keyword": "接线端子排 杜邦线 硅胶线 热缩管 套装",
        "spec": "端子排 + 公母杜邦线 + 22AWG/24AWG 线材 + 热缩管",
        "qty": 1,
        "use": "用于控制箱内部布线、传感器连接、执行器接线和调试维护，是样机装配的基础耗材。",
    },
    {
        "name": "保险丝座、TVS/续流二极管保护件套装",
        "keyword": "保险丝座 TVS 二极管 续流二极管 电子元件套装",
        "spec": "含保险丝座、保险丝、1N5819/1N4007、TVS 管等",
        "qty": 1,
        "use": "用于电源输入保护、继电器/风扇等感性负载续流和浪涌抑制，降低 MCU 复位和损坏风险。",
    },
]

COMPETITIONS = [
    "天津市大学生创新训练计划（大创）",
    "中国国际大学生创新大赛（互联网+）",
    "“挑战杯”大学生课外学术科技作品竞赛",
    "仁爱杯校级创新创业大赛",
    "开放原子开源大赛",
    "卓越杯校级科技竞赛",
    "中国高校计算机大赛--嵌入式设计赛",
    "全国大学生电子设计竞赛",
    "全国大学生物联网设计竞赛",
    "华北五省（市、自治区）大学生计算机应用大赛天津赛区竞赛",
]


def set_run_font(run, size=None, bold=None):
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_heading(doc, text, size=14, center=False):
    paragraph = doc.add_paragraph()
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)
    return paragraph


def build_docx():
    doc = Document()
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    add_heading(doc, "物料采购申请报告", 18, True)
    add_heading(doc, "采购硬件及使用方式", 14)

    paragraphs = [
        "本次采购面向“数牧增效”肉牛养殖智能监测项目，项目采用 ESP32-S3 主控节点、Sipeed MaixCam 边缘 AI 摄像头、环境传感器、声音采集模块和执行器联动模块，构建牛舍环境监测、视觉识别、声音异常采集、本地告警和云端看板展示的一体化原型系统。",
        "ESP32-S3 开发板是系统设备端的核心控制平台，负责采集温湿度、氨气、光照、火焰/烟雾等牛舍环境数据，并通过继电器或 MOSFET 驱动风扇、雾化器、声光报警器等执行器。同时，ESP32-S3 通过 WiFi/MQTT 与后端平台通信，并通过 UART 与 MaixCam 交换视觉事件和抓拍触发指令。",
        "Sipeed MaixCam 用于替代普通摄像头模块，作为独立视觉边缘 AI 节点部署在牛舍高位视角，完成肉牛目标检测、卧倒/低活动等行为识别、事件截图留证和本地缓存。配套采购 microSD 卡、摄像头支架和补光灯，可保证模型文件、图片日志和弱光环境下的识别稳定性。",
        "AHT20/SHT31、MQ-137、BH1750、烟雾传感器和火焰传感器用于构建牛舍环境采集层。其中 MQ-137 主要用于氨气趋势告警，后续正式试点可通过 TTL 转 RS485 模块扩展工业级 NH3 传感器；INMP441 I2S 数字麦克风用于采集咳嗽、喘鸣和异常叫声，为多模态健康风险判断提供声音证据。",
        "继电器模块、MOSFET 驱动模块、12V 风扇、雾化器和声光报警器用于搭建本地联动控制闭环。当温湿度、氨气、烟雾、火情或视觉/声音事件达到阈值时，系统可在本地完成通风、雾化、报警和抓拍联动，减少完全依赖云端响应带来的延迟。",
        "12V 防水开关电源、5V/3.3V DC-DC 降压模块、防水控制箱、端子排、线材、保险丝、TVS 和续流二极管等物料用于完成样机供电、防护和现场布线。该类基础物料虽不直接参与算法功能，但关系到牛舍粉尘、水汽和感性负载环境下的设备稳定性，是原型落地不可省略的配套部分。",
        "本清单按 1 套可演示原型系统配置，并对易损或低价传感器类物料预留 1 个备件，便于焊接调试、现场替换和数据对比标定。具体采购时可根据预算选择同规格国产模块，链接以采购搜索入口和规格关键词为准，最终价格需以下单平台实时价格为准。",
    ]

    for text in paragraphs:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5)

    add_heading(doc, "核心物料简表", 14)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for col, header in enumerate(["序号", "物料名称", "数量", "主要用途"]):
        table.rows[0].cells[col].text = header

    for index, item in enumerate(ITEMS, 1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = item["name"]
        cells[2].text = str(item["qty"])
        cells[3].text = item["use"]

    add_heading(doc, "预计参与比赛", 14)
    competition_table = doc.add_table(rows=len(COMPETITIONS), cols=1)
    competition_table.style = "Table Grid"
    for row_index, competition in enumerate(COMPETITIONS):
        competition_table.rows[row_index].cells[0].text = competition

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, size=10.5)

    doc.save(OUT_DOC)


def build_xlsx():
    workbook = load_workbook(str(XLSX_TEMPLATE))
    worksheet = workbook.active
    worksheet.title = "数牧增效采购清单"

    for row in worksheet.iter_rows(
        min_row=2, max_row=max(worksheet.max_row, len(ITEMS) + 3), min_col=1, max_col=4
    ):
        for cell in row:
            cell.value = None

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, header in enumerate(["名称", "链接", "颜色分类", "个数"], 1):
        cell = worksheet.cell(1, col)
        cell.value = header
        cell.font = Font(name="宋体", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_index, item in enumerate(ITEMS, 2):
        link = (
            f"淘宝搜索：https://s.taobao.com/search?q={quote(item['keyword'])}\n"
            f"京东搜索：https://search.jd.com/Search?keyword={quote(item['keyword'])}\n"
            f"用途：{item['use']}"
        )
        values = [item["name"], link, item["spec"], item["qty"]]
        for col, value in enumerate(values, 1):
            cell = worksheet.cell(row_index, col)
            cell.value = value
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            cell.font = Font(name="宋体", size=10)
        worksheet.row_dimensions[row_index].height = 82

    for col, width in {"A": 46, "B": 92, "C": 44, "D": 10}.items():
        worksheet.column_dimensions[col].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:D{len(ITEMS) + 1}"

    note_row = len(ITEMS) + 3
    worksheet.cell(note_row, 1).value = "备注"
    worksheet.cell(note_row, 2).value = (
        "本清单按 1 套可演示原型系统配置；链接列为采购搜索入口，最终型号、店铺和价格以下单平台实时信息为准。"
        "若预算紧张，可暂缓 4G 路由器和工业扩展接口，优先采购 ESP32-S3、MaixCam、环境传感器、电源、防护和基础线材。"
    )
    worksheet.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=4)
    for col in range(1, 5):
        cell = worksheet.cell(note_row, col)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="宋体", size=10, bold=(col == 1))
        cell.border = border
    worksheet.row_dimensions[note_row].height = 58

    workbook.save(OUT_XLSX)


def main():
    build_docx()
    build_xlsx()
    print(OUT_DOC)
    print(OUT_XLSX)
    print(f"items={len(ITEMS)}")


if __name__ == "__main__":
    main()
