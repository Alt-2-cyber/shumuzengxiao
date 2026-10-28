from pathlib import Path

from docx import Document


REFERENCE = Path("D:/table/yanglao/\u57fa\u4e8e\u7aef\u8fb9\u4e91\u670d\u534f\u540c\u7684\u9002\u8001AI\u4e3b\u52a8\u9632\u62a4\u7a7f\u6234\u7cfb\u7edf_\u8bbe\u8ba1\u6587\u6863.docx")


def pt(value):
    return None if value is None else round(value.pt, 2)


def cm(value):
    return None if value is None else round(value.cm, 2)


def main():
    doc = Document(REFERENCE)
    print("sections")
    for idx, section in enumerate(doc.sections):
        print(
            idx,
            {
                "page_w_cm": cm(section.page_width),
                "page_h_cm": cm(section.page_height),
                "top_cm": cm(section.top_margin),
                "bottom_cm": cm(section.bottom_margin),
                "left_cm": cm(section.left_margin),
                "right_cm": cm(section.right_margin),
                "header_cm": cm(section.header_distance),
                "footer_cm": cm(section.footer_distance),
            },
        )

    print("paragraph_style_profile")
    seen = set()
    for para in doc.paragraphs[:120]:
        style = para.style
        if style.name in seen:
            continue
        seen.add(style.name)
        fmt = style.paragraph_format
        font = style.font
        print(
            {
                "style": style.name,
                "font": font.name,
                "size_pt": pt(font.size),
                "bold": font.bold,
                "color": None if font.color.rgb is None else str(font.color.rgb),
                "space_before_pt": pt(fmt.space_before),
                "space_after_pt": pt(fmt.space_after),
                "line_spacing": fmt.line_spacing,
                "left_indent_cm": cm(fmt.left_indent),
                "first_line_indent_cm": cm(fmt.first_line_indent),
                "alignment": str(fmt.alignment),
            }
        )

    print("first_paragraph_shapes")
    for idx, para in enumerate(doc.paragraphs[:35]):
        text_len = len(para.text.strip())
        print(
            idx,
            {
                "style": para.style.name,
                "text_len": text_len,
                "runs": len(para.runs),
                "alignment": str(para.alignment),
                "space_before_pt": pt(para.paragraph_format.space_before),
                "space_after_pt": pt(para.paragraph_format.space_after),
            },
        )

    print("tables", len(doc.tables))
    for idx, table in enumerate(doc.tables[:10]):
        print(idx, {"rows": len(table.rows), "cols": len(table.columns), "style": table.style.name if table.style else None})

    image_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            image_count += 1
            print("image", image_count, {"target": rel.target_ref, "content_type": rel.target_part.content_type, "bytes": len(rel.target_part.blob)})
    print("image_count", image_count)


if __name__ == "__main__":
    main()
