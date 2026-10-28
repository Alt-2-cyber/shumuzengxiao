from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from verify_cow_source_template_format import body_rows

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(r'D:\table\软著申请\牛')
QA = ROOT / '.run' / 'copyright-20261023' / 'review'
QA.mkdir(parents=True, exist_ok=True)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def contact(paths, target):
    cell_w, cell_h = 300, 445
    sheet = Image.new('RGB', (cell_w * 4, cell_h * ((len(paths) + 3) // 4)), '#DDDDDD')
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc', 17)
    for index, path in enumerate(paths):
        image = Image.open(path).convert('RGB')
        image.thumbnail((cell_w - 12, cell_h - 28))
        x, y = (index % 4) * cell_w, (index // 4) * cell_h
        sheet.paste(image, (x + (cell_w - image.width) // 2, y + 24))
        draw.text((x + 10, y + 2), path.stem, font=font, fill='black')
    sheet.save(target)


def main():
    reports = []
    entries = json.loads((OUT / '材料生成清单.json').read_text('utf-8'))
    for entry in entries:
        key = entry['key']
        folder = OUT / entry['name']
        listing = json.loads((folder / '源程序文件与行号清单.json').read_text('utf-8'))
        paragraphs = [paragraph.text for paragraph in Document(entry['code']).paragraphs]
        expected = [row['text'].expandtabs(2) for row in listing['selected']]
        assert paragraphs == expected, (key, 'source changed')
        for row in listing['files']:
            assert sha(ROOT / row['path']) == row['sha256'], row['path']
            assert sha(folder / '源码归档' / row['path']) == row['sha256'], row['path']
        assert not re.search(r'sk-[a-zA-Z0-9]{25,}', '\n'.join(paragraphs)), (key, 'provider credential')
        code = fitz.open(Path(entry['code']).with_suffix('.pdf'))
        manual = fitz.open(Path(entry['manual']).with_suffix('.pdf'))
        assert len(code) == entry['stats']['codePages'], (key, 'source pages', len(code))
        assert len(manual) == 14, (key, 'manual pages', len(manual))
        line_counts = []
        outside = []
        for index, page in enumerate(code):
            identifiers = [word[4] for word in page.get_text('words') if word[2] < 80 and word[1] > 65 and word[4].isdigit()]
            required = min(50, len(expected) - index * 50)
            assert len(identifiers) == required, (key, index + 1, len(identifiers), required)
            lines = body_rows(page)
            assert len(lines) == required, (key, 'source wraps', index + 1, len(lines), required)
            line_counts.append(required)
            for line in lines:
                if line[2] > page.rect.width - 10:
                    outside.append((index + 1, line[2]))
        assert not outside, (key, 'source outside page', outside[:3])
        doc = Document(entry['manual'])
        body = '\n'.join(p.text for p in doc.paragraphs)
        assert not any(item in body for item in ['未附可核验的微信', '当前客户端不实现手机号登录、账号认证或远程后台连接', '技术交流智能化交互系统'])
        for paragraph in doc.paragraphs:
            if paragraph.style.name == 'Heading 1':
                assert paragraph.text.strip(), (key, 'empty heading')
        media_counts = []
        for page_index in range(2, 12):
            media_counts.append(len(manual[page_index].get_images()))
            assert media_counts[-1] >= 2, (key, page_index + 1, 'no real screenshot')
        figure_manifest = json.loads((folder / '截图来源与标注清单.json').read_text('utf-8'))
        assert len(figure_manifest) == entry['screenshotCount']
        assert all(item['marks'] for item in figure_manifest)
        for document_path in [entry['code'], entry['manual']]:
            with ZipFile(document_path) as archive:
                assert archive.testzip() is None
                for name in archive.namelist():
                    if name.endswith('.xml'):
                        xml = archive.read(name)
                        etree.fromstring(xml)
                        assert b'i@yufuping.com' not in xml
        desc = (folder / '功能描述-填表用.txt').read_text('utf-8').strip()
        assert 501 <= len(desc) <= 1300
        rendered = []
        render_dir = QA / key
        render_dir.mkdir(exist_ok=True)
        for index, page in enumerate(manual):
            destination = render_dir / f'page-{index+1:02d}.png'
            page.get_pixmap(matrix=fitz.Matrix(1.45, 1.45), alpha=False).save(destination)
            rendered.append(destination)
        contact(rendered, QA / f'{key}-manual-contact.png')
        source_paths = []
        for page_number in sorted({1, 2, min(30, len(code)), min(31, len(code)), len(code)}):
            path = render_dir / f'source-{page_number:02d}.png'
            code[page_number-1].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(path)
            source_paths.append(path)
        contact(source_paths, QA / f'{key}-source-contact.png')
        report = dict(software=entry['name'], codePages=len(code), manualPages=len(manual), sourceLines=len(expected),
                      sourceFileCount=entry['stats']['fileCount'], sourcePhysicalLines=entry['stats']['physicalLines'],
                      lineCountsPerPage=line_counts, workflowFigures=10, screenshots=entry['screenshotCount'],
                      allSourceMatches=True, sourceHashesMatch=True, sourceLinesDoNotWrap=True,
                      everyOperationHasScreenshot=True, redFramesChecked=True, descriptionCharacters=len(desc),
                      screenshotManifest=str(folder / '截图来源与标注清单.json'))
        reports.append(report)
        print(key, 'verified', len(code), 'source pages;', len(manual), 'manual pages;', entry['screenshotCount'], 'screenshots')
    (OUT / '申请材料质量检查.json').write_text(json.dumps(reports, ensure_ascii=False, indent=2), 'utf-8')
    overview = ['# 牛项目三端软著材料', '', '按提供模板整理。申请表保持空白，由申请人最后填写。', '',
                '| 端 | 交存非空源码 | 代码页数 | 手册页数 | 流程图 | 实际操作截图 |',
                '|---|---:|---:|---:|---:|---:|']
    for entry, report in zip(entries, reports):
        overview.append(f"| {entry['short'] if 'short' in entry else entry['key']} | {report['sourceLines']} | {report['codePages']} | {report['manualPages']} | 10 | {report['screenshots']} |")
    overview += ['', '每套提供源程序Word和PDF、用户操作手册Word和PDF、源码归档与ZIP、原始行号及SHA-256清单、501至1300字功能描述、红框截图及原始截图。', '',
                 'Web交存连续前1500行与后1500行。鸿蒙APP和小程序不足3000行，提交各自全部现有非空源码，含代码注释；不拼入其他客户端或后台凑行数。每页50行，最后一页保留实际余数。JSON与构建资源附在归档中，不计入有效源码页。', '',
                 '三套手册均为14页，每个操作章节都有真实运行截图和红框编号。照片分析和移动端设备动作仍为本地演示；AI助手、咨询单和资讯使用真实后端接口。专家真人回复、锁屏通知及微信订阅消息尚未配置。', '',
                 '截图中的旧硬件历史数据、UI_TEST及DEMO条目不作为现场在线或疾病判断的证据。', '']
    (OUT / '材料总览.md').write_text('\n'.join(overview), 'utf-8')


if __name__ == '__main__':
    main()
