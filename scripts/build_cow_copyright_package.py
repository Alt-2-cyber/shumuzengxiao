from __future__ import annotations

import csv
import hashlib
import json
import math
import shutil
import zipfile
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from cow_copyright_content import APPS
from cow_source_template_format import write_source_document


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(r"D:\table\软著申请\牛")
TEMPLATES = Path(r"D:\Users\q4984\xwechat_files\wxid_rt7txr6f1vcl22_d661\msg\file\2026-10")
CODE_TEMPLATE = TEMPLATES / "代码模板(1)(2).docx"
MANUAL_TEMPLATE = TEMPLATES / "技术交流智能化交互系统手册的副本.docx"
FORM_TEMPLATE = TEMPLATES / "申请表模板.docx"
WORK = ROOT / ".run" / "cow-copyright"
VERSION = "V1.0"
TODAY = "2026-10-23"
WIDTH_PT = 420


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def font(run, size=10.5, bold=False, east="宋体", latin="Calibri", color=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def no_grid(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    elem = OxmlElement("w:snapToGrid")
    elem.set(qn("w:val"), "0")
    ppr.append(elem)


def new_from_template(template: Path, software: str):
    d = Document(template)
    section_properties = deepcopy(d.sections[-1]._sectPr)
    for node in list(d._element.body):
        d._element.body.remove(node)
    d._element.body.append(section_properties)
    for rid, relationship in list(d.part.rels.items()):
        if relationship.reltype.endswith("/image"):
            d.part.drop_rel(rid)
    # Remove template comments, headers and unrelated metadata from the new files.
    for rid, relationship in list(d.part.rels.items()):
        if relationship.reltype.endswith(("/comments", "/commentsExtended", "/people")):
            d.part.drop_rel(rid)
    section = d.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for automatic_line_number in list(section._sectPr.findall(qn("w:lnNumType"))):
        section._sectPr.remove(automatic_line_number)
    for area in (section.header, section.footer, section.first_page_header, section.first_page_footer):
        for paragraph in area.paragraphs:
            paragraph.clear()
    d.core_properties.title = software + " " + VERSION
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.core_properties.subject = "软件著作权申请材料"
    d.core_properties.comments = "按提供模板及现有源码整理"
    return d


def header_footer(d, name, label):
    s = d.sections[0]
    s.header_distance, s.footer_distance = Cm(1.15), Cm(1.15)
    p = s.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(name + " " + VERSION), 9, east="宋体", latin="Arial")
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(label + "  |  第 "), 8.5)
    for field_name, suffix in (("PAGE", " 页 / 共 "), ("NUMPAGES", " 页")):
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        p._p.append(field)
        font(p.add_run(suffix), 8.5)


def ptext(d, value, size=10.5, bold=False, after=6, before=0, alignment=None):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(17)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.widow_control = False
    if alignment is not None:
        p.alignment = alignment
    no_grid(p)
    font(p.add_run(value), size, bold)
    return p


def heading(d, text, page_break=False):
    p = d.add_paragraph()
    p.style = d.styles["Heading 1"]
    p.paragraph_format.page_break_before = page_break
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = Pt(21)
    p.paragraph_format.keep_with_next = True
    no_grid(p)
    font(p.add_run(text), 15, True, east="黑体", latin="Arial")
    return p


def table(d, rows):
    t = d.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(3.35), Cm(13.6)
    for key, value in rows:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Cm(3.35), Cm(13.6)
        for cell, text in zip(cells, (key, value)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            pp = cell.paragraphs[0]
            pp.paragraph_format.space_after = Pt(3)
            pp.paragraph_format.space_before = Pt(3)
            pp.paragraph_format.line_spacing = Pt(15)
            font(pp.add_run(text), 9.5)
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + side)
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "888888")
                borders.append(border)
            tcpr.append(borders)
    return t


def source_files(app):
    files = []
    for directory, suffixes in app["roots"]:
        files.extend(p for p in (ROOT / directory).rglob("*") if p.is_file() and p.suffix in suffixes)
    if app["key"] == "web":
        java_main = ROOT / "smart-cattle-backend/src/main/java/com/shumuzengxiao/smartcattle"
        first = [java_main / "SmartCattleApplication.java", java_main / "service/TelemetryService.java",
                 java_main / "service/HealthScreeningService.java"]
        order = {p: index for index, p in enumerate(first)}
        java = sorted((p for p in files if p.suffix == ".java"), key=lambda p: (order.get(p, 100), p.as_posix()))
        frontend = [p for p in files if p.suffix != ".java"]
        def front_key(p):
            relative = p.relative_to(ROOT / "smart-cattle-frontend/src").as_posix()
            group = 0 if relative.startswith("assets/") else 1 if relative.startswith(("components/", "effects/")) else 3 if relative.startswith("views/") else 2
            return group, relative
        return java + sorted(frontend, key=front_key)
    return sorted(files, key=lambda p: p.relative_to(ROOT).as_posix())


def prepare_source(app, folder):
    all_records, manifest = [], []
    full_lines = 0
    archive_root = folder / "源码归档"
    archive_root.mkdir(exist_ok=True)
    for f in source_files(app):
        data = f.read_bytes()
        rel = f.relative_to(ROOT).as_posix()
        lines = data.decode("utf-8-sig").splitlines()
        full_lines += len(lines)
        entry = dict(path=rel, physicalLines=len(lines), nonemptyLines=sum(bool(line.strip()) for line in lines), sha256=digest(data))
        manifest.append(entry)
        dest = archive_root / rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(data)
        for original_line, line in enumerate(lines, 1):
            if line.strip():
                all_records.append(dict(file=rel, originalLine=original_line, text=line, index=len(all_records)+1))
    # Supplemental application descriptors and project configuration are archived, not used to pad source pages.
    supplemental = []
    if app["key"] == "hm":
        supplemental = [p for root in (ROOT / "hm/AppScope", ROOT / "hm/entry/src/main/resources") for p in root.rglob("*") if p.is_file()]
        supplemental += [ROOT / p for p in ("hm/entry/src/main/module.json5", "hm/build-profile.json5", "hm/oh-package.json5", "hm/entry/oh-package.json5", "hm/hvigorfile.ts", "hm/entry/hvigorfile.ts")]
    elif app["key"] == "wx":
        supplemental = list((ROOT / "wx/miniprogram").rglob("*.json"))
        supplemental += [ROOT / p for p in ("wx/project.config.json", "wx/tsconfig.json", "wx/package.json")]
    else:
        supplemental = [ROOT / p for p in ("smart-cattle-backend/pom.xml", "smart-cattle-frontend/package.json", "smart-cattle-frontend/vite.config.js")]
    for f in supplemental:
        if f.exists():
            dest = archive_root / f.relative_to(ROOT)
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, dest)
    selected = all_records if len(all_records) < 3000 else all_records[:1500] + all_records[-1500:]
    pages = [selected[i:i+50] for i in range(0, len(selected), 50)]
    app["stats"] = dict(fileCount=len(manifest), physicalLines=full_lines, nonemptyLines=len(all_records),
                        selectedLines=len(selected), codePages=len(pages), submittedAll=len(all_records)<3000,
                        descriptionsCharacters=len(app["description"]), workflowFigures=10)
    (folder / "源程序文件与行号清单.json").write_text(json.dumps(dict(stats=app["stats"], files=manifest, selected=selected), ensure_ascii=False, indent=2), encoding="utf-8")
    with (folder / "源码文件统计.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["path", "physicalLines", "nonemptyLines", "sha256"])
        writer.writeheader()
        writer.writerows(manifest)
    with zipfile.ZipFile(folder / "完整源码归档.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(archive_root.rglob("*")):
            if path.is_file():
                zf.write(path, path.relative_to(archive_root).as_posix())
    return pages


def source_document(app, folder, pages):
    output = folder / (app["name"] + "V1.0-源程序.docx")
    return write_source_document(app['name'], output, pages)


def wrap_label(text, maximum=12):
    return [text[i:i+maximum] for i in range(0,len(text),maximum)]


def draw_flow(app, number, flow, path):
    width, height = 1600, 440
    image = Image.new("RGB", (width,height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 29)
    bold = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 33)
    small = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 24)
    draw.text((38,24), f"图 {number}  {flow['title']}流程", font=bold, fill="#111111")
    left, y, box_w, box_h, gap = 35, 125, 335, 105, 61
    for i,node in enumerate(flow["nodes"]):
        x = left+i*(box_w+gap)
        draw.rectangle((x,y,x+box_w,y+box_h), outline="#333333", width=3, fill="#F3F7F5")
        labels = wrap_label(node,11)
        for j,line in enumerate(labels):
            size = draw.textbbox((0,0),line,font=regular)
            draw.text((x+(box_w-(size[2]-size[0]))/2,y+18+j*37),line,font=regular,fill="#111111")
        if i<3:
            x0=x+box_w+5; x1=x+box_w+gap-8; yy=y+box_h/2
            draw.line((x0,yy,x1,yy),fill="#333333",width=4)
            draw.polygon([(x1,yy),(x1-14,yy-9),(x1-14,yy+9)],fill="#333333")
    draw.line((left+box_w/2,y+box_h,left+box_w/2,315,left+3*(box_w+gap)+box_w/2,315,left+3*(box_w+gap)+box_w/2,y+box_h+2),fill="#777777",width=2)
    draw.text((left+box_w/2+24,335),"异常、取消或条件不足：按本页处理说明提示或停止；不伪造成功状态。",font=small,fill="#555555")
    if app["key"] != "web":
        draw.text((38,390),"本图对应现有客户端代码；分析结果与设备动作均按本地演示范围说明。",font=small,fill="#555555")
    else:
        draw.text((38,390),"本图对应现有浏览器前端及后端接口；现场执行以设备回执与实际数据为准。",font=small,fill="#555555")
    image.save(path)


def add_picture(d,path,width_cm=16.8,height_cm=None):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r=p.add_run()
    if height_cm:
        r.add_picture(str(path),height=Cm(height_cm))
    else:
        r.add_picture(str(path),width=Cm(width_cm))


def manual(app,folder):
    d = new_from_template(MANUAL_TEMPLATE,app["name"])
    s=d.sections[0]
    s.top_margin,s.bottom_margin=Cm(2.54),Cm(2.54)
    s.left_margin,s.right_margin=Cm(1.91),Cm(1.91)
    header_footer(d,app["name"],"用户操作手册")
    figures=folder / "配图"
    figures.mkdir(exist_ok=True)
    ptext(d,app["name"],size=20,bold=True,before=82,after=22,alignment=WD_ALIGN_PARAGRAPH.CENTER)
    ptext(d,VERSION,size=16,bold=True,after=38,alignment=WD_ALIGN_PARAGRAPH.CENTER)
    ptext(d,"用户操作手册",size=22,bold=True,after=48,alignment=WD_ALIGN_PARAGRAPH.CENTER)
    ptext(d,"客户端类型："+app["short"],size=12,after=12,alignment=WD_ALIGN_PARAGRAPH.CENTER)
    ptext(d,app["summary"],size=11,after=12)
    ptext(d,"版本范围："+app["status"],size=10.5)
    ptext(d,"本文档根据现有源程序整理，操作入口、结果和限制与当前代码对应。",size=10.5)
    heading(d,"软件概述、运行条件与阅读顺序",True)
    ptext(d,app["description"][:app["description"].find("。")+1])
    table(d,[("版本",VERSION),("开发语言",app["language"]),("开发工具",app["tools"]),
             ("运行平台",app["platform"]),("支撑环境",app["runtime"]),("软件范围",app["boundaries"])])
    ptext(d,"使用前请先确认设备或运行环境符合上表。以下十个章节按现有页面与数据处理顺序编排，每章配一张流程图；不同的异常状态和演示数据范围在对应章节说明。",before=8)
    ptext(d,"操作章节："+"；".join(f"{i}、{f['title']}" for i,f in enumerate(app["workflows"],1))+"。")
    ptext(d,"当前实现说明："+app["status"])
    ptext(d,"截图中的品牌短名为“数牧增效”，本手册的软件全称以封面和页眉为准；示例数值不作为运行准确率或临床有效性的证明。")
    for i,flow in enumerate(app["workflows"],1):
        heading(d,f"{i}、{flow['title']}",True)
        ptext(d,flow["purpose"],after=9)
        image_path=figures/f"流程图{i:02d}-{flow['title']}.png"
        draw_flow(app,i,flow,image_path)
        add_picture(d,image_path)
        ptext(d,f"图{i}  {flow['title']}流程图",size=9,after=9,alignment=WD_ALIGN_PARAGRAPH.CENTER)
        for j,step in enumerate(flow["steps"],1):
            ptext(d,f"（{j}）{step}",after=5)
        ptext(d,"操作结果："+flow["result"],after=7)
        ptext(d,"异常处理："+flow["exception"],after=7)
        ptext(d,"代码对应："+flow["source"],size=9,after=0)
    heading(d,"11、界面及模块对应",True)
    if app["screenshots"]:
        ptext(d,"以下图片为本项目当前运行环境的真实界面截图。红框和红色编号表示下一步操作位置；截图中的日期、设备在线状态、示例指标和网络状态按截图采集时的实际页面显示，不能替代现场设备实测。", after=8)
        screenshot_table=d.add_table(rows=1 if app["key"]=="hm" else 2,cols=4 if app["key"]=="hm" else 2)
        screenshot_table.autofit=False
        for idx,(filename,caption) in enumerate(app["screenshots"]):
            path=WORK/"screens"/"annotated"/filename
            if not path.exists():
                raise FileNotFoundError(path)
            shutil.copy2(path,figures/filename)
            if app["key"]=="hm":
                cell=screenshot_table.rows[0].cells[idx]
                image_w=3.9
            else:
                cell=screenshot_table.rows[idx//2].cells[idx%2]
                image_w=8.25
            cell.text=""
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path),width=Cm(image_w))
            p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=Pt(13)
            font(p.add_run(f"图{11+idx} {caption}（红框为点击位置）"),8.5)
        ptext(d,"截图说明：界面标题、按钮和提示保留实际程序文案。软件名称不意味着额外功能，设备执行、识别有效性和联网能力均以对应章节的真实实现范围为准。",before=9)
    else:
        ptext(d,"当前未附可核验的微信模拟器截图，以下模块表依据实际WXML、TypeScript和WXSS源码整理，不将网页或鸿蒙截图冒作小程序界面。前十章的十张流程图已覆盖本版本主要交互。")
        table(d,[("总览", "pages/index：本地记录摘要、最新结果、示例设备与告警"),
                 ("AI识牛", "pages/diagnosis：chooseMedia选图、进度、模拟结果及重选"),
                 ("记录", "pages/records：缓存列表、空状态和清空确认"),
                 ("管理", "pages/admin：静态设备、刷新与会话操作日志"),
                 ("导航", "custom-tab-bar：switchTab跳转与选中高亮"),
                 ("数据", "utils/data.ts：示例数据、StorageSync、格式化与映射")])
        ptext(d,"源码中页面标签为“总览、AI识牛、记录、管理”，与app.json声明的路径一致；页面名称含AI不代表已经接入真实图像推理。",before=12)
    heading(d,"12、常见问题与版本说明",True)
    if app["key"]=="web":
        entries=[("页面没有数据", "核对前后端地址、端口、数据库、API权限和设备上报；当前截图中的空状态不能作为设备故障结论。"),
                 ("设备指令无效", "先核对控制链路和设备供电，再查命令记录、ACK及实际遥测，未接通按钮不应反复点击。"),
                 ("图片或音频不可用", "检查原件是否归档、文件路径与访问角色，元数据存在不代表原件一定可回放。"),
                 ("健康或模型状态不明", "核对数据覆盖、基线、证据新鲜度和审批版本。平台为风险辅助筛查，不能代替兽医确诊。")]
    else:
        entries=[("没有识别记录", "先在照片页面选择图片并完成一次本地演示。记录仅保存在当前客户端运行环境，清理数据后不会从服务器恢复。"),
                 ("数值总是相同", "设备列表、在线率和部分指标是固定示例，并不进行实时采样；刷新只重读本地数据。"),
                 ("提示已下发但设备未动", "当前按钮只写会话日志并显示提示，没有网络控制请求和执行回执，不能据此认定硬件故障。"),
                 ("照片结果能否诊断疾病", "不能。当前createMockAnalysis按路径和时间生成示例结果，未运行真实图像模型。需另行开发和验证实际推理链路。")]
    for title,answer in entries:
        ptext(d,title,size=11,bold=True,after=5)
        ptext(d,answer,after=10)
    ptext(d,"版本说明：本手册对应"+app["name"]+" "+VERSION+"的当前源码范围。软件全称、版本与源程序材料一致，未使用模板原系统代码、截图或业务说明。")
    ptext(d,"数据保护：清空、卸载和清理缓存前自行保留必要资料；不要把真实密码、密钥或个人资料写入公开演示内容。")
    ptext(d,"代码与文档对应：程序材料附有文件哈希、有效行统计及原文件行号清单。后续功能调整应同步更新手册，不能仅更换版本文字。")
    output=folder/(app["name"]+"V1.0-用户操作手册.docx")
    d.save(output)
    return output


def auxiliary(app,folder):
    stats=app["stats"]
    assert 501 <= len(app["description"]) <= 1300
    (folder/"功能描述-填表用.txt").write_text(app["description"]+"\n",encoding="utf-8")
    (folder/"功能描述字数与技术信息.txt").write_text(
        f"软件全称：{app['name']}\n版本号：{VERSION}\n软件分类：应用软件\n"
        f"功能描述字数：{len(app['description'])}（不含末尾换行）\n"
        f"源程序物理行数：{stats['physicalLines']}\n非空源代码行数：{stats['nonemptyLines']}\n"
        f"代码材料页数：{stats['codePages']}；每页50行，最后一页不足50行据实保留。\n"
        f"开发语言：{app['language']}\n开发工具：{app['tools']}\n运行环境：{app['platform']}\n"
        f"支撑环境：{app['runtime']}\n技术特点：{app['technical']}\n"
        "申请表按用户要求保持空白，申请人、权利人、完成日期和发表状态由本人最后填写。\n",
        encoding="utf-8")
    form_copy = folder/"申请表模板-待本人填写.docx"
    if not form_copy.exists():
        shutil.copy2(FORM_TEMPLATE,form_copy)
    mode=("全部客户端源程序" if stats["submittedAll"] else "前1500行与后1500行，共3000行")
    (folder/"源代码数量及交存说明.md").write_text(
        f"# {app['name']} {VERSION}\n\n"
        f"- 当前源文件：{stats['fileCount']} 个。\n- 原始物理行：{stats['physicalLines']}；非空源代码：{stats['nonemptyLines']}。\n"
        f"- 本次材料：{mode}；{stats['selectedLines']}行，{stats['codePages']}页。\n"
        "- 空行不作为有效代码；没有复制其他客户端、后台或模板代码凑3000行。长行以字体横向压缩保持单行，不截断代码。\n"
        f"- 源码边界：{app['boundaries']}\n- 说明：{app['status']}\n"
        "- 原文件行号和SHA-256见JSON/CSV清单，完整源码见源码归档及ZIP。构建、资源和JSON配置单列归档，不借其凑有效代码行。\n"
        "- 本次代码、说明书和流程图均按当前源码生成；当前版本尚待本人确认完成日期与权利信息，未回填旧申请表日期。\n"
        "- 操作手册按12–15页范围编排、包含10张源码流程图和当前客户端真实截图；PDF由本机Word导出后检查实际页数。\n"
        "- 空白申请表为提供模板原件副本，未填写；功能描述在独立TXT中供最终填表使用。\n",
        encoding="utf-8")


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    WORK.mkdir(parents=True,exist_ok=True)
    created=[]
    for app in APPS:
        folder=OUT/app["name"]
        folder.mkdir(exist_ok=True)
        pages=prepare_source(app,folder)
        code=source_document(app,folder,pages)
        book=manual(app,folder)
        auxiliary(app,folder)
        created.append(dict(name=app["name"],key=app["key"],stats=app["stats"],code=str(code),manual=str(book)))
        print(json.dumps(created[-1],ensure_ascii=False))
    (OUT/"材料生成清单.json").write_text(json.dumps(created,ensure_ascii=False,indent=2),encoding="utf-8")
    template_dir=OUT/"原始参考模板"
    template_dir.mkdir(exist_ok=True)
    for source in [CODE_TEMPLATE,MANUAL_TEMPLATE,FORM_TEMPLATE]:
        if not (template_dir/source.name).exists():
            shutil.copy2(source,template_dir/source.name)
    (OUT/"材料总览.md").write_text(
        "# 牛项目三项软件著作权材料\n\n"
        "申请表按要求保持空白，由本人最后填写；各文件夹另有501–1300字功能描述。\n\n"
        +"\n".join(f"- {a['name']} {VERSION}：源码{a['stats']['physicalLines']}物理行，非空{a['stats']['nonemptyLines']}行；交存{a['stats']['selectedLines']}行/{a['stats']['codePages']}页；手册计划14页/10张流程图。" for a in APPS)
        +"\n\nWeb包含后台Java与Web客户端；鸿蒙和微信代码量不足3000行，提交各自全部现有客户端代码。未修改业务程序、未增加虚假代码。\n\n"
        "鸿蒙与小程序现有照片分析和设备管理为本地演示流程；文档、流程图和功能描述已如实注明。\n\n"
        "每套提供可编辑Word、由Word导出的PDF、源码归档、行数及文件哈希清单、带红框编号的真实界面截图、流程图和空白申请表副本。\n",
        encoding="utf-8")


if __name__=="__main__":
    main()
