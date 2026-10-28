from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\table\shumuzengxiao")
OUT = ROOT / "outputs"
TEMPLATE = OUT / "template_challenge_converted.docx"
FINAL_ASCII = OUT / "muyunzhishe_challenge_cup_business_plan.docx"
FINAL = OUT / "牧云智舍-挑战杯商业计划书.docx"
POSTERS = OUT / "imagegen"

FONT = "微软雅黑"
BLUE = "1F4E79"
DARK_BLUE = "17365D"
LIGHT_BLUE = "EAF3F8"
LIGHT_GRAY = "F6F8FA"
INK = "111111"
MUTED = "666666"
TABLE_WIDTH = 9360


def set_east_asia_font(run, name: str = FONT) -> None:
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def style_run(run, size: float = 10.5, bold: bool = False, color: str = INK) -> None:
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, color=color)


def clear_doc(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)

    heading_tokens = {
        "Heading 1": (16, DARK_BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, "2F5597", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def add_footer(section) -> None:
    section.footer_distance = Inches(0.49)
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("牧云智舍：多模态智能牛舍监测与闭环预警平台")
    style_run(run, size=9, color=MUTED)


def add_para(
    doc: Document,
    text: str = "",
    *,
    size: float = 10.5,
    bold: bool = False,
    color: str = INK,
    align=None,
    before: float = 0,
    after: float = 8,
    first_line: bool = True,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        style_run(run, size=size, bold=bold, color=color)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_east_asia_font(run)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE if level == 1 else BLUE)
        run.font.size = Pt(16 if level == 1 else 13)


def add_key_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    widths = [2200, TABLE_WIDTH - 2200]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for idx, (key, value) in enumerate(rows):
        c0, c1 = table.rows[idx].cells
        set_cell_text(c0, key, bold=True, color="FFFFFF", size=9.8)
        set_cell_text(c1, value, size=9.8)
        set_cell_shading(c0, BLUE)
        if idx % 2 == 0:
            set_cell_shading(c1, LIGHT_BLUE)
    add_para(doc, "", after=2, first_line=False)


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=9.2)
        set_cell_shading(cell, BLUE)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for col_idx, value in enumerate(row):
            set_cell_width(cells[col_idx], widths[col_idx])
            set_cell_margins(cells[col_idx])
            set_cell_text(cells[col_idx], value, size=9)
            if row_idx % 2 == 0:
                set_cell_shading(cells[col_idx], LIGHT_GRAY)
    set_table_geometry(table, widths)
    add_para(doc, "", after=2, first_line=False)


def add_points_table(doc: Document, title: str, rows: list[tuple[str, str]]) -> None:
    add_heading(doc, title, 2)
    add_matrix(doc, ["要点", "说明"], [[a, b] for a, b in rows], [2300, TABLE_WIDTH - 2300])


def add_image(doc: Document, path: Path, caption: str, width: float = 6.45) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_para(doc, caption, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=8)


def add_cover(doc: Document) -> None:
    add_para(
        doc,
        '2026年“挑战杯”大学生创业计划竞赛',
        size=18,
        bold=True,
        color=BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
        before=18,
        after=18,
    )
    add_para(doc, "牧云智舍", size=30, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=28, after=4)
    add_para(
        doc,
        "多模态智能牛舍监测与闭环预警平台",
        size=20,
        bold=True,
        color=BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
        after=12,
    )
    add_para(
        doc,
        "低成本硬件 + 边缘 AI + 多端协同",
        size=13,
        color="3A6EA5",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
        after=14,
    )
    add_image(doc, POSTERS / "01-cover.png", "图0-1 项目整体概念与端边云协同闭环示意", width=6.55)
    add_key_table(
        doc,
        [
            ("项目名称", "牧云智舍：多模态智能牛舍监测与闭环预警平台"),
            ("项目方向", "数字农业 / 智慧畜牧 / 动物健康风险早筛 / 低成本物联网"),
            ("核心方案", "ESP32-S3 环境控制节点、MaixCam 边缘视觉、INMP441 声音采集、Spring Boot 云端服务、Web/鸿蒙 APP/微信小程序多端协同"),
            ("应用对象", "中小规模肉牛养殖场、家庭牧场、教学实训基地与示范牛舍"),
            ("文档类型", "商业计划书"),
            ("日期", "2026 年 6 月"),
        ],
    )
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_heading(doc, "目录", 1)
    for item in [
        "一、执行摘要",
        "二、产品介绍",
        "三、行业与市场分析",
        "四、商业模式",
        "五、团队与组织计划",
        "六、财务预测",
        "七、教育维度与社会价值",
        "八、风险控制与实施计划",
        "九、附件：技术成果与证明材料",
    ]:
        add_para(doc, item, size=11.5, bold=True, color=BLUE, first_line=False, after=3)
    doc.add_page_break()


def build_document() -> None:
    doc = Document(str(TEMPLATE))
    clear_doc(doc)
    setup_styles(doc)
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    add_footer(section)

    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "一、执行摘要", 1)
    add_heading(doc, "1. 项目背景", 2)
    add_para(doc, "肉牛养殖正在从经验巡栏走向数据化、连续化和智能化管理。传统牛舍主要依赖人工观察牛只精神状态、采食情况和环境变化，巡检频次有限，异常发现滞后，证据留存不足。温湿度、氨气、烟雾、通风状态等环境因素会影响牛只健康和生产效率，咳嗽、喘鸣、长时间卧倒、活动减少、眼鼻分泌物等信号也需要多模态联合判断。")
    add_para(doc, "牧云智舍面向中小规模养殖场和普通农户，以低成本硬件、边缘 AI 和多端应用构成可持续运行的主动监测平台。项目定位不是替代兽医诊断，而是为养殖者提供“早发现、可复核、可追溯、可处置”的辅助巡检工具。")
    add_heading(doc, "2. 产品概述", 2)
    add_para(doc, "平台由端侧采集、边缘推理、云端服务、业务应用和运维复盘五个层次构成。端侧由 ESP32-S3 采集温湿度、氨气、光照、烟雾等环境数据并控制风扇、雾化、声光报警等执行器；边缘侧由 MaixCam 运行轻量视觉模型，由 INMP441 等声音节点采集并筛选牛舍声音事件；云端采用 Spring Boot、MySQL/H2、告警策略和模型版本管理；应用层提供 Web 驾驶舱、鸿蒙 APP 和微信小程序。")
    add_matrix(
        doc,
        ["能力", "产品实现", "用户价值"],
        [
            ["连续感知", "环境、声音、视觉与设备状态持续采集", "减少人工巡检盲区"],
            ["边缘初判", "MaixCam 本地视觉事件、声音特征筛选、阈值判断", "降低视频长期上传成本"],
            ["风险预警", "规则评分与多模态证据链生成", "让异常更早进入复核流程"],
            ["闭环处置", "告警确认、命令下发、处置回执、模型迭代", "形成可追溯管理闭环"],
        ],
        [1700, 4050, 3610],
    )
    add_points_table(
        doc,
        "3. 产品优势与市场",
        [
            ("低成本部署", "采用 ESP32-S3、MaixCam、I2S 麦克风与常见环境传感器构建原型，适合中小牛舍渐进式部署。"),
            ("多模态融合", "环境阈值、视觉事件、声音事件、历史趋势和人工复核共同形成风险评分，避免单一指标误判。"),
            ("证据链可追溯", "告警记录关联环境曲线、截图、音频片段、控制命令和处置回执，便于复盘。"),
            ("多端协同", "Web 面向管理人员，鸿蒙 APP 和微信小程序面向现场巡检人员，覆盖不同使用场景。"),
        ],
    )
    add_heading(doc, "4. 商业模式", 2)
    add_para(doc, "项目以“硬件节点 + 平台服务 + 运维支持”为核心收入结构，先从教学实训、示范牛舍和中小养殖场试点切入，再向规模化场区和区域农服机构扩展。早期以样机销售、试点实施和平台部署服务为主，后续通过 SaaS 订阅、模型升级、设备 OTA、数据报表和运维托管形成持续收入。")
    add_heading(doc, "5. 财务与投资", 2)
    add_para(doc, "项目处于原型到试点阶段，资金重点投入硬件样机、传感器标定、样本采集、模型训练、平台部署和试点运维。预计通过学校创新创业资源、竞赛孵化资金、地方农业数字化项目和企业试点合作获得启动资金。")

    add_heading(doc, "二、产品介绍", 1)
    add_heading(doc, "1. 项目背景", 2)
    add_para(doc, "智慧畜牧的核心不是简单安装摄像头，而是把环境、动物行为、声音事件、设备状态和人工处置连接为闭环。对中小牛舍而言，高成本全场景视频系统和复杂工业控制系统往往部署门槛高、维护难度大。本项目选择以低成本节点和轻量 AI 为切入点，让普通牛舍也能逐步具备连续感知和主动预警能力。")
    add_image(doc, POSTERS / "02-background-pain-points.png", "图2-1 背景与痛点：从人工巡检走向连续感知")
    add_heading(doc, "2. 产品架构", 2)
    add_image(doc, POSTERS / "03-architecture-loop.png", "图2-2 总体架构：端、边、云、服、应用闭环")
    add_matrix(
        doc,
        ["层级", "组成", "关键功能"],
        [
            ["端侧采集", "ESP32-S3、环境传感器、继电器、执行器", "环境采样、本地阈值判断、风扇/雾化/报警控制"],
            ["边缘推理", "MaixCam、INMP441、轻量模型、断点缓存", "视觉事件、声音事件、抓拍留证、断网补传"],
            ["云端服务", "Spring Boot API、MySQL/H2、告警策略、模型版本", "数据持久化、风险融合、接口服务、策略管理"],
            ["业务应用", "Web 驾驶舱、鸿蒙 APP、微信小程序", "总览分析、现场巡检、告警处置、设备控制"],
            ["运维复盘", "告警确认、命令下发、处置回执、训练迭代", "形成持续优化闭环"],
        ],
        [1700, 3500, 4160],
    )
    add_heading(doc, "3. 核心硬件与算法", 2)
    add_image(doc, POSTERS / "04-hardware-multimodal-ai.png", "图2-3 硬件与多模态 AI：低成本节点支撑主动预警")
    add_matrix(
        doc,
        ["模块", "技术实现", "作用"],
        [
            ["ESP32-S3 主控", "WiFi、MQTT、UART、I2S、GPIO/PWM", "环境采集与执行器控制核心"],
            ["MaixCam 视觉节点", "轻量 YOLO、事件截图、UART/MQTT 上报", "生成牛只姿态、异常候选和视觉证据"],
            ["INMP441 声音节点", "I2S 采集、VAD、Log-Mel/MFCC、轻量分类", "识别咳嗽、喘鸣、异常叫声等声音事件"],
            ["多模态融合", "环境阈值、视觉事件、声音事件、历史趋势、人工复核", "输出风险等级、证据链和处置建议"],
        ],
        [1800, 3900, 3660],
    )
    add_heading(doc, "4. 软件成果", 2)
    add_image(doc, POSTERS / "05-software-results.png", "图2-4 软件成果展示：Web、鸿蒙 APP、小程序协同")
    add_para(doc, "当前项目已完成 Spring Boot 后端、Vue Web 看板、鸿蒙 APP、微信小程序和数据库脚本。Web 驾驶舱展示存栏、在线设备、环境趋势、风险告警和 3D 牛舍态势；移动端覆盖告警列表、证据查看、状态确认、设备控制和闭环记录。")

    add_heading(doc, "三、行业与市场分析", 1)
    add_heading(doc, "1. 市场现状", 2)
    add_para(doc, "数字农业、智慧畜牧、动物疫病防控和绿色养殖共同推动养殖场基础设施升级。规模化养殖场具备较强的信息化投入能力，但大量中小牛舍仍处于人工巡检、分散记录和经验判断阶段。项目目标市场首先聚焦“愿意数字化但预算有限”的中小养殖主体，并向教学实训、示范基地和区域农服机构延伸。")
    add_matrix(
        doc,
        ["目标客户", "痛点", "切入价值"],
        [
            ["中小牛舍/家庭牧场", "人工巡检不连续、异常发现滞后、缺少记录", "低成本节点与移动端巡检闭环"],
            ["规模化场区", "多栏位管理、设备联动和证据归档需求高", "多点部署、Web 驾驶舱和告警策略"],
            ["教学实训基地", "需要可演示、可讲解的智慧畜牧平台", "端边云服完整链路和可视化教学"],
            ["农业示范基地", "需要展示数字农业成果和数据化运营能力", "可视化大屏、巡检记录和闭环管理"],
        ],
        [2200, 3600, 3560],
    )
    add_heading(doc, "2. 竞争与差异化", 2)
    add_para(doc, "传统方案多集中在单一视频监控、单点环境采集或高成本工业控制系统。本项目差异化在于将低成本硬件、边缘 AI、声音识别、环境联动和多端闭环组合为一体，并明确定位为“辅助巡检与预警”而非夸大为自动诊断系统。")
    add_matrix(
        doc,
        ["比较维度", "传统单点方案", "牧云智舍方案"],
        [
            ["部署成本", "视频监控或工业传感器一次性投入较高", "低成本节点先试点，按栏位扩展"],
            ["数据类型", "环境或视频单一数据源", "环境、视觉、声音、历史趋势融合"],
            ["处置闭环", "多为展示和人工记录", "告警确认、命令下发、回执归档"],
            ["可解释性", "告警原因不完整", "保留截图、音频、曲线和策略依据"],
        ],
        [1900, 3700, 3760],
    )
    add_points_table(
        doc,
        "3. 市场进入策略",
        [
            ("试点验证", "以学校实验牛舍、合作养殖户和示范基地作为首批试点，验证数据稳定性和巡检流程。"),
            ("分层产品", "形成基础环境监测包、AI 视觉增强包、声音事件增强包和平台闭环服务包。"),
            ("渠道合作", "通过农技服务机构、畜牧设备经销商、智慧农业集成商建立渠道合作。"),
            ("案例传播", "在真实样本积累后，以试点数据、看板截图和闭环记录形成可复制推广材料。"),
        ],
    )

    add_heading(doc, "四、商业模式", 1)
    add_matrix(
        doc,
        ["产品/服务", "内容", "收入方式"],
        [
            ["硬件节点包", "ESP32-S3 环境节点、声音节点、MaixCam 视觉节点、执行器控制模块", "设备销售/租赁"],
            ["平台部署服务", "后端、数据库、Web 驾驶舱、移动端接入与现场配置", "项目实施费"],
            ["SaaS 平台订阅", "告警策略、数据看板、闭环记录、模型版本管理", "按场区/设备数订阅"],
            ["模型与运维服务", "传感器标定、模型迭代、OTA 升级、数据报表", "年度服务费"],
        ],
        [2100, 4600, 2660],
    )
    add_heading(doc, "1. 定价思路", 2)
    add_para(doc, "早期采用试点友好型定价，降低养殖户尝试门槛；成熟后按“硬件成本 + 部署服务 + 平台订阅 + 运维增值”组合收费。基础版覆盖环境监测和设备控制，增强版增加视觉、声音和风险融合，高级版面向规模化场区提供多栏位管理、模型迭代和运维报表。")
    add_points_table(
        doc,
        "2. 渠道策略",
        [
            ("校企合作", "通过农业院校、实训基地和挑战杯项目展示获得首批验证场景。"),
            ("设备渠道", "与畜牧设备商、传感器供应商、弱电施工团队合作，降低现场部署成本。"),
            ("农服渠道", "与区域畜牧服务机构合作，为中小牧场提供数字化改造包。"),
            ("示范传播", "通过真实试点案例、数据看板和闭环记录形成可复制推广材料。"),
        ],
    )

    add_heading(doc, "五、团队与组织计划", 1)
    add_para(doc, "项目团队建议按“硬件与嵌入式、AI 算法、后端与数据库、Web 与移动端、畜牧场景与试点运营、商业拓展”六类角色组织。当前技术报告已覆盖端边云服完整原型，后续应补充真实牛舍试点负责人、兽医/畜牧顾问和供应链合作伙伴。")
    add_matrix(
        doc,
        ["角色", "主要职责", "阶段重点"],
        [
            ["项目负责人", "统筹产品方向、竞赛材料、试点合作和资源协调", "完成商业化路线与试点落地"],
            ["硬件负责人", "ESP32-S3、传感器、执行器、电源与外壳设计", "提升现场稳定性与安全性"],
            ["AI 负责人", "视觉、声音、多模态风险评分和样本标注流程", "形成可验证模型指标"],
            ["软件负责人", "Spring Boot、数据库、Web、鸿蒙 APP、小程序", "实现多端数据一致和告警闭环"],
            ["运营负责人", "养殖户调研、培训、售后与示范基地展示", "验证用户使用流程和付费意愿"],
        ],
        [1800, 4300, 3260],
    )

    add_heading(doc, "六、财务预测", 1)
    add_heading(doc, "1. 财务假设", 2)
    add_para(doc, "由于项目仍处于原型验证阶段，以下预测采用保守假设：第一年以样机试点和教学示范为主，第二年扩大区域试点并形成订阅收入，第三年建立稳定硬件供应与平台运维体系。具体数值需在真实 BOM、部署人力和试点转化率确定后进一步校准。")
    add_matrix(
        doc,
        ["年度", "重点目标", "收入来源", "成本重点"],
        [
            ["第 1 年", "完成 3-5 个试点牛舍和教学示范场景", "样机部署费、竞赛/孵化支持、少量服务费", "硬件打样、传感器标定、样本采集"],
            ["第 2 年", "扩展至 20-50 个栏位/点位", "硬件销售、平台订阅、运维服务", "生产小批量、售后支持、模型迭代"],
            ["第 3 年", "形成区域合作渠道和标准化产品包", "设备包、SaaS、OTA/模型升级、数据报表", "渠道建设、云服务、客服运维"],
        ],
        [1400, 3300, 2800, 1860],
    )
    add_matrix(
        doc,
        ["用途", "说明"],
        [
            ["硬件样机与备件", "ESP32-S3、MaixCam、声音节点、传感器、继电器、电源、防护外壳和安装材料"],
            ["现场试点", "真实牛舍部署、网络调试、环境标定、巡检培训和售后维护"],
            ["数据与模型", "图像/声音样本采集、标注、训练、评估和误报分析"],
            ["平台研发", "Web/移动端优化、权限体系、告警策略、模型管理和 OTA 能力"],
            ["市场推广", "示范基地展示、路演材料、渠道合作和竞赛成果转化"],
        ],
        [2300, 7060],
    )

    add_heading(doc, "七、教育维度与社会价值", 1)
    add_para(doc, "项目具有明显的跨学科实践价值，融合物联网、嵌入式开发、边缘 AI、后端平台、Web 可视化、移动应用、畜牧管理和创新创业训练。作为挑战杯项目，它不仅产出一个软件原型，也能带动学生理解真实农业场景中的工程约束、成本约束和用户使用习惯。")
    add_matrix(
        doc,
        ["价值维度", "说明"],
        [
            ["教学实训", "可作为智慧农业、嵌入式、人工智能和软件工程课程的综合实践平台。"],
            ["乡村振兴", "帮助中小养殖户以较低门槛获得数字化巡检和风险预警能力。"],
            ["绿色养殖", "通过环境联动控制改善通风、氨气和湿度管理，减少资源浪费。"],
            ["就业带动", "未来可形成硬件装配、现场运维、数据标注、模型服务和区域销售岗位。"],
        ],
        [2300, 7060],
    )

    add_heading(doc, "八、风险控制与实施计划", 1)
    add_matrix(
        doc,
        ["风险", "表现", "应对措施"],
        [
            ["传感器漂移", "氨气、温湿度等数据长期偏移", "建立标定流程，预留工业传感器替换接口"],
            ["模型误报/漏报", "牛舍噪声、遮挡、光照变化影响识别", "保留人工复核，持续采集本地样本迭代模型"],
            ["网络不稳定", "农村场景 WiFi/4G 波动影响上报", "边缘缓存关键事件，恢复后补传"],
            ["现场安全", "执行器带载、电源、防水防尘存在风险", "使用隔离继电器、低压控制和防护外壳"],
            ["用户接受度", "养殖户担心复杂难用或维护成本高", "移动端简化流程，提供安装培训和运维服务"],
        ],
        [1800, 3600, 3960],
    )
    add_heading(doc, "1. 路线图", 2)
    add_image(doc, POSTERS / "06-innovation-plan.png", "图8-1 创新应用与计划：从原型到试点部署")
    add_matrix(
        doc,
        ["阶段", "任务", "交付物"],
        [
            ["第 1 阶段", "ESP32-S3 采集、控制、MQTT 上报和后端环境接口联调", "环境节点样机、传感器数据、控制命令闭环"],
            ["第 2 阶段", "MaixCam 跑通轻量 YOLO 模型，完成视觉事件上报", "视觉节点样机、事件截图、推理指标"],
            ["第 3 阶段", "采集本地牛舍图像和声音样本，进行标注、训练和误报分析", "数据集、模型版本、评估报告"],
            ["第 4 阶段", "Web/移动端接入真实数据，完成告警、抓拍、控制和记录归档", "试点系统原型、多端演示材料"],
            ["第 5 阶段", "现场试运行，记录环境、疾病风险、设备故障和人工复核结果", "试点总结、模型优化计划、正式版 BOM"],
        ],
        [1500, 4900, 2960],
    )

    add_heading(doc, "九、附件：技术成果与证明材料", 1)
    add_para(doc, "附件建议随申报材料补充：系统源代码截图、Web/APP/小程序运行截图、后端 API 文档、硬件连接图、模型训练记录、数据集来源说明、试点照片、演示视频二维码、软著/专利/竞赛证明等。当前技术报告已包含系统设计、硬件方案、AI 算法方案、数据库与接口设计、多端实现和测试计划，可作为本商业计划书的技术支撑材料。")
    add_matrix(
        doc,
        ["材料类别", "当前状态", "后续补充"],
        [
            ["软件原型", "Spring Boot 后端、Vue Web、鸿蒙 APP、微信小程序已完成演示原型", "补充部署截图、接口测试结果和演示视频"],
            ["硬件方案", "ESP32-S3、MaixCam、INMP441、传感器与执行器拓扑明确", "补充实物照片、BOM 和安全测试记录"],
            ["算法方案", "视觉、声音、多模态融合方案已设计", "补充真实数据训练集、模型指标和误报分析"],
            ["试点材料", "已形成部署路线图", "补充真实牛舍试点协议、现场记录和用户反馈"],
        ],
        [1900, 4300, 3160],
    )

    doc.save(str(FINAL_ASCII))
    try:
        FINAL.write_bytes(FINAL_ASCII.read_bytes())
    except OSError:
        pass
    print(FINAL_ASCII)


if __name__ == "__main__":
    build_document()
