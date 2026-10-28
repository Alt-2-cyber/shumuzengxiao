from pathlib import Path

from docx import Document


REPORT = Path("D:/table/shumuzengxiao/\u6570\u7267\u589e\u6548\u667a\u80fd\u725b\u820d\u76d1\u6d4b\u5e73\u53f0\u6280\u672f\u62a5\u544a.docx")

REQUIRED = [
    "\u6570\u7267\u589e\u6548",
    "\u667a\u80fd\u725b\u820d\u76d1\u6d4b\u5e73\u53f0\u6280\u672f\u62a5\u544a",
    "ESP32-S3",
    "MaixCam",
    "Spring Boot",
    "Vue3",
    "\u9e3f\u8499 APP",
    "\u5fae\u4fe1\u5c0f\u7a0b\u5e8f",
    "\u591a\u6a21\u6001",
    "\u521b\u65b0\u70b9",
    "\u6d4b\u8bd5\u9a8c\u8bc1",
]


def main():
    doc = Document(REPORT)
    text = "\n".join(p.text for p in doc.paragraphs)
    missing = [s for s in REQUIRED if s not in text]
    if missing:
        raise RuntimeError("missing text: " + ", ".join(missing))
    image_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)
    if image_count < 7:
        raise RuntimeError(f"expected at least 7 images, found {image_count}")
    if len(doc.tables) < 9:
        raise RuntimeError(f"expected at least 9 tables, found {len(doc.tables)}")
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    print("DOCX verification passed")
    print(f"paragraphs={len(doc.paragraphs)}")
    print(f"tables={len(doc.tables)}")
    print(f"images={image_count}")
    print(f"headings={len(headings)}")
    print(f"file={REPORT}")


if __name__ == "__main__":
    main()
