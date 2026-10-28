from __future__ import annotations

import os
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(r"D:\table")
XLSX_PATH = BASE_DIR / "数牧增效-硬件采购清单2.xlsx"
TEMPLATE_PATH = BASE_DIR / "物料申请表(3)(2).docx"
OUTPUT_PATH = BASE_DIR / "数牧增效-物料采购申请报告2.docx"


def read_purchase_rows():
    from openpyxl import load_workbook

    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        name, link_text, variant, qty = row
        if not name:
            continue
        text = link_text or ""
        purpose = ""
        match = re.search(r"用途：(.+)", text, flags=re.S)
        if match:
            purpose = match.group(1).strip()
            purpose = purpose.replace("\n", " ")
        rows.append(
            {
                "idx": idx,
                "name": str(name).strip(),
                "variant": str(variant).strip() if variant else "标准模块/按采购页面默认配置",
                "qty": int(qty) if isinstance(qty, (int, float)) and float(qty).is_integer() else qty,
                "purpose": purpose,
            }
        )
    return rows


DEFAULT_PURPOSES = {
    "MAX30102心率血氧传感器": "作为项目扩展健康体征采集验证模块，用于预研单牛心率、血氧等体征信号接入方式，后续可与视觉、声音、运动数据共同参与健康风险评估。",
    "杜邦线": "用于 ESP32-S3、传感器、继电器、显示屏等模块之间的快速连线和调试，便于原型阶段反复调整硬件连接。",
    "MB-102面包板": "用于原型电路搭建和传感器联调，减少早期焊接成本，方便进行主控、传感器和执行器组合验证。",
    "新款 ADXL345": "作为牛只运动行为采集模块，可安装在腿环或项圈位置，采集三轴加速度数据，用于行走、站立、卧倒和步频异常识别。若后续改用 MPU6500，I2C/SPI 数据链路和算法流程可保持一致。",
}


def clear_body(doc: Document) -> None:
    body = doc.element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is sect_pr:
            continue
        body.remove(child)


def set_cell_text(cell, text, bold=False, size=9, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for para in cell.paragraphs:
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.05


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="C9CDD4", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_paragraph(doc, text="", style=None, align=None, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(5)
    marker = p.add_run("· ")
    marker.font.name = "宋体"
    marker._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    marker.font.size = Pt(10.5)
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    run.bold = True
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return p


def add_competition_table(doc):
    competitions = [
        "天津市大学生创新训练计划（大创）",
        "中国国际大学生创新大赛（互联网+）",
        '"挑战杯"大学生课外学术科技作品竞赛',
        "仁爱杯校级创新创业大赛",
        "开放原子开源大赛",
        "卓越杯校级科技竞赛",
        "中国高校计算机大赛--嵌入式设计赛",
        "全国大学生电子设计竞赛",
        "全国大学生物联网设计竞赛",
        "华北五省（市、自治区）大学生计算机应用大赛天津赛区竞赛",
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed_layout(table)
    widths = [1.1, 14.8]
    headers = ["序号", "赛事名称"]
    for i, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], text, bold=True, size=10, color="FFFFFF")
        shade_cell(table.rows[0].cells[i], "2F5597")
        set_cell_width(table.rows[0].cells[i], widths[i])
    set_repeat_table_header(table.rows[0])
    for idx, name in enumerate(competitions, start=1):
        cells = table.add_row().cells
        set_cell_text(cells[0], idx, size=9)
        set_cell_text(cells[1], name, size=9)
        for i, width in enumerate(widths):
            set_cell_width(cells[i], width)
    set_table_borders(table)


def configure_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    title.font.name = "黑体"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)

    for style_name, size, color in [
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 12, "1F4E79"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Heading 3" else 6)
        style.paragraph_format.space_after = Pt(6 if style_name != "Heading 3" else 4)
        style.paragraph_format.line_spacing = 1.15


def build():
    rows = read_purchase_rows()
    for row in rows:
        if not row["purpose"]:
            row["purpose"] = DEFAULT_PURPOSES.get(row["name"], "用于项目硬件原型搭建、现场调试或功能扩展验证。")

    doc = Document(TEMPLATE_PATH)
    clear_body(doc)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("物料采购申请报告")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = (
        "本次物料采购依据《数牧增效-硬件采购清单2.xlsx》整理，面向“数牧增效--农户普惠型肉牛养殖监测专家”"
        "项目原型搭建与竞赛展示需求。采购内容覆盖 ESP32-S3 主控、MaixCam 边缘 AI 摄像头、环境传感器、声音采集模块、"
        "加速度行为采集模块、执行器联动、电源转换与基础布线辅材等部分，用于形成“现场采集--边缘识别--本地联动--云端展示--告警复核”的闭环系统。"
    )
    add_heading(doc, "一、采购硬件及使用方式", 1)
    add_paragraph(doc, intro)
    add_paragraph(
        doc,
        "系统现场端以 ESP32-S3 为环境与控制节点，负责温湿度、氨气、光照、烟雾、火焰、声音和加速度数据接入，并通过 WiFi/MQTT 上传到后端平台；"
        "MaixCam 作为视觉 AI 节点，负责牛只检测、姿态识别、异常事件抓拍和本地证据留存。两类节点协同工作，可在网络不稳定时先完成本地判断和缓存，网络恢复后再补传关键数据。",
    )
    add_paragraph(
        doc,
        "环境与安全传感器用于监测牛舍温湿度、氨气趋势、光照条件和火情风险；执行器模块用于联动风扇、水泵/雾化、声光报警和补光设备；"
        "INMP441 麦克风用于采集咳嗽、喘鸣、异常叫声等声音事件；ADXL345 加速度模块用于腿环或项圈式运动监测，支撑行走、站立、卧倒和步频异常识别。"
    )
    add_paragraph(
        doc,
        "本采购方案以学生创新训练项目的可落地性为优先目标，尽量选用价格低、接口通用、资料丰富、易调试的模块。"
        "算法侧优先复现公开数据集和开源模型，硬件侧优先完成可演示样机和数据链路闭环，后续再根据真实牛舍环境逐步替换为工业级传感器或防护外壳。"
    )

    add_heading(doc, "二、核心物料明细", 1)
    add_paragraph(doc, "下表为本次采购清单中的核心物料、规格数量及主要用途。采购链接和店铺信息以 Excel 原表为准，正文中不展开长链接，避免影响申请报告版式。")

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed_layout(table)
    headers = ["序号", "物料名称", "规格/颜色分类", "数量", "主要用途"]
    widths = [0.9, 3.4, 4.0, 0.9, 6.7]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=9, color="FFFFFF")
        shade_cell(table.rows[0].cells[i], "2F5597")
        set_cell_width(table.rows[0].cells[i], widths[i])
    set_repeat_table_header(table.rows[0])

    for row in rows:
        cells = table.add_row().cells
        values = [row["idx"], row["name"], row["variant"], row["qty"], row["purpose"]]
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=8.3)
            set_cell_width(cells[i], widths[i])
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_table_borders(table)

    add_heading(doc, "三、采购必要性与落地说明", 1)
    necessity = [
        "保证系统原型完整：清单同时覆盖感知、计算、通信、显示、供电和执行器，能够支撑从牛舍数据采集到前端看板展示的完整闭环，而不是单点功能演示。",
        "降低学生项目实施风险：ESP32-S3、MaixCam、SHT31、BH1750、INMP441、ADXL345 等模块资料成熟、接口通用，便于快速联调、复现实验和后续维护。",
        "支撑多模态健康评估：视觉节点判断牛只姿态和活动异常，声音节点捕捉咳嗽、喘鸣等事件，加速度模块分析步频和卧倒时长，环境节点提供温湿度、氨气和安全背景数据，多源信息融合后可形成疑似异常提示。",
        "适合比赛展示与扩展：OLED、风扇、水泵、继电器和 MOSFET 模块可以直观展示现场联动效果；多块 ESP32-S3 和多个传感器便于分组调试、备份替换和扩展多节点部署。",
        "符合渐进式落地路线：原型阶段以低成本通用模块验证技术路线，正式部署阶段可将 MQ 类气体传感器、接线方式和外壳防护升级为工业级方案，软件接口和数据流向保持基本不变。",
    ]
    for item in necessity:
        add_bullet(doc, item)

    add_heading(doc, "四、预计参与比赛", 1)
    add_competition_table(doc)

    add_heading(doc, "五、申请结论", 1)
    add_paragraph(
        doc,
        "综上，本次采购物料能够支撑“数牧增效”项目完成硬件原型、数据采集、边缘 AI 识别、本地联动控制和平台展示等关键环节。"
        "采购清单整体以低成本、易获取、易调试为原则，符合学生项目经费使用特点，也具备后续升级为真实牛舍试点系统的扩展空间，建议予以采购支持。"
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
