from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path("D:/table/shumuzengxiao")
ASSETS = BASE / "report_assets"
OUTPUT = BASE / "\u6570\u7267\u589e\u6548\u667a\u80fd\u725b\u820d\u76d1\u6d4b\u5e73\u53f0\u6280\u672f\u62a5\u544a.docx"


GREEN = "2F6F46"
DARK = "102216"
LIGHT_GREEN = "E8F6EC"
LIGHT_GRAY = "F7FAF6"
GOLD = "F6E7BC"


def set_run_font(run, font_name="宋体", size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font_size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, "宋体", font_size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9072):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, "宋体", 9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, "宋体", 9)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.9)
    section.bottom_margin = Cm(3.8)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, before, after, align in [
        ("Title", 16, 12, 3, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 16, 17, 16.5, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 14, 13, 13, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12.5, 10, 8, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        st = styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(DARK)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.alignment = align

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header.add_run("数牧增效智能牛舍监测平台技术报告")
    set_run_font(r, "宋体", 9, color="666666")
    add_page_number(section.footer.paragraphs[0])


def add_paragraph(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, "宋体", 12, True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, "宋体", 12)
    else:
        r = p.add_run(text)
        set_run_font(r, "宋体", 12)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_run_font(r, "黑体", {1: 16, 2: 14, 3: 12.5}.get(level, 12), True, DARK)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run_font(r, "宋体", 10, False, "555555")
    return p


def add_picture(doc, path, width_cm=15.8, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if caption:
        add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_GREEN)
        set_cell_margins(cell, 140, 140, 140, 140)
        set_cell_text(cell, h, True, WD_ALIGN_PARAGRAPH.CENTER, 10.5)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_margins(cells[idx], 120, 140, 120, 140)
            set_cell_text(cells[idx], str(value), False, WD_ALIGN_PARAGRAPH.LEFT, 10.5)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("（{}）".format(item[0]))
        set_run_font(r, "宋体", 12, True)
        r2 = p.add_run(item[1])
        set_run_font(r2, "宋体", 12)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    r = p.add_run("数牧增效")
    set_run_font(r, "黑体", 28, True, DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    r = p.add_run("智能牛舍监测平台技术报告")
    set_run_font(r, "黑体", 22, True, DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    r = p.add_run("基于 ESP32-S3 与 MaixCam 的多模态边缘 AI 主动监测系统")
    set_run_font(r, "宋体", 14, False, "444444")

    add_picture(doc, ASSETS / "ai-generated-smart-barn.png", 14.5, "图 0-1 智能牛舍端边云协同场景示意图（Image Gen 生成）")

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, 130, 160, 130, 160)
    meta = [
        ("项目名称", "数牧增效智能牛舍监测平台"),
        ("报告类型", "系统设计与实现技术报告"),
        ("技术路线", "ESP32-S3 + MaixCam + Spring Boot + Vue3 + HarmonyOS/微信小程序"),
        ("应用场景", "肉牛养殖环境监测、健康风险识别、设备联动控制"),
        ("报告日期", "2026 年 6 月"),
    ]
    for i, (k, v) in enumerate(meta):
        shade_cell(table.rows[i].cells[0], LIGHT_GREEN)
        set_cell_text(table.rows[i].cells[0], k, True, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.rows[i].cells[1], v)
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    r = p.add_run("目录")
    set_run_font(r, "黑体", 16, True, DARK)
    entries = [
        "1 绪论",
        "2 需求分析",
        "3 系统总体设计",
        "4 现场硬件与边缘节点设计",
        "5 多模态 AI 算法设计",
        "6 平台软件与数据库设计",
        "7 Web 与移动端实现",
        "8 创新点与技术优势",
        "9 测试验证与部署计划",
        "10 总结",
    ]
    for e in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.left_indent = Cm(1.2)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(e)
        set_run_font(r, "宋体", 12)
    doc.add_page_break()


def build():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "1 绪论", 1)
    add_heading(doc, "1.1 项目背景", 2)
    add_paragraph(doc, "肉牛养殖正在从经验巡栏向数据化、智能化管理转型。传统牛舍主要依赖人工观察牛只精神状态、采食情况和环境变化，存在巡检频次有限、异常发现滞后、证据留存不足等问题。温湿度、氨气、烟雾、通风状态等环境因素会直接影响牛只呼吸道健康和应激水平；卧倒时间、活动量、咳嗽、喘鸣、眼鼻分泌物、皮肤异常等信号则需要与环境数据共同判断。")
    add_paragraph(doc, "数牧增效项目面向中小规模养殖场和普通农户，目标不是建设高成本集中式视频监控系统，而是以低成本边缘节点、轻量 AI 模型和多端应用平台构成可持续运行的主动监测系统。系统通过端侧采集、边缘初判、云端融合和服务端闭环处置，提高异常发现效率并降低人工巡检压力。")

    add_heading(doc, "1.2 建设目标", 2)
    add_bullets(doc, [
        ("1", "实现牛舍环境数据的连续采集与趋势展示，覆盖温度、湿度、氨气、光照、烟雾及执行器状态。"),
        ("2", "在 MaixCam 端完成牛只目标、姿态和局部异常的轻量化视觉识别，减少原始视频长期上传压力。"),
        ("3", "利用 I2S 麦克风采集牛舍声音，识别咳嗽、喘鸣、异常叫声等健康风险信号。"),
        ("4", "通过 Spring Boot 后端汇聚环境、声音、视觉和历史数据，形成风险等级、证据链和处置建议。"),
        ("5", "提供 Web 驾驶舱、鸿蒙 APP 和微信小程序，支撑现场巡检、远程控制、记录归档和模型迭代。"),
    ])

    add_heading(doc, "1.3 技术路线", 2)
    add_paragraph(doc, "系统采用“端—边—云—服”协同架构。端侧由 ESP32-S3 负责环境采集、执行器控制和 MQTT 通信；边缘侧由 MaixCam 承担视觉 AI 推理，由 I2S 声音节点完成有效声音筛选和特征提取；云端由 Spring Boot、JPA、H2/MySQL 组成数据与接口层；服务端由 Vue3 Web 看板、鸿蒙 APP 和微信小程序面向不同使用场景提供交互。")

    add_heading(doc, "2 需求分析", 1)
    add_heading(doc, "2.1 业务需求", 2)
    add_table(doc, ["需求类别", "具体需求", "项目实现方式"], [
        ("环境监测", "实时掌握温湿度、氨气、光照、烟雾等牛舍环境状态。", "ESP32-S3 环境控制节点采集数据，后端提供 latest/trend API，前端展示 48 小时趋势。"),
        ("健康风险识别", "发现疑似呼吸道异常、活动异常、局部外观异常。", "融合声音 AI、视觉 AI 和环境趋势，生成 HealthRisk 记录和复核建议。"),
        ("设备联动控制", "环境异常时可控制风扇、雾化器、声光报警和抓拍。", "Web/移动端写入控制命令，ESP32-S3 订阅或轮询执行。"),
        ("证据留存", "异常需要截图、音频、环境曲线和处理记录支撑复核。", "VisionEvent、AudioEvent、AlertRecord、ControlCommandRecord 形成事件链。"),
        ("多端使用", "管理人员在电脑端查看态势，巡检人员在手机端拍照识别。", "Vue Web 驾驶舱、鸿蒙 APP、微信小程序分别服务管理和巡检。"),
    ])

    add_heading(doc, "2.2 非功能需求", 2)
    add_table(doc, ["指标", "设计要求"], [
        ("低成本", "使用 ESP32-S3、MaixCam、MQ 系列传感器等原型级硬件先跑通试点，正式版预留工业传感器接口。"),
        ("实时性", "环境异常在本地先执行阈值判断和控制联动，云端用于汇聚、复核和展示。"),
        ("可靠性", "断网时保留关键事件摘要和截图，网络恢复后补传；执行器控制避免由 GPIO 直接带载。"),
        ("可维护性", "后端模块按环境、设备、事件、模型分层；前端页面按场景拆分，便于后续扩展。"),
        ("可解释性", "健康风险不只给出分数，还要给出咳嗽、鼻部异常、氨气偏高等证据。"),
    ])

    add_heading(doc, "3 系统总体设计", 1)
    add_heading(doc, "3.1 总体架构", 2)
    add_picture(doc, ASSETS / "architecture-diagram.png", 15.8, "图 3-1 数牧增效端边云服协同总体架构")
    add_paragraph(doc, "系统整体架构分为四层。端侧负责采集和执行，保证高温、高湿、氨气偏高、火情烟雾等异常能够在本地快速响应；边缘侧负责视觉和声音事件生成，将连续视频和原始音频转化为结构化事件；云端负责数据持久化、指标聚合、风险融合和接口服务；服务端负责将风险、趋势和控制能力提供给管理人员和巡检人员。")

    add_heading(doc, "3.2 数据闭环", 2)
    add_paragraph(doc, "系统闭环链路为：采集感知、边缘判断、事件上报、融合评分、多端告警、控制命令、现场执行、数据回流优化。与单纯展示型监控不同，本项目强调异常发现后的处置闭环。以氨气趋势偏高为例，ESP32-S3 可先开启风扇并上报告警，同时触发 MaixCam 抓拍形成视觉证据；后端将环境曲线、抓拍事件和声音事件汇总为复核建议。")

    add_heading(doc, "4 现场硬件与边缘节点设计", 1)
    add_heading(doc, "4.1 ESP32-S3 环境控制节点", 2)
    add_paragraph(doc, "ESP32-S3 作为环境监测与控制核心，替代传统 STM32 + ESP8266 组合。其优势在于 WiFi、MQTT、UART、I2S 和较强 MCU 资源集成在同一主控上，可减少外接通信模块、降低布线复杂度，并为后续 TinyML 或本地声音特征处理保留余量。")
    add_table(doc, ["模块", "连接/接口", "主要作用"], [
        ("温湿度传感器", "I2C", "采集牛舍温湿度，作为通风、雾化和健康风险评估依据。"),
        ("MQ-137 氨气模块", "ADC / 预留 RS485", "用于氨气趋势告警；正式部署可替换为工业 NH3 传感器。"),
        ("烟雾/火焰传感器", "ADC / GPIO", "用于火情或烟雾快速识别，并触发本地声光报警。"),
        ("继电器/MOSFET", "GPIO / PWM", "控制风扇、雾化片、报警器等低压负载，避免主控直接带载。"),
        ("OLED 显示", "I2C", "显示联网状态、传感器读数和告警信息，便于现场调试。"),
    ])

    add_heading(doc, "4.2 MaixCam 视觉 AI 节点", 2)
    add_paragraph(doc, "MaixCam 作为独立视觉节点运行轻量 YOLO 模型，完成牛体检测、姿态识别、局部异常候选识别和事件截图留证。该设计避免将原始视频长期上传到服务器，减少网络带宽占用，并能在牛舍现场更快生成结构化视觉事件。MaixCam 与 ESP32-S3 之间通过 UART JSON 进行事件同步和抓拍触发。")

    add_heading(doc, "4.3 声音采集节点", 2)
    add_paragraph(doc, "声音侧使用 INMP441 等 I2S 数字麦克风采集牛舍声音，通过分帧、能量检测、VAD、Log-Mel/MFCC 特征提取，识别咳嗽、喘鸣、异常叫声和背景噪声。声音事件不单独作为诊断结论，而是作为多模态健康风险判断的一类证据。")

    add_heading(doc, "5 多模态 AI 算法设计", 1)
    add_picture(doc, ASSETS / "multimodal-dataflow.png", 15.8, "图 5-1 多模态健康风险识别数据流")
    add_heading(doc, "5.1 视觉识别算法", 2)
    add_paragraph(doc, "视觉算法以 YOLOv8n、YOLOv5n 或 MaixCam 支持的轻量 YOLO 模型为基线，检测牛体、站立、卧倒、人员进入等目标和状态。算法输出目标框、类别、置信度、跟踪 ID 和截图路径；后续结合区域 ROI、时间状态机和目标跟踪结果，判断长时间卧倒、疑似跌倒、活动减少、局部皮肤或眼鼻异常。")

    add_heading(doc, "5.2 声音识别算法", 2)
    add_paragraph(doc, "声音算法先通过 VAD 过滤无效背景噪声，再提取 Log-Mel 或 MFCC 特征，使用轻量 CNN、YAMNet 迁移学习或后端分类模型识别咳嗽、喘鸣和异常叫声。声音事件记录包含类型、置信度、时长、SNR 和音频片段路径，便于后续人工复核和模型再训练。")

    add_heading(doc, "5.3 健康风险融合", 2)
    add_table(doc, ["输入模态", "特征示例", "风险贡献"], [
        ("环境数据", "温湿度、氨气、烟雾、通风状态", "判断牛舍应激和呼吸道风险背景。"),
        ("声音数据", "咳嗽、喘鸣、异常叫声、SNR", "发现呼吸道异常和群体异常声音。"),
        ("姿态行为", "卧倒时长、活动量、跛行、异常聚集", "判断活动减少、跌倒或应激行为。"),
        ("局部视觉", "眼鼻分泌物、皮肤纹理异常、局部损伤", "形成疾病筛查的图像证据。"),
        ("历史趋势", "同牛只、同栏位、同时间段变化", "降低单次误报，校准风险阈值。"),
    ])
    add_paragraph(doc, "融合模型在原型阶段可采用规则引擎加权评分，形成正常、关注、告警、需复核四级风险；在数据积累后可进一步引入 Logistic Regression、XGBoost 或轻量时序模型。系统输出应明确标注“用于人工复核，不替代兽医诊断”。")

    add_heading(doc, "6 平台软件与数据库设计", 1)
    add_heading(doc, "6.1 后端架构", 2)
    add_paragraph(doc, "后端采用 Spring Boot 3.3.5、Java 17、Spring Data JPA、Validation、WebSocket、H2/MySQL。当前代码已实现环境数据、设备列表、控制命令、视觉事件、声音事件、健康风险、告警记录和模型信息接口。演示环境默认使用 H2 内存数据库并自动写入样例数据，MySQL profile 可用于试点部署。")
    add_table(doc, ["实体/接口", "功能说明"], [
        ("EnvironmentReading", "保存温湿度、氨气、光照、烟雾、风扇和雾化状态。"),
        ("Device", "保存 ESP32-S3、MaixCam、声音节点和执行器模块状态。"),
        ("VisionEvent", "保存视觉异常事件、置信度、目标 ID、截图路径和证据描述。"),
        ("AudioEvent", "保存声音事件类型、置信度、时长、SNR 和音频路径。"),
        ("HealthRisk", "保存融合风险等级、风险分数、证据链和处置建议。"),
        ("ControlCommandRecord", "记录风扇、雾化、报警、抓拍等控制命令历史。"),
    ])

    add_heading(doc, "6.2 API 设计", 2)
    add_table(doc, ["接口类别", "代表接口", "说明"], [
        ("总览", "GET /api/dashboard", "返回最新环境、48 小时趋势、设备、事件、健康风险、告警和统计指标。"),
        ("环境", "GET /api/environment/latest；POST /api/environment", "支持环境数据查询和 ESP32-S3/MQTT 网关转写。"),
        ("设备", "GET /api/devices；POST /api/devices/command", "获取设备清单并写入控制命令队列。"),
        ("事件", "GET/POST /api/events/vision、audio、health-risks", "记录视觉、声音和融合风险事件。"),
        ("模型", "GET /api/models", "展示视觉、声音、多模态融合模型版本、任务和指标。"),
    ])

    add_heading(doc, "7 Web 与移动端实现", 1)
    add_heading(doc, "7.1 Web 平台", 2)
    add_picture(doc, ASSETS / "web-contact-sheet.png", 15.8, "图 7-1 Web 平台运行界面截图")
    add_paragraph(doc, "Web 前端采用 Vue3、Vite、Pinia、ECharts、Three.js 和 lucide 图标库，实现项目首页、总览驾驶舱、环境监测、健康风险、设备控制和 AI 模型页面。驾驶舱展示存栏数量、在线设备、未闭环告警、边缘推理 FPS、3D 牛舍态势、环境数据和风险队列；设备控制页支持风扇、雾化、声光报警、视觉抓拍等命令下发。")
    add_picture(doc, ASSETS / "web-dashboard.png", 15.2, "图 7-2 总览驾驶舱页面截图")
    add_picture(doc, ASSETS / "web-health.png", 15.2, "图 7-3 健康风险证据链页面截图")

    add_heading(doc, "7.2 鸿蒙 APP", 2)
    add_picture(doc, ASSETS / "hm-app-contact-sheet.png", 15.8, "图 7-4 鸿蒙 APP 核心界面截图")
    add_paragraph(doc, "鸿蒙端采用 ArkTS 实现移动巡检闭环，包含总览、AI 识牛、识别记录、后台管理和设备链路界面。巡检人员可上传或拍摄牛只照片，系统生成风险等级、特征判断和巡检建议；管理页提供风扇、抓拍、声光提醒、模型同步等操作入口；设备链路页展示 ESP32-S3、MaixCam、声音节点和继电器控制箱状态。当前移动端识别流程采用本地模拟分析验证交互闭环，后续可替换为真实模型服务或端侧推理结果。")

    add_heading(doc, "7.3 微信小程序", 2)
    add_paragraph(doc, "微信小程序与鸿蒙端功能相近，提供移动端总览、照片诊断、历史记录和管理入口。小程序使用本地缓存保存最近诊断记录，适合在项目演示阶段快速验证巡检流程，也便于后续接入统一后端账号体系和真实推理服务。")

    add_heading(doc, "8 创新点与技术优势", 1)
    add_table(doc, ["创新点", "技术说明", "应用价值"], [
        ("端边云服协同闭环", "端侧采集执行、边缘 AI 初判、云端融合评分、多端告警处置。", "从单纯监测升级为异常处置闭环。"),
        ("ESP32-S3 一体化主控", "集成 WiFi、MQTT、UART、I2S 和控制能力，替代 STM32 + 外接 WiFi 模块。", "降低硬件复杂度和部署成本。"),
        ("MaixCam 本地视觉 AI", "端侧运行轻量 YOLO，只上传结构化事件和截图证据。", "减少带宽压力，提高响应速度和隐私安全。"),
        ("环境触发抓拍机制", "氨气、高温、烟雾异常可触发摄像头抓拍和事件关联。", "形成环境异常与视觉证据联动。"),
        ("多模态健康风险证据链", "融合环境、声音、姿态、局部视觉和历史趋势。", "提高风险判断可解释性，降低单模态误报。"),
        ("面向农户低成本落地", "原型版使用低成本硬件，正式版预留工业传感器和 4G 网关接口。", "适合中小规模养殖场逐步部署。"),
    ])

    add_heading(doc, "9 测试验证与部署计划", 1)
    add_heading(doc, "9.1 当前原型验证", 2)
    add_paragraph(doc, "当前项目已完成 Spring Boot 后端、Vue Web 看板、鸿蒙 APP、微信小程序和数据库脚本。后端接口 /api/dashboard 已可返回设备、环境、事件、风险和告警数据；前端页面可通过本地服务访问并完成截图验证；移动端原型已覆盖巡检、诊断、记录和管理流程。")
    add_table(doc, ["验证项", "当前状态", "后续要求"], [
        ("后端服务", "本地 /api/dashboard 返回 200，演示数据可自动注入。", "接入 MQTT Broker 和真实设备上报。"),
        ("Web 看板", "项目首页、驾驶舱、环境、健康、设备、模型页面可访问。", "增加实时推送和告警闭环状态。"),
        ("移动端", "鸿蒙 APP 和微信小程序完成核心交互原型。", "接入真实模型接口、用户权限和后端同步。"),
        ("AI 模型", "报告和页面展示 YOLO、声音分类和融合模型方案。", "训练端侧模型并在 MaixCam/后端验证指标。"),
    ])

    add_heading(doc, "9.2 测试指标", 2)
    add_table(doc, ["测试类别", "建议指标"], [
        ("硬件通信", "ESP32-S3 与传感器稳定采样；MaixCam UART/MQTT 事件上报可靠；断网重连后可补传关键事件。"),
        ("环境控制", "氨气、高温、烟雾等阈值触发后，风扇、雾化和报警执行延迟满足现场要求。"),
        ("视觉 AI", "牛体、站立、卧倒、疑似跌倒等类别在本地数据集上达到可演示准确率，并记录 FPS。"),
        ("声音 AI", "咳嗽、喘鸣、背景噪声分类在牛舍噪声条件下具备可用召回率。"),
        ("融合风险", "输出风险等级、证据链和建议，人工复核后可回流样本优化阈值。"),
        ("多端体验", "Web、鸿蒙、微信端页面无明显遮挡，命令记录和诊断记录可追溯。"),
    ])

    add_heading(doc, "9.3 阶段计划", 2)
    add_table(doc, ["阶段", "任务", "交付物"], [
        ("第 1 阶段", "ESP32-S3 采集、控制、MQTT 上报和后端环境接口联调。", "环境节点样机、传感器数据、控制命令闭环。"),
        ("第 2 阶段", "MaixCam 跑通轻量 YOLO 模型，完成 UART/MQTT 视觉事件上报。", "视觉节点样机、事件截图、推理指标。"),
        ("第 3 阶段", "采集本地牛舍图像和声音样本，进行标注、训练和误报分析。", "数据集、模型版本、评估报告。"),
        ("第 4 阶段", "Web/移动端接入真实数据，完成告警、抓拍、控制和记录归档。", "试点系统原型、多端演示材料。"),
        ("第 5 阶段", "现场试运行，记录环境、疾病风险、设备故障和人工复核结果。", "试点总结、模型优化计划、正式版 BOM。"),
    ])

    add_heading(doc, "10 总结", 1)
    add_paragraph(doc, "数牧增效项目围绕肉牛养殖场景构建了从现场感知、边缘 AI、云端融合到多端服务的完整技术路线。系统通过 ESP32-S3 降低采集与控制复杂度，通过 MaixCam 将视觉 AI 前移到牛舍现场，通过声音、视觉、环境和历史趋势融合形成可解释的健康风险证据链，并通过 Web、鸿蒙 APP 和微信小程序支撑管理与巡检。")
    add_paragraph(doc, "从当前实现看，项目已经具备完整软件原型和清晰硬件落地方案。后续重点应放在真实 ESP32-S3/MaixCam 数据接入、MQTT 稳定通信、端侧模型训练部署、现场样本回流和风险阈值校准上。完成这些环节后，系统可从演示原型推进为面向中小养殖场的低成本智能牛舍监测试点方案。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
