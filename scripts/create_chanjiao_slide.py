# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "产教融合-实践育人.pptx"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def inches(value: float):
    return Inches(value)


prs = Presentation()
prs.slide_width = inches(13.333)
prs.slide_height = inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

C = {
    "ink": "17231C",
    "forest": "1F6B4D",
    "moss": "8FB95E",
    "leaf": "DDECCB",
    "cream": "F5F7F2",
    "white": "FFFFFF",
    "gold": "D9A441",
    "teal": "157C89",
    "blue": "355C7D",
    "rust": "BC5A3C",
    "muted": "5E6B63",
    "line": "D7DFD2",
    "soft": "E6ECE1",
}


def fill(shape, color: str, transparency: float = 0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def line(shape, color: str, width: float = 0.75, transparency: float = 0):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def no_line(shape):
    shape.line.fill.background()


def add_rect(x, y, w, h, color, outline=None, radius=False, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, inches(x), inches(y), inches(w), inches(h))
    fill(shp, color, transparency)
    if outline:
        line(shp, outline, 0.8)
    else:
        no_line(shp)
    return shp


def add_oval(x, y, w, h, color, outline=None, lw=0.75):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, inches(x), inches(y), inches(w), inches(h))
    fill(shp, color)
    if outline:
        line(shp, outline, lw)
    else:
        no_line(shp)
    return shp


def add_text(text, x, y, w, h, size=12, color=None, bold=False, align="left", valign="top"):
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }[valign]
    p = tf.paragraphs[0]
    p.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }[align]
    run = p.add_run()
    run.text = text
    font = run.font
    font.name = "Microsoft YaHei"
    font.size = Pt(size)
    font.bold = bold
    if color:
        font.color.rgb = rgb(color)
    return box


def add_line(x1, y1, x2, y2, color=C["line"], width=1.2):
    shp = slide.shapes.add_connector(
        1, inches(x1), inches(y1), inches(x2), inches(y2)
    )
    line(shp, color, width)
    return shp


def arrow_h(x1, y, x2, color=C["line"], width=1.2):
    add_line(x1, y, x2 - 0.12, y, color, width)
    tri = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE, inches(x2 - 0.12), inches(y - 0.06), inches(0.14), inches(0.12)
    )
    tri.rotation = 90
    fill(tri, color)
    no_line(tri)


def arrow_v(x, y1, y2, color=C["line"], width=1.2):
    add_line(x, y1, x, y2 - 0.12, color, width)
    tri = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE, inches(x - 0.06), inches(y2 - 0.12), inches(0.12), inches(0.14)
    )
    tri.rotation = 180
    fill(tri, color)
    no_line(tri)


def add_pill(text, x, y, w, txt_color, bg):
    add_rect(x, y, w, 0.34, bg, radius=True)
    add_text(text, x, y + 0.08, w, 0.12, 8.4, txt_color, True, "center", "middle")


def add_mini_barn(x, y):
    roof = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE, inches(x + 0.11), inches(y), inches(0.5), inches(0.28)
    )
    fill(roof, C["white"])
    no_line(roof)
    add_rect(x + 0.17, y + 0.22, 0.38, 0.29, C["white"])
    add_rect(x + 0.28, y + 0.35, 0.16, 0.16, C["forest"])


def add_data_icon(x, y, color):
    points = [(0.12, 0.40), (0.25, 0.24), (0.42, 0.32), (0.62, 0.12)]
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        add_line(x + x1, y + y1, x + x2, y + y2, C["white"], 1.3)
    for dx, dy in points:
        add_oval(x + dx - 0.04, y + dy - 0.04, 0.08, 0.08, color, C["white"], 1.1)
    add_line(x + 0.1, y + 0.49, x + 0.68, y + 0.49, C["white"], 1.2)


def add_users_icon(x, y):
    add_oval(x + 0.28, y + 0.07, 0.2, 0.2, C["white"])
    add_oval(x + 0.08, y + 0.16, 0.16, 0.16, C["white"])
    add_oval(x + 0.56, y + 0.16, 0.16, 0.16, C["white"])
    add_rect(x + 0.22, y + 0.35, 0.34, 0.16, C["white"], radius=True)
    add_rect(x + 0.02, y + 0.40, 0.26, 0.12, C["white"], radius=True)
    add_rect(x + 0.50, y + 0.40, 0.26, 0.12, C["white"], radius=True)


def add_docs_icon(x, y):
    add_rect(x + 0.21, y + 0.11, 0.34, 0.46, C["white"])
    for dy, ww in [(0.24, 0.2), (0.34, 0.22), (0.44, 0.14)]:
        add_line(x + 0.28, y + dy, x + 0.28 + ww, y + dy, C["rust"], 1)


def add_card(x, y, w, h, tag, title, desc, color, tint, icon):
    add_rect(x + 0.06, y + 0.06, w, h, "CBD7C6", radius=True, transparency=0.45)
    add_rect(x, y, w, h, C["white"], C["soft"], radius=True)
    add_rect(x, y, 0.12, h, color)
    add_oval(x + 0.34, y + 0.36, 0.72, 0.72, color)
    if icon == "barn":
        add_mini_barn(x + 0.38, y + 0.45)
    elif icon == "data":
        add_data_icon(x + 0.31, y + 0.43, color)
    elif icon == "users":
        add_users_icon(x + 0.31, y + 0.42)
    elif icon == "docs":
        add_docs_icon(x + 0.31, y + 0.42)
    add_rect(x + 1.22, y + 0.36, 0.92, 0.3, tint, radius=True)
    add_text(tag, x + 1.33, y + 0.44, 0.7, 0.12, 7.6, color, True, "center", "middle")
    add_text(title, x + 1.22, y + 0.78, w - 1.55, 0.34, 18, C["ink"], True)
    add_text(desc, x + 1.22, y + 1.23, w - 1.55, 0.56, 12.4, C["muted"])


# Background and title band
add_rect(0, 0, 13.333, 7.5, C["cream"])
add_rect(10.95, -0.38, 1.9, 1.9, "DDECCB", radius=True, transparency=0.15)
add_rect(11.55, 0.23, 1.18, 1.18, "F0C667", radius=True, transparency=0.12)
add_rect(0, 0, 13.333, 1.68, C["ink"])
add_rect(0, 1.55, 13.333, 0.13, C["moss"])

add_text("产教融合 · 实践育人", 0.72, 0.40, 6.4, 0.55, 30, C["white"], True)
add_text(
    "以真实牛舍现场为课堂，把验证、模型、系统和成果转化贯通起来",
    0.75,
    1.12,
    8.2,
    0.22,
    10.6,
    C["leaf"],
)
add_pill("协同验证", 9.08, 0.48, 0.92, C["ink"], C["leaf"])
add_pill("课程实践", 10.18, 0.48, 0.92, C["ink"], "F0C667")
add_pill("迭代共创", 11.28, 0.48, 0.92, C["ink"], "BBD7E5")

# Left narrative rail
add_text("四方联动闭环", 0.82, 2.06, 2.4, 0.34, 17.5, C["ink"], True)
add_text(
    "产业现场提出问题，教学科研给出方法；\n师生共创完成迭代，最终沉淀为\n可展示、可测试、可转化的成果材料。",
    0.84,
    2.55,
    2.84,
    0.98,
    10.8,
    C["muted"],
)

flow_x, flow_y = 0.94, 3.72
arrow_v(flow_x + 0.26, flow_y + 0.16, flow_y + 2.5, "C8D6C3", 1.25)
steps = [
    ("现场问题", C["rust"]),
    ("数据模型", C["teal"]),
    ("平台验证", C["forest"]),
    ("成果转化", C["gold"]),
]
for i, (label, color) in enumerate(steps):
    yy = flow_y + i * 0.63
    add_oval(flow_x, yy, 0.52, 0.52, color, C["white"], 1.4)
    add_text(str(i + 1), flow_x + 0.18, yy + 0.15, 0.16, 0.12, 9.6, C["white"], True, "center", "middle")
    add_text(label, flow_x + 0.72, yy + 0.15, 1.15, 0.16, 10.5, C["ink"], True)

# Main cards and relationship arrows
card_w, card_h = 4.28, 1.94
left_x, right_x = 4.05, 8.68
top_y, bottom_y = 2.0, 4.32
add_card(
    left_x,
    top_y,
    card_w,
    card_h,
    "产业端",
    "产教融合",
    "与农服机构、畜牧设备渠道和示范基地协同验证。",
    C["forest"],
    C["leaf"],
    "barn",
)
add_card(
    right_x,
    top_y,
    card_w,
    card_h,
    "教学端",
    "科教融汇",
    "将真实牛舍数据、算法模型和大屏系统融入课程实践。",
    C["teal"],
    "D9EEF0",
    "data",
)
add_card(
    left_x,
    bottom_y,
    card_w,
    card_h,
    "共创端",
    "师生共创",
    "围绕现场问题持续迭代硬件、算法和平台。",
    C["blue"],
    "DCE7F0",
    "users",
)
add_card(
    right_x,
    bottom_y,
    card_w,
    card_h,
    "成果端",
    "产学研联动",
    "形成样机、平台、测试报告、案例材料和商业计划。",
    C["rust"],
    "F3DED7",
    "docs",
)

arrow_h(left_x + card_w + 0.12, top_y + card_h / 2, right_x - 0.12, "B9CDB1", 1.25)
arrow_h(left_x + card_w + 0.12, bottom_y + card_h / 2, right_x - 0.12, "B9CDB1", 1.25)
arrow_v(left_x + card_w / 2, top_y + card_h + 0.08, bottom_y - 0.08, "B9CDB1", 1.25)
arrow_v(right_x + card_w / 2, top_y + card_h + 0.08, bottom_y - 0.08, "B9CDB1", 1.25)

# Footer output line
add_rect(0.78, 6.92, 11.76, 0.01, "D4DEC8")
add_text("输出物：样机 · 平台 · 测试报告 · 案例材料 · 商业计划", 0.84, 7.03, 6.6, 0.2, 9.8, C["muted"])
add_text("数牧增效智能牛舍监测平台", 10.05, 7.03, 2.5, 0.2, 9.3, C["muted"], False, "right")

prs.core_properties.title = "产教融合 实践育人"
prs.core_properties.subject = "产教融合、科教融汇、师生共创、产学研联动"
prs.core_properties.author = "Codex"
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
