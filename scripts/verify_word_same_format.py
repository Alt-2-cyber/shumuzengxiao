from pathlib import Path

from docx import Document


BASE = Path("D:/table")
OUTPUT_PATH = BASE / "\u7269\u6599\u7533\u8bf7\u8868(3)-\u6570\u7267\u589e\u6548\u5b8c\u6210\u7248.docx"


EXPECTED_ITEMS = [
    "ESP32-S3 / N16R8 \u5f00\u53d1\u677f",
    "Sipeed MaixCam \u8fb9\u7f18 AI \u6444\u50cf\u5934\u5f00\u53d1\u677f",
    "\u706d\u706b\u98ce\u6247\u6a21\u5757 \u964d\u6e29\u98ce\u673a",
    "SHT31\u6e29\u6e7f\u5ea6\u4f20\u611f\u5668\u6a21\u5757",
    "MQ-135 \u7a7a\u6c14\u8d28\u91cf\u4f20\u611f\u5668\u6a21\u5757",
    "MQ137\u6c28\u6c14\u6a21\u5757",
    "INMP441\u9ea6\u514b\u98ce\u6a21\u5757",
    "BH1750 \u5149\u7167\u4f20\u611f\u5668\u6a21\u5757",
    "1.3\u5bf8OLED4\u9488\uff08\u84dd\u8272\u663e\u793a\uff09",
    "\u65b0\u6b3e ADXL345 \u52a0\u901f\u5ea6\u4f20\u611f\u5668",
    "\u706b\u7130\u4f20\u611f\u5668\u6a21\u5757",
    "GY-MAX30102",
    "\u5b89\u670d\u4f18B-1\u7ea2\u5916\u6d4b\u6e29\u4f20\u611f\u5668\u6a21\u5757",
    "\u805a\u5408\u7269\u9502\u7535\u6c60",
    "\u4e07\u7528\u8868",
    "MOSFET \u5927\u529f\u7387\u5f00\u5173\u9a71\u52a8\u6a21\u5757",
    "DC-DC \u964d\u538b\u6a21\u5757 12V \u8f6c 5V",
    "\u56db\u8def\u7ee7\u7535\u5668\u6a21\u5757\uff0c\u5e26\u5149\u8026\u9694\u79bb",
    "830\u5b54\u9762\u5305\u677f \u5e26\u80cc\u80f6\u53ef\u62fc\u63a5",
    "\u675c\u90a6\u7ebf21CM \u516c\u5bf9\u516c 2.54mm\uff081\u639240P\uff09",
    "\u675c\u90a6\u7ebf21CM \u6bcd\u5bf9\u6bcd 2.54mm\uff081\u639240P\uff09",
    "\u675c\u90a6\u7ebf21CM \u516c\u5bf9\u6bcd 2.54mm\uff081\u639240P\uff09",
    "\u7535\u70d9\u94c1BY-55W",
    "60w\u5200\u5934",
]

EXPECTED_COMPETITIONS = [
    "\u5929\u6d25\u5e02\u5927\u5b66\u751f\u521b\u65b0\u8bad\u7ec3\u8ba1\u5212\uff08\u5927\u521b\uff09",
    "\u4e2d\u56fd\u56fd\u9645\u5927\u5b66\u751f\u521b\u65b0\u5927\u8d5b\uff08\u4e92\u8054\u7f51+\uff09",
    '"\u6311\u6218\u676f"\u5927\u5b66\u751f\u8bfe\u5916\u5b66\u672f\u79d1\u6280\u4f5c\u54c1\u7ade\u8d5b',
    "\u4ec1\u7231\u676f\u6821\u7ea7\u521b\u65b0\u521b\u4e1a\u5927\u8d5b",
    "\u5f00\u653e\u539f\u5b50\u5f00\u6e90\u5927\u8d5b",
    "\u5353\u8d8a\u676f\u6821\u7ea7\u79d1\u6280\u7ade\u8d5b",
    "\u4e2d\u56fd\u9ad8\u6821\u8ba1\u7b97\u673a\u5927\u8d5b--\u5d4c\u5165\u5f0f\u8bbe\u8ba1\u8d5b",
    "\u5168\u56fd\u5927\u5b66\u751f\u7535\u5b50\u8bbe\u8ba1\u7ade\u8d5b",
    "\u5168\u56fd\u5927\u5b66\u751f\u7269\u8054\u7f51\u8bbe\u8ba1\u7ade\u8d5b",
    "\u534e\u5317\u4e94\u7701\uff08\u5e02\u3001\u81ea\u6cbb\u533a\uff09\u5927\u5b66\u751f\u8ba1\u7b97\u673a\u5e94\u7528\u5927\u8d5b\u5929\u6d25\u8d5b\u533a\u7ade\u8d5b",
]


def main():
    doc = Document(OUTPUT_PATH)
    paragraphs = doc.paragraphs
    text = "\n".join(p.text for p in paragraphs)
    missing_items = [name for name in EXPECTED_ITEMS if name not in text]
    if missing_items:
        raise RuntimeError("Missing material paragraphs: " + " / ".join(missing_items))

    item_positions = [text.index(name) for name in EXPECTED_ITEMS]
    if item_positions != sorted(item_positions):
        raise RuntimeError("Material item order does not match the Excel order.")

    if len(doc.tables) != 1:
        raise RuntimeError(f"Expected 1 competition table, found {len(doc.tables)}")
    competitions = [row.cells[0].text.strip() for row in doc.tables[0].rows]
    if competitions != EXPECTED_COMPETITIONS:
        raise RuntimeError("Competition table changed unexpectedly.")

    first_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "\u91c7\u8d2d\u786c\u4ef6\u53ca\u4f7f\u7528\u65b9\u5f0f")
    second_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "\u9884\u8ba1\u53c2\u4e0e\u6bd4\u8d5b")
    material_paras = paragraphs[first_heading + 1 : second_heading]
    bold_leads = [bool(p.runs and p.runs[0].bold) for p in material_paras]
    if len(material_paras) != 24 or not all(bold_leads):
        raise RuntimeError("Material paragraph count/bold lead formatting check failed.")

    print("DOCX structural verification passed")
    print(f"file={OUTPUT_PATH}")
    print(f"paragraphs={len(paragraphs)}")
    print(f"materials={len(material_paras)}")
    print(f"competition_rows={len(competitions)}")
    print("first_material=" + material_paras[0].text)
    print("last_material=" + material_paras[-1].text)


if __name__ == "__main__":
    main()
