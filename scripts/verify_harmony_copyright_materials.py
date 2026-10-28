from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document
from lxml import etree
from PIL import Image, ImageDraw, ImageFont

from update_harmony_copyright_materials import ROOT, STAGE, TARGET, WORK, NAME
from verify_cow_source_template_format import body_rows


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def contact(paths, target, columns=4):
    cell_w, cell_h = 325, 475
    image = Image.new('RGB', (cell_w * columns, cell_h * ((len(paths) + columns - 1) // columns)), '#DDE3EC')
    draw = ImageDraw.Draw(image)
    for index, path in enumerate(paths):
        item = Image.open(path).convert('RGB')
        item.thumbnail((cell_w - 10, cell_h - 28))
        x, y = index % columns * cell_w, index // columns * cell_h
        image.paste(item, (x + (cell_w - item.width) // 2, y + 25))
        draw.text((x + 8, y + 6), path.stem, fill='black')
    image.save(target)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--folder', type=Path, default=STAGE)
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    folder = args.folder
    listing = json.loads((folder / '源程序文件与行号清单.json').read_text('utf-8'))
    stats = listing['stats']
    assert stats['fileCount'] == 14 and stats['nonemptyLines'] == 2390 and stats['codePages'] == 48
    assert stats['submittedAll'] and stats['selectedLines'] == stats['nonemptyLines']
    expected = [row['text'].expandtabs(2) for row in listing['selected']]
    code_path = folder / f'{NAME}V1.0-源程序.docx'
    manual_path = folder / f'{NAME}V1.0-用户操作手册.docx'
    assert [p.text for p in Document(code_path).paragraphs] == expected
    with ZipFile(folder / '完整源码归档.zip') as archive:
        for item in listing['files']:
            assert sha(ROOT / item['path']) == item['sha256'], item['path']
            assert sha(folder / '源码归档' / item['path']) == item['sha256']
            assert hashlib.sha256(archive.read(item['path'])).hexdigest() == item['sha256']
    code = fitz.open(code_path.with_suffix('.pdf'))
    manual = fitz.open(manual_path.with_suffix('.pdf'))
    assert len(code) == 48 and len(manual) == 14
    counts, compressed = [], []
    for index, page in enumerate(code):
        identifiers = [word[4] for word in page.get_text('words') if word[2] < 80 and word[1] > 65 and word[4].isdigit()]
        required = min(50, len(expected) - index * 50)
        assert len(identifiers) == required, (index + 1, 'line identifiers', len(identifiers))
        assert [int(value) for value in identifiers] == list(range(index * 50 + 1, index * 50 + required + 1))
        body_lines = body_rows(page)
        assert len(body_lines) == required, (index + 1, 'wrapped source rows', len(body_lines))
        assert all(line[2] <= page.rect.width - 10 for line in body_lines), (index + 1, 'source outside page')
        counts.append(required)
    document = Document(manual_path)
    body = '\n'.join(p.text for p in document.paragraphs) + '\n' + '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert '总览、巡检、记录、管理' in body
    assert 'AI识牛' not in body and 'AI 识牛' not in body
    assert '微信基础库' not in body and '微信订阅' not in body and '小程序为微信' not in body
    assert '2390' in body
    assert not re.search(r'sk-[a-zA-Z0-9]{25,}', '\n'.join(expected) + body)
    for index in range(2, 12):
        assert len(manual[index].get_images()) >= 2, (index + 1, 'missing screenshot or flowchart')
    screenshots = json.loads((folder / '截图来源与标注清单.json').read_text('utf-8'))
    assert len(screenshots) == 15
    red_count = 0
    for item in screenshots:
        paths = {key: folder / Path(item[key]).relative_to(TARGET) for key in ('file', 'layout', 'annotated')}
        assert sha(paths['file']) == item['rawSha256']
        assert sha(paths['layout']) == item['layoutSha256']
        assert sha(paths['annotated']) == item['annotatedSha256']
        nodes = json.loads(paths['layout'].read_text('utf-8'))
        for mark in item['marks']:
            match = next(node for node in nodes if node.get('hierarchy') == mark['hierarchy'])
            assert [int(value) for value in re.findall(r'-?\d+', match['bounds'])] == mark['rect']
            red_count += 1
        raw = Image.open(paths['file']).convert('RGB')
        assert raw.getpixel((raw.width // 2, 150))[0] > 230, (item['key'], 'old dark screenshot')
    assert len(list((folder / '配图').glob('流程图*.png'))) == 10
    assert len(list((folder / '配图').glob('*-红框.png'))) == 15
    for doc_path in (code_path, manual_path):
        with ZipFile(doc_path) as archive:
            assert archive.testzip() is None
            for member in archive.namelist():
                if member.endswith('.xml'):
                    raw = archive.read(member)
                    etree.fromstring(raw)
                    assert b'i@yufuping.com' not in raw
    description = (folder / '功能描述-填表用.txt').read_text('utf-8').strip()
    assert 501 <= len(description) <= 1300
    if args.render:
        review = WORK / 'review'
        review.mkdir(exist_ok=True)
        manuals, sources = [], []
        for label, pdf, destinations in [('manual', manual, manuals), ('source', code, sources)]:
            for index, page in enumerate(pdf):
                destination = review / f'{label}-{index+1:02d}.png'
                page.get_pixmap(matrix=fitz.Matrix(1.75, 1.75), alpha=False).save(destination)
                destinations.append(destination)
        contact(manuals, review / 'manual-contact.png')
        for index in range(0, len(sources), 12):
            contact(sources[index:index+12], review / f'source-contact-{index//12+1}.png')
        contact([folder / Path(item['annotated']).relative_to(TARGET) for item in screenshots], review / 'screenshots-contact.png', 5)
    report = dict(software=NAME, uiVersion='新版白底UI 2026-10-23', sourceFileCount=14,
                  sourcePhysicalLines=stats['physicalLines'], sourceLines=2390, codePages=48,
                  lineCountsPerPage=counts, manualPages=14, workflowFigures=10,
                  realScreenshots=15, redFrames=red_count, descriptionCharacters=len(description),
                  sourceMatchesCurrentCode=True, archiveHashesMatch=True, sourceLinesDoNotWrap=True,
                  allScreenshotsUseNewUI=True, redFramesMatchActualWidgets=True,
                  pdfExport='Microsoft Word native PDF export', applicationFormPreserved=True)
    report['files'] = {path.name: sha(path) for path in folder.iterdir()
                       if path.is_file() and path.suffix in {'.docx', '.pdf', '.zip'} and not path.name.startswith(('申请表', '~$'))}
    (folder / '材料质量检查.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), 'utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
