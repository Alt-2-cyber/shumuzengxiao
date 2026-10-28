from __future__ import annotations

import hashlib
import json
import math
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

import build_cow_copyright_package as base
from cow_copyright_content import APPS
from cow_manual_walkthrough_content import mobile_flows, web_flows

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '.run' / 'copyright-20261023'
SCREENS = WORK / 'screens'
DATE = '2026-10-23'
FONT = r'C:\Windows\Fonts\msyh.ttc'
BOLD = r'C:\Windows\Fonts\msyhbd.ttc'


def manifest():
    entries = {}
    for name in ['web', 'native', 'mini']:
        for entry in json.loads((SCREENS / f'{name}-manifest.json').read_text('utf-8')):
            entries[entry['key']] = entry
    # Match the red frames to the interactive widget rather than its nearby label.
    for key, hint, widget_type, number in [('hm-consultation-top', '问题标题 *', 'TextInput', '2'),
                                           ('hm-assistant', '发送问题', 'Button', '2')]:
        nodes = json.loads((SCREENS / f'{key}-layout.json').read_text('utf-8'))
        candidates = [node for node in nodes if node.get('type') == widget_type]
        if key == 'hm-consultation-top':
            node = next(node for node in candidates if node.get('hint') == '例如：牛只连续两天行走跛行')
        else:
            node = next(node for node in candidates if node['bounds'].startswith('[1057,2389]'))
        import re
        bounds = [int(value) for value in re.findall(r'-?\d+', node['bounds'])]
        entries[key]['marks'] = [mark for mark in entries[key]['marks'] if mark['label'] != number]
        entries[key]['marks'].append(dict(rect=bounds, label=number, text=hint))
    nodes = json.loads((SCREENS / 'hm-subscriptions-layout.json').read_text('utf-8'))
    controls = [node for node in nodes if node.get('type') in {'Toggle', 'Checkbox'}]
    entries['hm-subscriptions']['marks'] = [dict(rect=[int(value) for value in re.findall(r'-?\d+', node['bounds'])],
        label=str(index + 1), text=['新闻提醒开关', '每日知识开关', '关注方向'][index]) for index, node in enumerate(controls[:3])]
    entries['wx-subscriptions']['marks'] = [dict(rect=[352, 108, 413, 148], label='1', text='新闻提醒开关'),
                                            dict(rect=[352, 189, 413, 230], label='2', text='每日知识开关'),
                                            dict(rect=[378, 304, 410, 337], label='3', text='关注方向')]
    entries['hm-records']['marks'] = [mark for mark in entries['hm-records']['marks'] if mark['label'] == '1']
    entries['hm-knowledge-detail']['marks'] = [mark for mark in entries['hm-knowledge-detail']['marks'] if mark['label'] != '2']
    return entries


def annotate(entry, folder):
    image = Image.open(entry['file']).convert('RGB')
    width, height = image.size
    crop_top = 0
    if entry['key'].startswith(('hm-', 'wx-')):
        marks = entry['marks']
        first = min(mark['rect'][1] for mark in marks)
        last = max(mark['rect'][3] for mark in marks)
        wanted = max(1500 if width > 600 else 550, last - first + (220 if width > 600 else 80))
        if wanted < height:
            crop_top = max(0, min(height - wanted, first - (150 if width > 600 else 50)))
            image = image.crop((0, int(crop_top), width, int(crop_top + wanted)))
            height = image.height
    draw = ImageDraw.Draw(image)
    size = 30 if width < 600 else 75 if width < 1300 else 36
    font = ImageFont.truetype(BOLD, size)
    border = max(4, round(width / 155))
    for mark in entry['marks']:
        x1, y1, x2, y2 = mark['rect']
        y1, y2 = y1 - crop_top, y2 - crop_top
        x1, y1 = max(3, x1), max(3, y1)
        x2, y2 = min(width - 4, x2), min(height - 4, y2)
        assert x2 > x1 and y2 > y1, (entry['key'], mark)
        draw.rectangle((x1, y1, x2, y2), outline='#E0162C', width=border)
        diameter = size + 12
        bx = min(width - diameter - 2, x1)
        by = max(2, y1 - diameter - 4)
        draw.rectangle((bx, by, bx + diameter, by + diameter), fill='#E0162C')
        box = draw.textbbox((0, 0), mark['label'], font=font)
        draw.text((bx + (diameter - box[2]) / 2, by + 3), mark['label'], font=font, fill='white')
    destination = folder / f"{entry['key']}-红框.png"
    image.save(destination)
    return destination


def flowchart(flow, number, folder):
    image = Image.new('RGB', (1600, 205), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, 29)
    for index, node in enumerate(flow['nodes']):
        x = 22 + index * 402
        draw.rectangle((x, 42, x + 330, 150), outline='#333333', width=3)
        rows = [node[i:i + 10] for i in range(0, len(node), 10)]
        for row, text in enumerate(rows):
            box = draw.textbbox((0, 0), text, font=font)
            draw.text((x + (330 - box[2]) / 2, 64 + row * 37), text, font=font, fill='#111111')
        if index < 3:
            draw.line((x + 342, 96, x + 385, 96), fill='#333333', width=4)
            draw.polygon([(x + 385, 96), (x + 371, 88), (x + 371, 104)], fill='#333333')
    destination = folder / f"流程图{number:02d}-{flow['title']}.png"
    image.save(destination)
    return destination


def paragraph(document, text, size=10.5, after=5, bold=False, center=False):
    p = base.ptext(document, text, size=size, after=after, bold=bold,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER if center else None)
    p.paragraph_format.line_spacing = Pt(16)
    return p


def pictures(document, app, flow, entries, folder, number):
    screenshots = flow['screenshots']
    t = document.add_table(rows=1, cols=len(screenshots))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    total_width = 16.8
    for cell, (key, caption) in zip(t.rows[0].cells, screenshots):
        cell.width = Cm(total_width / len(screenshots))
        for tc_margin in ['top', 'bottom', 'left', 'right']:
            tcpr = cell._tc.get_or_add_tcPr()
            margins = tcpr.find(qn('w:tcMar'))
            if margins is None:
                margins = OxmlElement('w:tcMar'); tcpr.append(margins)
            element = OxmlElement('w:' + tc_margin)
            element.set(qn('w:w'), '30'); element.set(qn('w:type'), 'dxa'); margins.append(element)
        entry = entries[key]
        assert entry['marks'], f'Missing verified red boxes: {key}'
        png = annotate(entry, folder)
        raw = Path(entry['file'])
        raw_dest = folder / '原始截图' / raw.name
        raw_dest.parent.mkdir(exist_ok=True)
        shutil.copy2(raw, raw_dest)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1
        width, height = Image.open(png).size
        if app['key'] == 'web':
            w = total_width - 0.2; h = w * height / width
        else:
            h = min(11.5, (8.0 if len(screenshots) > 1 else 11.0) * height / width)
            w = h * width / height
        p.add_run().add_picture(str(png), width=Cm(w), height=Cm(h))
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = Pt(13)
        p.paragraph_format.space_after = Pt(2)
        suffix = '（界面局部）' if app['key'] != 'web' else ''
        base.font(p.add_run(f'图{number}-{screenshots.index((key, caption))+1} {caption}{suffix}'), 9)
        rowprops = t.rows[0]._tr.get_or_add_trPr()
        if rowprops.find(qn('w:cantSplit')) is None:
            rowprops.append(OxmlElement('w:cantSplit'))
    return t


def manual(app, folder, entries):
    document = base.new_from_template(base.MANUAL_TEMPLATE, app['name'])
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(1.905)
    base.header_footer(document, app['name'], '用户操作手册')
    figures = folder / '配图'
    figures.mkdir(exist_ok=True)
    p = paragraph(document, app['name'], 21, 20, True, True)
    if 'Title' not in document.styles:
        style = document.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles['Normal']
    p.style = document.styles['Title']
    p.paragraph_format.space_before = Pt(110)
    paragraph(document, 'V1.0', 16, 35, True, True)
    paragraph(document, '用户操作手册', 23, 40, True, True)
    paragraph(document, f"客户端类型：{app['short']}", 12, 16, False, True)
    paragraph(document, '依据当前源码与实际运行界面编制', 11, 12, False, True)
    paragraph(document, f'材料整理日期：{DATE}', 11, 12, False, True)
    paragraph(document, '正文使用真实运行截图，红框内编号与相邻操作步骤对应。', 10.5, 10, False, True)

    base.heading(document, '软件概述与使用准备', True)
    paragraph(document, app['summary'], after=10)
    base.table(document, [('软件全称', app['name'] + ' V1.0'), ('运行平台', app['platform']),
                           ('支撑环境', app['runtime']), ('页面入口', app.get('navigation', 'Web通过侧边导航进入业务页面。' if app['key'] == 'web' else '总览、AI识牛、记录、管理；首页进入AI助手、专家咨询、养牛资讯。')),
                           ('开始前检查', app.get('preflight', '确认平台地址可访问，后端与数据库运行；硬件数据以最近上报时间和现场情况为准。')),
                           ('截图阅读', '红框编号标明操作区域；红框若圈出结果字段表示核对该字段。界面中的UI_TEST、DEMO和手册演示条目不是真实病例。')])
    paragraph(document, '操作章节', bold=True, after=5)
    for index in range(0, 10, 2):
        paragraph(document, '；'.join(f"{j+1} {app['workflows'][j]['title']}" for j in range(index, min(index+2,10))), size=9.5, after=4)
    paragraph(document, '版本范围：' + app['status'], size=9.5, after=3)

    for index, flow in enumerate(app['workflows'], 1):
        base.heading(document, f"{index} {flow['title']}", True)
        paragraph(document, flow['purpose'], size=10, after=4)
        chart = flowchart(flow, index, figures)
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1
        p.add_run().add_picture(str(chart), width=Cm(16.6))
        pictures(document, app, flow, entries, figures, index)
        for step, text in enumerate(flow['steps'], 1):
            paragraph(document, f'（{step}）{text}', size=10, after=3)
        paragraph(document, '操作结果：' + flow['result'], size=9.5, after=3)
        paragraph(document, '异常处理：' + flow['exception'], size=9.5, after=3)
        paragraph(document, '源码对应：' + flow['source'], size=8.5, after=0)

    base.heading(document, '11 常见问题与处理', True)
    faq = [
        ('页面无数据或数据不更新', '先核对平台地址、网络、服务进程与数据时间。硬件未上报时显示等待或离线；数据库旧记录和客户端示例数值不等于持续在线。'),
        ('照片分析能否确认疾病', '移动端现有照片分析是本地演示逻辑。Web健康结果也是风险筛查线索，疾病需要兽医结合现场检查及必要检测确认。'),
        ('AI助手连接失败或切换知识库', '进入服务设置测试平台连接。供应商密钥由后端管理；回答标识显示来自大模型或知识库。不要把平台访问凭证与DeepSeek密钥混填。'),
        ('咨询提交失败或没有专家回复', '检查标题与详情字数、网络和访问身份；失败保留草稿。当前仅保存咨询单、支持查询和撤回，未配置真人专家接单、排班和回复。'),
        ('资讯更新失败或新闻较旧', '查看来源同步状态和最后更新时间。后台保存原发布日期，不能把旧新闻重新标为今日新闻；断网时可阅读缓存和收藏。'),
        ('为什么关闭应用后收不到资讯通知', '当前实现应用内推荐、消息与未读提醒。鸿蒙Push Kit和微信订阅消息尚未配置，应用进程关闭后不会发送锁屏通知。'),
    ]
    if app['key'] == 'web':
        faq = [
            ('硬件离线但仍有环境数值', '环境和设备页面可能读取数据库中的历史记录。核对最后记录时间、实时链路、节点标识和现场供电，不把历史值视为持续在线。'),
            ('WebSocket已连接但没有新数据', '这只说明浏览器与平台的连接已建立。ESP32或MaixCAM仍需独立连接并上报，最新遥测时间不更新时先排查硬件链路。'),
            ('控制按钮不可用或设备没有动作', '离线或未建立控制链路时按钮可能禁用。命令成功提交也需要ACK和现场核对，未接线的雾化器、抓拍等不代表已经支持控制。'),
            ('图片或音频无法打开', '先检查原件是否归档和访问角色权限。事件元数据存在并不保证音频、图片文件可用，不应以占位图冒充现场证据。'),
            ('健康风险和模型标签如何解释', '平台展示辅助筛查及数据资格。当前牛只检测、公开声音基线和姿态模型不能自动确诊口蹄疫或跛行疾病，模型范围以实际版本和发布门禁为准。'),
            ('页面报错或接口被拒绝', '核对前后端地址、数据库连接与平台访问权限；生产启用角色密钥时使用被授权的读取和操作角色。')
        ]
    for title, answer in app.get('faq', faq):
        paragraph(document, title, bold=True, after=4)
        paragraph(document, answer, after=12)

    base.heading(document, '12 数据与版本说明', True)
    paragraph(document, '本手册与源程序材料均使用软件全称和V1.0标识。文档整理日期不替代实际开发完成或首次发表日期，权利信息由申请人填写。', after=12)
    if app['key'] == 'web':
        rows = [('业务数据', '后端数据库存储环境、设备、事件、风险和控制记录，图片与音频按部署目录保存。'),
                ('实时状态', 'WebSocket连接与硬件在线是不同状态。截图时硬件离线、暂无新帧，历史记录可用于查看页面结构。'),
                ('角色权限', '启用访问控制后，平台按viewer、operator、clinical、admin等角色限制操作。')]
    else:
        rows = [('本机记录', '照片演示记录、助手会话、咨询草稿、资讯缓存、收藏和订阅偏好保存在本机。'),
                ('服务端数据', '问答由平台调用供应商或知识库，咨询单持久化到后端，资讯由平台同步官方来源。'),
                ('身份范围', '当前咨询按客户端随机身份隔离，不含账号体系；清理数据或更换设备不会自动同步原列表。')]
    rows += [('截图来源', app.get('screenshot_origin', 'Web为本机浏览器，鸿蒙为真实HarmonyOS模拟器，小程序为微信开发者工具中的实际小程序页面。只添加红框编号，不篡改业务数据。')),
             ('资料保护', '公开演示资料不包含真实个人电话、供应商密钥或数据库口令；清空、卸载或清理缓存前保留必要记录。'),
             ('源码交存', f"提交{app['stats']['selectedLines']}行非空源程序，按50行每页编排，最后一页余数据实保留；原始行号和文件哈希另附清单。")]
    base.table(document, rows)
    paragraph(document, '资料对应关系', bold=True, after=6)
    paragraph(document, '源程序Word和PDF、当前源码归档、逐文件行数及SHA-256清单、带红框操作截图与流程图共同构成本版材料。新增功能与页面变更后需同步重新导出。', after=10)
    paragraph(document, '说明：本手册没有复用参考模板原系统的登录、需求发布等业务截图或操作说明。', size=9.5)
    output = folder / f"{app['name']}V1.0-用户操作手册.docx"
    document.save(output)
    return output


def main():
    entries = manifest()
    created = []
    for original in APPS:
        app = deepcopy(original)
        app['workflows'] = web_flows() if app['key'] == 'web' else mobile_flows(app['key'])
        if app['key'] == 'hm':
            app['summary'] = '鸿蒙移动巡检客户端，包含照片本地演示与记录、DeepSeek养殖问答、专家咨询单、新闻与每日小知识。'
        if app['key'] == 'wx':
            app['summary'] = '微信小程序提供巡检页面、照片本地演示与记录、AI问答、专家咨询单以及养牛新闻和每日知识。'
        folder = base.OUT / app['name']
        pages = base.prepare_source(app, folder)
        code = base.source_document(app, folder, pages)
        book = manual(app, folder, entries)
        base.auxiliary(app, folder)
        screenshot_keys = [key for flow in app['workflows'] for key, _ in flow['screenshots']]
        figure_manifest = [entries[key] for key in screenshot_keys]
        (folder / '截图来源与标注清单.json').write_text(json.dumps(figure_manifest, ensure_ascii=False, indent=2), 'utf-8')
        created.append(dict(name=app['name'], key=app['key'], stats=app['stats'], code=str(code), manual=str(book),
                            screenshotCount=len(screenshot_keys), workflowFigures=10))
        print(app['key'], 'source', app['stats']['selectedLines'], 'screenshots', len(screenshot_keys))
    (base.OUT / '材料生成清单.json').write_text(json.dumps(created, ensure_ascii=False, indent=2), 'utf-8')
    (base.OUT / '材料总览.md').write_text('# 牛项目三端软著材料\n\n本次按所提供Word模板更新源程序与操作手册。手册逐章包含真实界面截图、红框编号、步骤和操作结果。申请表保持空白。\n\n'
        + '\n'.join(f"- {e['name']}：代码{e['stats']['codePages']}页，手册按14页编排，10张流程图，{e['screenshotCount']}张操作截图。" for e in created)
        + '\n\n最终页数以Word分页与PDF导出检查.json为准。Web交存前1500和后1500行；APP和小程序不足3000行，提交各自全部现有非空源码。\n', 'utf-8')


if __name__ == '__main__':
    main()
