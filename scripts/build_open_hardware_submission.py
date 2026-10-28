# -*- coding: utf-8 -*-
"""Build the editable report and its original figures from verified content."""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from open_hardware_submission_content import DATE, MEMBERS, PAGES, SUBTITLE, TITLE


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "open_hardware_submission_20261026"
WORK = OUT / "work"
ASSETS = WORK / "assets"
STEM = "数牧增效-开源智能硬件解决方案-作品介绍文档"
FONT = Path("C:/Windows/Fonts/msyh.ttc")
BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")
COL = {"ink": "15221E", "green": "1D7054", "teal": "166F83", "grey": "F1F4F3", "line": "CBD9D3", "muted": "52625B", "gold": "A97820"}


def image_font(size, bold=False):
    return ImageFont.truetype(str(BOLD if bold else FONT), size)


def wrapped(draw, text, font, max_width):
    lines = []
    for paragraph in text.split("\n"):
        acc = ""
        for char in paragraph:
            if acc and draw.textlength(acc + char, font=font) > max_width:
                lines.append(acc)
                acc = char
            else:
                acc += char
        lines.append(acc)
    return lines


def img_text(draw, text, x, y, width, size=37, bold=False, color=None, align="left"):
    ft = image_font(size, bold)
    lines = wrapped(draw, text, ft, width)
    line_height = int(size * 1.42)
    for i, item in enumerate(lines):
        xx = x + (width - draw.textlength(item, font=ft)) / 2 if align == "center" else x
        draw.text((xx, y + i * line_height), item, font=ft, fill="#" + (color or COL["ink"]), anchor="lt")
    return len(lines) * line_height


def node(draw, xy, title, body, accent="green", title_size=42, body_size=34):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=16, fill="#FFFFFF", outline="#" + COL[accent], width=4)
    img_text(draw, title, x1 + 27, y1 + 24, x2 - x1 - 54, title_size, True, COL[accent])
    img_text(draw, body, x1 + 27, y1 + 88, x2 - x1 - 54, body_size)


def arrow(draw, coords, color="green", end=True):
    draw.line(coords, fill="#" + COL[color], width=6, joint="curve")
    if not end:
        return
    x, y = coords[-1]
    px, py = coords[-2]
    angle = math.atan2(y - py, x - px)
    draw.polygon([(x, y), (x - 23 * math.cos(angle - .5), y - 23 * math.sin(angle - .5)), (x - 23 * math.cos(angle + .5), y - 23 * math.sin(angle + .5))], fill="#" + COL[color])


def figures():
    ASSETS.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", (1800, 1000), "#F4F7F5")
    d = ImageDraw.Draw(im)
    img_text(d, "牛舍现场", 55, 32, 800, 42, True)
    img_text(d, "平台与使用者", 1040, 32, 700, 42, True)
    node(d, (50, 105, 850, 315), "环境控制节点", "ESP32-S3  温湿度与光照  可选气体\n本地阈值控制与风扇声光反馈")
    node(d, (50, 380, 850, 590), "牛身佩戴节点", "ESP32-S3  颈背腿运动  INMP441\n体表趋势  有效窗口与模块健康", "teal")
    node(d, (50, 655, 850, 865), "MaixCam 视觉节点", "牛只检测  事件与关键截图\n模型版本  身份关联与证据归档")
    node(d, (1040, 230, 1750, 520), "统一接入与数据服务", "Spring Boot  数据库存储\nREST 与 WebSocket\n证据原件及 SHA-256", "teal")
    node(d, (1040, 640, 1750, 890), "管理与复核", "Web 与移动端  控制回执\n人工检查  审核与发布记录")
    arrow(d, [(850, 210), (930, 210), (930, 310), (1035, 310)])
    arrow(d, [(850, 485), (970, 485), (970, 390), (1035, 390)], "teal")
    arrow(d, [(850, 760), (930, 760), (930, 470), (1035, 470)])
    arrow(d, [(1395, 525), (1395, 635)], "teal")
    img_text(d, "采集与事件上行  /  控制命令与回执双向核对", 150, 935, 1500, 33, color=COL["muted"], align="center")
    im.save(ASSETS / "architecture.png")

    im = Image.new("RGB", (1800, 720), "#F4F7F5")
    d = ImageDraw.Draw(im)
    node(d, (55, 50, 565, 240), "正常采集", "核对 ID 与配置\n累计真实有效样本")
    node(d, (645, 50, 1155, 240), "发现缺件或读取异常", "对应位置标为不可用\n不重启其他正常总线", "gold", 38)
    node(d, (1235, 50, 1745, 240), "终止旧窗口", "清空连续统计\n原始会话记录失败", "teal")
    arrow(d, [(570, 145), (640, 145)])
    arrow(d, [(1160, 145), (1230, 145)], "teal")
    arrow(d, [(1490, 245), (1490, 365)], "teal")
    node(d, (645, 370, 1745, 570), "模块恢复后重新开始", "自检通过后重收完整窗口  20 Hz × 10 秒 = 200 点\n首帧不与旧帧求差  不将中断拼接为连续数据", "teal", 40, 33)
    arrow(d, [(640, 470), (310, 470), (310, 245)])
    img_text(d, "软件验证覆盖故障与恢复逻辑  实际信号与佩戴质量另行验收", 120, 630, 1560, 35, color=COL["muted"], align="center")
    im.save(ASSETS / "reliability.png")

    im = Image.new("RGB", (1800, 860), "#F4F7F5")
    d = ImageDraw.Draw(im)
    img_text(d, "现有牛只检测路线", 55, 30, 1600, 39, True)
    node(d, (55, 110, 555, 320), "现场画面", "MaixCam 图像采集\n基础画面稳定后再推理", title_size=40, body_size=34)
    node(d, (650, 110, 1150, 320), "YOLOv8n COCO", "检测 cow 与位置\n不直接解释疾病", title_size=40, body_size=34)
    node(d, (1245, 110, 1745, 320), "平台事件", "截图与版本归档\n供人员回看核对", "teal", 40, 34)
    arrow(d, [(560, 215), (645, 215)])
    arrow(d, [(1155, 215), (1240, 215)], "teal")
    img_text(d, "独立研究分类路线", 55, 395, 1600, 39, True)
    node(d, (55, 480, 555, 730), "公开外观标签", "去重与图像组分区\n隔离 413 张留出测试", "teal", 40, 34)
    node(d, (650, 480, 1150, 730), "YOLOv8n-cls", "92.01% 来源标签准确率\nONNX 数值校验通过", "teal", 40, 32)
    node(d, (1245, 480, 1745, 730), "设备部署待完成", "NPU 转换与量化核对\n实机加载  时延与画面验收", "gold", 38, 32)
    arrow(d, [(560, 605), (645, 605)], "teal")
    arrow(d, [(1155, 605), (1240, 605)], "gold")
    img_text(d, "研究标签与疾病结论分开管理  两条路线各自验收", 90, 790, 1620, 34, color=COL["muted"], align="center")
    im.save(ASSETS / "vision.png")


def set_font(run, size=None, bold=None, color="000000", face="Microsoft YaHei"):
    run.font.name = face
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + attr), face)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def field(paragraph, instruction, cached="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = cached
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    set_font(run, 9, color="59645E")


def paragraph(doc, text, style=None):
    para = doc.add_paragraph(style=style)
    set_font(para.add_run(text))
    return para


def make_table(doc, headers, records, widths=None):
    widths = widths or [6.8 / len(headers)] * len(headers)
    total = sum(widths)
    widths = [v * 6.8 / total for v in widths]
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    tw = pr.find(qn("w:tblW"))
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int(6.8 * 1440)))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:" + side)
        for k, v in {"val": "single", "sz": "5", "color": "D9D9D9"}.items():
            b.set(qn("w:" + k), v)
        borders.append(b)
    pr.append(borders)
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for ridx, values in enumerate([headers] + records):
        row = t.rows[0] if ridx == 0 else t.add_row()
        rp = row._tr.get_or_add_trPr()
        rp.append(OxmlElement("w:cantSplit"))
        if ridx == 0:
            rp.append(OxmlElement("w:tblHeader"))
        for idx, val in enumerate(values):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "34443C" if ridx == 0 else ("F3F5F4" if ridx % 2 == 0 else "FFFFFF"))
            shade.set(qn("w:val"), "clear")
            cp.append(shade)
            margins = OxmlElement("w:tcMar")
            for side, v in (("top", 90), ("bottom", 90), ("left", 110), ("right", 110)):
                m = OxmlElement("w:" + side)
                m.set(qn("w:w"), str(v))
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            cp.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.line_spacing = Pt(14.5)
            if idx == 0 or len(str(val)) < 13:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(para.add_run(str(val)), 10, ridx == 0, "FFFFFF" if ridx == 0 else "18231E")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    set_font(spacer.add_run(""), 2)
    return t


def evidence():
    files = [
        "README.md", "API.md", "smart-cattle-backend/pom.xml", "smart-cattle-frontend/package.json",
        "硬件代码/README.md", "硬件代码/牛身节点/README.md", "硬件代码/牛身节点/src/main.cpp",
        "硬件代码/MaixCam节点/README.md", "硬件代码/MaixCam节点/models/MODEL_MANIFEST.txt",
        "docs/2026-10-26-hardware-integration-status.md", "docs/2026-10-26-vision-yolo-training.md",
        "docs/2026-10-26-calf-cough-research.md", "docs/2026-10-25-imu-data-evidence.md",
        "docs/2026-10-25-audio-data-evidence.md", "疾病识别生产落地/生产部署与审计清单.md",
        "疾病识别生产落地/疾病实体与证据字典.md", "数牧增效-硬件采购清单5.xlsx",
    ]
    sources = []
    for path in files:
        f = ROOT / path
        sources.append({"path": path, "exists": f.exists(), "sha256": hashlib.sha256(f.read_bytes()).hexdigest() if f.exists() else None})
    reports = []
    for path in sorted((ROOT / "smart-cattle-backend/target/surefire-reports").glob("TEST-*.xml")):
        el = ET.parse(path).getroot()
        reports.append({k: el.attrib.get(k) for k in ("name", "tests", "errors", "failures", "skipped")})
    result = {
        "date": "2026-10-26", "project": TITLE, "members": MEMBERS,
        "scope": "Report preparation and nonhardware verification; no serial access or firmware upload",
        "sources": sources, "backend_surefire": reports,
        "checks": [
            {"command": "mvn -o -B test", "result": "57 passed, 0 failures, 0 errors, 0 skipped"},
            {"command": "npm run build", "result": "passed, Vite 6.4.2"},
            {"command": "python scripts/test_maixcam_lifecycle.py", "result": "4 passed"},
            {"command": "python scripts/test_maixcam_uart_console.py", "result": "4 passed"},
            {"command": "python scripts/train_calf_cough_research_20261026.py self-test", "result": "17 passed"},
            {"command": "python scripts/test_cow_sensor_recovery.py", "result": "6 host scenarios passed"},
            {"command": "python 硬件代码/MaixCam节点/tools/test_vision_logic.py", "result": "failed during import: mock maix module missing app; event assertions not executed"},
        ],
        "deferred": ["field deployment", "clinical validation", "mobile packaging", "public repository URL", "demo video URL"],
    }
    (WORK / "evidence.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def main(pages=None, out_dir=None, stem=None, illustrated=False):
    global OUT, WORK, ASSETS, STEM
    if out_dir is not None:
        OUT = Path(out_dir)
        WORK = OUT / "work"
        ASSETS = WORK / "assets"
    if stem is not None:
        STEM = stem
    pages = PAGES if pages is None else pages
    WORK.mkdir(parents=True, exist_ok=True)
    if not illustrated:
        figures()
    evidence()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.70)
    sec.left_margin = Inches(.8)
    sec.right_margin = Inches(.8)
    sec.header_distance = Inches(.25)
    sec.footer_distance = Inches(.3)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = Pt(17.2)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 32), ("Subtitle", 21), ("Heading 1", 18), ("Heading 2", 12), ("Heading 3", 11)]:
        st = doc.styles[name]
        st.font.name = "Microsoft YaHei"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.bold = name != "Subtitle"
        st.font.italic = False
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        st.paragraph_format.space_before = Pt(11 if name == "Heading 2" else 0)
        st.paragraph_format.space_after = Pt(6 if name == "Heading 2" else 16)
        st.paragraph_format.line_spacing = Pt(size * 1.5)
        st.paragraph_format.keep_with_next = True
        for border in list(st.element.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    for name in ("TOC 1", "TOC 2"):
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        st = doc.styles[name]
        st.font.name = "Microsoft YaHei"
        st.font.size = Pt(10.5)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = Pt(15)
    if "Hyperlink" in doc.styles:
        doc.styles["Hyperlink"].font.color.rgb = RGBColor(0, 0, 0)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("数牧增效  |  "), 8, color="59645E")
    field(footer, "PAGE")
    set_font(footer.add_run(" / "), 8, color="59645E")
    field(footer, "NUMPAGES")

    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(30)
    lead.paragraph_format.space_after = Pt(24)
    set_font(lead.add_run("开源智能硬件解决方案"), 12, color="4C6257")
    para = doc.add_paragraph(style="Title")
    set_font(para.add_run(TITLE), 34, True)
    for part in ("开源多模态智能牛舍监测", "与闭环预警系统"):
        para = doc.add_paragraph(style="Subtitle")
        para.paragraph_format.space_after = Pt(8)
        run = para.add_run(part)
        set_font(run, 24)
        run.italic = False
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(22)
    set_font(para.add_run("作品介绍文档"), 18, True)
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(15)
    set_font(para.add_run("面向中小牧场的环境与牛体监测方案，整合边缘视觉、运动和声音证据，支持本地控制、人工复核及可追溯的模型发布。"), 12)
    if illustrated:
        para = doc.add_paragraph()
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(5)
        para.add_run().add_picture(str(ASSETS / "cover_overview.png"), width=Inches(6.7))
    for text in ("团队成员", MEMBERS, "文档版本  " + ("2.0 图文详解版" if illustrated else "1.0"), "编制日期  " + DATE):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.space_before = Pt(12 if text == "团队成员" else 0)
        set_font(para.add_run(text), 11)
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(24)
    set_font(para.add_run("技术创新 30%    场景落地 30%    开源治理 20%    长期发展 20%"), 10, color="52625B")

    doc.add_page_break()
    para = doc.add_paragraph(style="Title")
    set_font(para.add_run("目录"), 24, True)
    paragraph(doc, "正文按系统实现、应用验证、开源治理与持续发展组织，附录提供运行方法、视频分镜和提交信息。")
    toc = doc.add_paragraph()
    field(toc, 'TOC \\o "1-1" \\h \\z \\u', "目录将在排版时更新")

    markdown = ["# " + TITLE, "## " + SUBTITLE, "团队成员：" + MEMBERS, "日期：" + DATE, ""]
    figure_index = 0
    figure_catalog = []
    for title, blocks in pages:
        para = doc.add_paragraph(style="Heading 1")
        para.paragraph_format.page_break_before = True
        para.paragraph_format.keep_with_next = True
        set_font(para.add_run(title), 18, True)
        markdown.extend(["", "# " + title, ""])
        for block in blocks:
            kind = block[0]
            if kind == "p":
                paragraph(doc, block[1])
                markdown.extend([block[1], ""])
            elif kind == "h":
                para = doc.add_paragraph(style="Heading 2")
                set_font(para.add_run(block[1]), 12, True)
                markdown.extend(["## " + block[1], ""])
            elif kind == "break_h":
                para = doc.add_paragraph(style="Heading 2")
                para.paragraph_format.page_break_before = True
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(13)
                set_font(para.add_run(block[1]), 15, True)
                markdown.extend(["## " + block[1], ""])
            elif kind == "table":
                make_table(doc, block[1], block[2], block[3])
                markdown.extend(["| " + " | ".join(block[1]) + " |", "| " + " | ".join(["---"] * len(block[1])) + " |"])
                markdown.extend("| " + " | ".join(str(v) for v in row) + " |" for row in block[2])
                markdown.append("")
            elif kind == "fig":
                figure_index += 1
                description = block[2]
                if illustrated:
                    if description.startswith("图 "):
                        description = description.split(" ", 2)[2]
                    description = f"图 {figure_index}  {description}"
                para = doc.add_paragraph()
                para.paragraph_format.keep_with_next = True
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(3)
                # Inline pictures must use automatic line height, not body text's exact height.
                para.paragraph_format.line_spacing = 1.0
                r = para.add_run()
                inline = r.add_picture(str(ASSETS / (block[1] + ".png")), width=Inches(block[3]))
                inline._inline.docPr.set("descr", description)
                caption = doc.add_paragraph()
                caption.paragraph_format.space_after = Pt(10)
                caption.paragraph_format.line_spacing = Pt(13)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(caption.add_run(description), 9, color="52625B")
                markdown.extend(["![" + description + "](work/assets/" + block[1] + ".png)", ""])
                figure_catalog.append({"number": figure_index, "asset": block[1], "caption": description, "chapter": title})
            elif kind == "code":
                for text in block[1]:
                    para = doc.add_paragraph()
                    para.paragraph_format.space_after = Pt(1)
                    para.paragraph_format.line_spacing = Pt(13)
                    set_font(para.add_run(text), 9.5, face="Microsoft YaHei")
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(3)
                spacer.paragraph_format.line_spacing = Pt(3)
                markdown.extend(["```powershell", *block[1], "```", ""])
    doc.core_properties.title = TITLE + SUBTITLE + "作品介绍文档"
    doc.core_properties.author = "数牧增效项目团队"
    doc.core_properties.subject = "开源智能硬件解决方案参赛作品介绍"
    doc.core_properties.keywords = "智能牛舍 物联网 边缘计算 开源硬件 多模态"
    doc.core_properties.comments = ""
    doc.save(OUT / (STEM + ".docx"))
    (OUT / "作品介绍文档正文.md").write_text("\n".join(markdown), encoding="utf-8")
    content_length = sum(len(str(b)) for _, blocks in pages for b in blocks)
    (WORK / "figure-catalog.json").write_text(json.dumps(figure_catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(OUT / (STEM + ".docx")), "chapters": len(pages), "content_characters": content_length, "figures": figure_index + (1 if illustrated else 0)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
