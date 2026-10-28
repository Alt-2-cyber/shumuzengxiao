from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document

from cow_source_template_format import REFERENCE, write_source_document

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '.run' / 'source-template-format-20261023'
OUT = Path(r'D:\table\软著申请\牛')


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    reports = []
    for source in sorted(OUT.glob('*/*V1.0-源程序.docx')):
        if source.name.startswith('~$'):
            continue
        folder = source.parent
        listing = json.loads((folder / '源程序文件与行号清单.json').read_text('utf-8'))
        selected = listing['selected']
        expected = [row['text'].expandtabs(2) for row in selected]
        actual = [paragraph.text for paragraph in Document(source).paragraphs]
        assert actual in (expected, [f"{row['index']:05d} | {row['text'].expandtabs(2)}" for row in selected]), source
        backup = WORK / '修改前备份' / folder.name
        backup.mkdir(parents=True, exist_ok=True)
        for path in (source, source.with_suffix('.pdf')):
            destination = backup / path.name
            if not destination.exists():
                shutil.copy2(path, destination)
        pages = [selected[index:index+50] for index in range(0, len(selected), 50)]
        staged = WORK / '新版源码' / folder.name / source.name
        write_source_document(folder.name, staged, pages)
        assert [paragraph.text for paragraph in Document(staged).paragraphs] == expected
        reports.append(dict(name=folder.name, sourceLines=len(selected), expectedPages=len(pages),
                            staged=str(staged), target=str(source), inlineNumbersRemoved=True,
                            automaticMarginLineNumbers=True, reference=str(REFERENCE)))
    assert len(reports) == 3
    if not (WORK / '参考代码模板.docx').exists():
        shutil.copy2(REFERENCE, WORK / '参考代码模板.docx')
    (WORK / '待检查清单.json').write_text(json.dumps(reports, ensure_ascii=False, indent=2), 'utf-8')
    print(json.dumps(reports, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
