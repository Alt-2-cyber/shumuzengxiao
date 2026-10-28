from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree as ET

ROOT = Path(r'D:\table\软著申请\牛')
WORK = Path(__file__).resolve().parents[1] / '.run' / 'application-form-fill-20261024'


def main() -> None:
    reports = json.loads((WORK / '填写检查.json').read_text('utf-8'))
    for report in reports:
        path = Path(report['path'])
        document = Document(path)
        table = document.tables[0]
        assert table.cell(2, 1).text.strip() == report['software']
        assert table.cell(4, 1).text.strip() == 'V1.0'
        assert table.cell(5, 1).text.strip() == '1 应用软件'
        assert len(table.cell(21, 1).text.strip()) == report['featureCharacters']
        assert len(table.cell(21, 1).text.strip()) >= 501
        assert len(table.cell(22, 1)._tc.xpath('.//w:drawing')) == 1
        assert '请输入详细的技术特点' not in table.cell(24, 1).text
        assert '标签选至少一个' not in table.cell(23, 1).text
        for row in report['identityAndContactRowsBlank']:
            assert table.cell(row, 1).text.strip() == '', (path, row, table.cell(row, 1).text)
        with ZipFile(path) as archive:
            assert archive.testzip() is None
            for name in archive.namelist():
                if name.endswith('.xml'):
                    ET.fromstring(archive.read(name))
        print(report['software'], 'verified', report['featureCharacters'], 'feature chars; project fields filled; personal fields blank')
    print('All three application forms passed.')


if __name__ == '__main__':
    main()
